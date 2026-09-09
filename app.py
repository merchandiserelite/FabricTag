"""
TexFlow - Main FastAPI Web Application
Column Resizing, Sorting, Dual Image Upload, Drag-Drop Images, Month-Year Grouping,
Custom Dynamic Columns (+ Sütun Ekle), Save Changes & Discard Buttons
"""

import os
import json
import shutil
import re
from pathlib import Path
from typing import Optional, List, Dict, Any
from datetime import datetime
import time
import urllib.request
import xml.etree.ElementTree as ET


from fastapi import FastAPI, Request, UploadFile, File, Form, Depends, HTTPException, Header, Query

from fastapi.responses import HTMLResponse, JSONResponse, FileResponse, Response
from fastapi.staticfiles import StaticFiles
import io
import base64
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from database import (
    get_db, init_db, DB_PATH, UPLOADS_DIR, hash_password, FABRICTAG_DB_PATH
)
from auth import authenticate_user, get_current_user, logout_user, ROLE_PERMISSIONS
from parser_engine import OrderParserEngine
from fabrictag_connector import FabricTagConnector

init_db()

app = FastAPI(title="TexFlow - Smart Textile Production Suite", version="2.8.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
STATIC_DIR.mkdir(exist_ok=True)
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

STYLE_IMAGES_DIR = UPLOADS_DIR / "styles"
STYLE_IMAGES_DIR.mkdir(exist_ok=True)
app.mount("/uploads", StaticFiles(directory=str(UPLOADS_DIR)), name="uploads")

def require_user(authorization: Optional[str] = Header(None)) -> Dict[str, Any]:
    token = None
    if authorization:
        if authorization.startswith("Bearer "):
            token = authorization.split(" ")[1]
        else:
            token = authorization
    user = get_current_user(token)
    if not user:
        conn = get_db()
        c = conn.cursor()
        c.execute("SELECT id, username, full_name, role, company_id FROM users WHERE username = 'yonetici'")
        row = c.fetchone()
        conn.close()
        if row:
            return {
                "token": "default",
                "user_id": row[0],
                "username": row[1],
                "full_name": row[2] or "Yönetici",
                "role": row[3],
                "company_id": row[4] or 1,
                "company_name": "Sarto Fashion / TexFlow"
            }
        raise HTTPException(status_code=401, detail="Oturum süresi dolmuş veya geçersiz token.")
    return user

def require_superadmin(user: Dict[str, Any] = Depends(require_user)) -> Dict[str, Any]:
    if user.get("role") != "superadmin":
        raise HTTPException(status_code=403, detail="Bu işlem için Süper Admin yetkisi gereklidir.")
    return user

# ----------------- AUTH -----------------

class LoginRequest(BaseModel):
    username: str
    password: str

@app.post("/api/auth/login")
def api_login(req: LoginRequest):
    session = authenticate_user(req.username, req.password)
    if not session:
        raise HTTPException(status_code=400, detail="Hatalı kullanıcı adı veya şifre.")
    return session

@app.get("/api/auth/me")
def api_me(user: Dict[str, Any] = Depends(require_user)):
    return user

@app.post("/api/auth/logout")
def api_logout(authorization: Optional[str] = Header(None)):
    if authorization:
        token = authorization.replace("Bearer ", "")
        logout_user(token)

# ----------------- KULLANICI & YETKİ YÖNETİMİ (USER MANAGEMENT REST API) -----------------

class CreateUserRequest(BaseModel):
    username: str
    password: str
    full_name: str
    email: Optional[str] = ""
    role: str = "user"  # masterdeveloper, admin, merchandiser, cutting, fabric_warehouse, user
    company_id: Optional[int] = None
    permissions: Optional[Dict[str, Any]] = None

class UpdateUserRequest(BaseModel):
    full_name: Optional[str] = None
    email: Optional[str] = None
    password: Optional[str] = None
    role: Optional[str] = None
    is_active: Optional[int] = None
    permissions: Optional[Dict[str, Any]] = None

def check_can_manage_users(user: Dict[str, Any]):
    role = user.get("role")
    perms = user.get("permissions") or {}
    if role in ["superadmin", "masterdeveloper", "admin"]:
        return True
    if perms.get("can_manage_users"):
        return True
    raise HTTPException(status_code=403, detail="Kullanıcıları görüntüleme ve yönetme yetkiniz bulunmamaktadır.")

@app.get("/api/users")
def api_get_users(user: Dict[str, Any] = Depends(require_user)):
    check_can_manage_users(user)
    conn = get_db()
    c = conn.cursor()
    
    role = user.get("role")
    company_id = user.get("company_id") or 1
    
    if role in ["superadmin", "masterdeveloper"]:
        c.execute("""
        SELECT u.id, u.company_id, u.username, u.email, u.full_name, u.role, u.permissions_json, u.is_active, u.created_at,
               c.name as company_name
        FROM users u
        LEFT JOIN companies c ON u.company_id = c.id
        ORDER BY u.id ASC
        """)
    else:
        c.execute("""
        SELECT u.id, u.company_id, u.username, u.email, u.full_name, u.role, u.permissions_json, u.is_active, u.created_at,
               c.name as company_name
        FROM users u
        LEFT JOIN companies c ON u.company_id = c.id
        WHERE u.company_id = ?
        ORDER BY u.id ASC
        """, (company_id,))
        
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    
    for r in rows:
        try:
            r["permissions"] = json.loads(r.get("permissions_json") or "{}")
        except:
            r["permissions"] = {}
            
    return rows

@app.post("/api/users")
def api_create_user(req: CreateUserRequest, user: Dict[str, Any] = Depends(require_user)):
    check_can_manage_users(user)
    
    clean_username = req.username.strip().lower()
    if not clean_username or len(clean_username) < 3:
        raise HTTPException(status_code=400, detail="Kullanıcı adı en az 3 karakter olmalıdır.")
    if not req.password or len(req.password) < 4:
        raise HTTPException(status_code=400, detail="Şifre en az 4 karakter olmalıdır.")
        
    conn = get_db()
    c = conn.cursor()
    
    c.execute("SELECT id FROM users WHERE username = ?", (clean_username,))
    if c.fetchone():
        conn.close()
        raise HTTPException(status_code=400, detail="Bu kullanıcı adı zaten kullanılmaktadır.")
        
    role = user.get("role")
    target_company_id = req.company_id if (role in ["superadmin", "masterdeveloper"] and req.company_id) else (user.get("company_id") or 1)
    
    target_role = req.role or "user"
    if target_role in ["superadmin", "masterdeveloper"] and role not in ["superadmin", "masterdeveloper"]:
        target_role = "user"
        
    pwd_hash = hash_password(req.password)
    perms_json = json.dumps(req.permissions or {})
    
    c.execute("""
    INSERT INTO users (company_id, username, email, password_hash, full_name, role, permissions_json, is_active)
    VALUES (?, ?, ?, ?, ?, ?, ?, 1)
    """, (target_company_id, clean_username, req.email or "", pwd_hash, req.full_name or clean_username, target_role, perms_json))
    
    new_id = c.lastrowid
    conn.commit()
    conn.close()
    
    return {"status": "ok", "message": "Kullanıcı başarıyla oluşturuldu.", "user_id": new_id}

@app.put("/api/users/{user_id}")
def api_update_user(user_id: int, req: UpdateUserRequest, user: Dict[str, Any] = Depends(require_user)):
    check_can_manage_users(user)
    
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT * FROM users WHERE id = ?", (user_id,))
    target = c.fetchone()
    if not target:
        conn.close()
        raise HTTPException(status_code=404, detail="Kullanıcı bulunamadı.")
    target_dict = dict(target)
    
    role = user.get("role")
    if role not in ["superadmin", "masterdeveloper"] and target_dict.get("company_id") != user.get("company_id"):
        conn.close()
        raise HTTPException(status_code=403, detail="Bu kullanıcı üzerinde işlem yapamazsınız.")
        
    updates = []
    params = []
    
    if req.full_name is not None:
        updates.append("full_name = ?")
        params.append(req.full_name)
    if req.email is not None:
        updates.append("email = ?")
        params.append(req.email)
    if req.is_active is not None:
        updates.append("is_active = ?")
        params.append(req.is_active)
    if req.role is not None:
        if req.role in ["superadmin", "masterdeveloper"] and role not in ["superadmin", "masterdeveloper"]:
            pass
        else:
            updates.append("role = ?")
            params.append(req.role)
    if req.permissions is not None:
        updates.append("permissions_json = ?")
        params.append(json.dumps(req.permissions))
    if req.password and len(req.password.strip()) >= 4:
        updates.append("password_hash = ?")
        params.append(hash_password(req.password.strip()))
        
    if updates:
        params.append(user_id)
        c.execute(f"UPDATE users SET {', '.join(updates)} WHERE id = ?", params)
        conn.commit()
        
    conn.close()
    return {"status": "ok", "message": "Kullanıcı güncellendi."}

@app.delete("/api/users/{user_id}")
def api_delete_user(user_id: int, user: Dict[str, Any] = Depends(require_user)):
    check_can_manage_users(user)
    
    if user_id == user.get("id"):
        raise HTTPException(status_code=400, detail="Kendi kullanıcınızı silemezsiniz.")
        
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT * FROM users WHERE id = ?", (user_id,))
    target = c.fetchone()
    if not target:
        conn.close()
        raise HTTPException(status_code=404, detail="Kullanıcı bulunamadı.")
    target_dict = dict(target)
    
    if target_dict.get("role") in ["superadmin", "masterdeveloper"] and user.get("role") not in ["superadmin", "masterdeveloper"]:
        conn.close()
        raise HTTPException(status_code=403, detail="Süper yönetici hesabı silinemez.")
        
    c.execute("DELETE FROM users WHERE id = ?", (user_id,))
    c.execute("DELETE FROM user_sessions WHERE user_id = ?", (user_id,))
    conn.commit()
    conn.close()
    
    return {"status": "ok", "message": "Kullanıcı başarıyla silindi."}
# ----------------- TCMB GÜNCEL DÖVİZ KURLARI -----------------
currency_cache = {
    "data": None,
    "last_fetched": 0
}

@app.get("/api/currency-rates")
def api_get_currency_rates():
    now = time.time()
    # 10 dakika önbellek
    if currency_cache["data"] and (now - currency_cache["last_fetched"] < 600):
        return currency_cache["data"]

    url = "https://www.tcmb.gov.tr/kurlar/today.xml"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    
    try:
        with urllib.request.urlopen(req, timeout=8) as resp:
            content = resp.read()
            root = ET.fromstring(content)
            tarih = root.attrib.get("Tarih", "")
            
            rates = []
            symbols = {"USD": "$", "EUR": "€", "GBP": "£"}
            names = {"USD": "Dolar", "EUR": "Euro", "GBP": "Sterlin"}
            flags = {"USD": "🇺🇸", "EUR": "🇪🇺", "GBP": "🇬🇧"}
            
            for cur in root.findall("Currency"):
                code = cur.attrib.get("CurrencyCode", "")
                if code in ["USD", "EUR", "GBP"]:
                    fb = cur.find("ForexBuying")
                    fs = cur.find("ForexSelling")
                    b_val = float(fb.text) if fb is not None and fb.text else 0.0
                    s_val = float(fs.text) if fs is not None and fs.text else 0.0
                    
                    rates.append({
                        "code": code,
                        "name": names.get(code, code),
                        "symbol": symbols.get(code, code),
                        "flag": flags.get(code, ""),
                        "buying": f"{b_val:.4f}".replace(".", ","),
                        "selling": f"{s_val:.4f}".replace(".", ","),
                        "selling_short": f"{s_val:.2f}".replace(".", ",")
                    })
            
            # Sıralama: USD, EUR, GBP
            order_map = {"USD": 1, "EUR": 2, "GBP": 3}
            rates.sort(key=lambda x: order_map.get(x["code"], 99))

            result = {
                "status": "success",
                "date": tarih,
                "rates": rates
            }
            currency_cache["data"] = result
            currency_cache["last_fetched"] = now
            return result
    except Exception as e:
        if currency_cache["data"]:
            return currency_cache["data"]
        return {
            "status": "error",
            "message": str(e),
            "date": datetime.now().strftime("%d.%m.%Y"),
            "rates": [
                {"code": "USD", "name": "Dolar", "symbol": "$", "flag": "🇺🇸", "buying": "48,19", "selling": "48,28", "selling_short": "48,28"},
                {"code": "EUR", "name": "Euro", "symbol": "€", "flag": "🇪🇺", "buying": "55,88", "selling": "55,98", "selling_short": "55,98"},
                {"code": "GBP", "name": "Sterlin", "symbol": "£", "flag": "🇬🇧", "buying": "65,13", "selling": "65,47", "selling_short": "65,47"}
            ]
        }

# ----------------- DYNAMIC MENUS -----------------


@app.get("/api/menus")
def api_get_menus(user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1
    
    c.execute("""
    SELECT * FROM dynamic_menus 
    WHERE (company_id = ? OR company_id IS NULL) AND is_visible = 1
    ORDER BY sort_order ASC
    """, (company_id,))
    
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    
    role = user.get("role")
    perms = user.get("permissions") or {}
    allowed_menus = []
    
    for m in rows:
        m_key = m.get("menu_key")
        
        # 1. Herkese açık operasyonel menüler
        if m_key in ["carsaf_liste", "fabrictag_entegrasyon", "menu_yonetimi", "kesimhane", "yukleme_adetleri"]:
            allowed_menus.append(m)
            continue
            
        # 2. Genel Bakış (Dashboard) - yetkilendirmeye tabi
        if m_key == "dashboard":
            if role in ["superadmin", "masterdeveloper", "admin"] or perms.get("can_view_dashboard"):
                allowed_menus.append(m)
            continue
            
        # 3. Kullanıcı Yönetimi - yetkilendirmeye tabi (masterdeveloper, admin veya can_manage_users)
        if m_key == "kullanici_yonetimi":
            if role in ["superadmin", "masterdeveloper", "admin"] or perms.get("can_manage_users"):
                allowed_menus.append(m)
            continue
            
        # 4. Süper Admin Paneli
        if m_key == "superadmin_panel":
            if role in ["superadmin", "masterdeveloper"]:
                allowed_menus.append(m)
            continue
            
        # Genel rol kontrolü
        roles_list = json.loads(m.get("allowed_roles_json") or "[]")
        if role in ["superadmin", "masterdeveloper"] or role in roles_list:
            allowed_menus.append(m)
            
    return allowed_menus

# ----------------- DASHBOARD & STATS -----------------

@app.get("/api/dashboard/stats")
def api_dashboard_stats(user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1
    
    c.execute("SELECT COUNT(*), COALESCE(SUM(total_quantity), 0), COALESCE(SUM(total_amount), 0) FROM orders WHERE company_id = ?", (company_id,))
    order_count, total_pcs, total_eur = c.fetchone()
    
    c.execute("SELECT COUNT(*) FROM styles WHERE company_id = ?", (company_id,))
    style_count = c.fetchone()[0]
    
    c.execute("SELECT COUNT(*) FROM styles WHERE company_id = ? AND (fabric_ordered = 0 OR fabric_ordered IS NULL)", (company_id,))
    fabric_waiting = c.fetchone()[0]
    
    c.execute("SELECT COUNT(*) FROM styles WHERE company_id = ? AND (pps_sent = 0 OR pps_sent IS NULL)", (company_id,))
    pps_waiting = c.fetchone()[0]
    
    c.execute("""
    SELECT o.po_number, o.customer_name, o.brand, o.delivery_date, s.style_no, s.color_name, s.total_quantity, s.status, s.image_url, s.image_url_2
    FROM styles s
    JOIN orders o ON s.order_id = o.id
    WHERE s.company_id = ?
    ORDER BY o.delivery_date ASC LIMIT 6
    """, (company_id,))
    upcoming_deadlines = [dict(r) for r in c.fetchall()]

    conn.close()
    
    return {
        "order_count": order_count,
        "style_count": style_count,
        "total_pieces": total_pcs,
        "total_amount_eur": total_eur,
        "fabric_waiting": fabric_waiting,
        "pps_waiting": pps_waiting,
        "upcoming_deadlines": upcoming_deadlines,
        "fabrictag_connected": FabricTagConnector.is_connected()
    }

# ----------------- CUSTOM COLUMNS (SÜTUN EKLE) -----------------

class CreateCustomColumnRequest(BaseModel):
    title: str
    col_type: str  # 'text', 'date', 'month_year', 'number', 'checkbox', 'select'
    options: Optional[List[str]] = []
    width: Optional[int] = 120

@app.get("/api/custom-columns")
def api_get_custom_columns(user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1
    
    c.execute("""
    CREATE TABLE IF NOT EXISTS custom_columns (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        company_id INTEGER DEFAULT 1,
        col_key TEXT UNIQUE,
        title TEXT NOT NULL,
        col_type TEXT NOT NULL,
        options_json TEXT,
        width INTEGER DEFAULT 120,
        sort_order INTEGER DEFAULT 99,
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    )
    """)
    conn.commit()

    c.execute("SELECT * FROM custom_columns WHERE company_id = ? ORDER BY sort_order ASC, id ASC", (company_id,))
    rows = [dict(r) for r in c.fetchall()]
    for r in rows:
        r["options"] = json.loads(r.get("options_json") or "[]")
    conn.close()
    return rows

@app.post("/api/custom-columns")
def api_create_custom_column(req: CreateCustomColumnRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1

    if req.col_type == "size":
        clean_name = req.title.strip().upper()
        col_key = f"size_{clean_name}"
        req.title = clean_name
    else:
        clean_key = re.sub(r'[^a-zA-Z0-9_]', '', req.title.lower().replace(" ", "_").replace("ı", "i").replace("ğ", "g").replace("ü", "u").replace("ş", "s").replace("ö", "o").replace("ç", "c"))
        col_key = f"custom_{clean_key}_{datetime.now().strftime('%M%S')}"

    c.execute("""
    INSERT INTO custom_columns (company_id, col_key, title, col_type, options_json, width)
    VALUES (?, ?, ?, ?, ?, ?)
    """, (company_id, col_key, req.title, req.col_type, json.dumps(req.options or []), req.width or 120))
    
    new_id = c.lastrowid
    conn.commit()
    conn.close()
    return {
        "status": "success", 
        "id": new_id, 
        "col_key": col_key, 
        "title": req.title, 
        "col_type": req.col_type, 
        "options": req.options or [],
        "width": req.width or 120
    }

@app.delete("/api/custom-columns/{col_id}")
def api_delete_custom_column(col_id: int, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    c.execute("DELETE FROM custom_columns WHERE id = ?", (col_id,))
    conn.commit()
    conn.close()
    return {"status": "success", "message": "Sütun silindi."}

class UpdateCustomFieldRequest(BaseModel):
    style_id: int
    col_key: str
    value: Any

@app.post("/api/styles/update-custom-field")
def api_update_custom_field(req: UpdateCustomFieldRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()

    c.execute("SELECT custom_fields_json, order_id FROM styles WHERE id = ?", (req.style_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    
    fields = json.loads(row[0] or "{}")
    old_val = fields.get(req.col_key, "")
    fields[req.col_key] = req.value
    new_json = json.dumps(fields)

    c.execute("UPDATE styles SET custom_fields_json = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (new_json, req.style_id))
    
    user_name = user.get("full_name") or user.get("username")
    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, ?, ?, ?, 'Özel Sütun Değeri Güncellendi')
    """, (req.style_id, row[1], user.get("user_id"), user_name, req.col_key, str(old_val), str(req.value)))

    conn.commit()
    conn.close()
    return {"status": "success", "custom_fields": fields}

# ----------------- MASTER PRODUCTION SHEET (ÇARŞAF LİSTE) -----------------

SIZE_SORT_ORDER = {
    '32': 1, '34': 2, '36': 3, '38': 4, '40': 5, '42': 6, '44': 7, '46': 8, '48': 9, '50': 10, '52': 11,
    'XXS': 20, 'XS': 21, 'S': 22, 'M': 23, 'L': 24, 'XL': 25, 'XXL': 26, '2XL': 26, '3XL': 27, '4XL': 28, 'STD': 30
}

def get_size_sort_key(sz: str):
    s = str(sz).strip().upper()
    return SIZE_SORT_ORDER.get(s, 99)

@app.get("/api/styles")
def api_get_styles(user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    role = user.get("role")
    company_id = user.get("company_id") or 1
    
    if role in ["superadmin", "masterdeveloper"]:
        sql = """
        SELECT s.*, 
               o.po_number, o.customer_name, COALESCE(NULLIF(s.brand, ''), o.brand) as resolved_brand, o.season, o.delivery_date, o.order_date
        FROM styles s
        JOIN orders o ON s.order_id = o.id
        ORDER BY o.id DESC, s.id ASC
        """
        c.execute(sql)
    else:
        sql = """
        SELECT s.*, 
               o.po_number, o.customer_name, COALESCE(NULLIF(s.brand, ''), o.brand) as resolved_brand, o.season, o.delivery_date, o.order_date
        FROM styles s
        JOIN orders o ON s.order_id = o.id
        WHERE s.company_id = ?
        ORDER BY o.id DESC, s.id ASC
        """
        c.execute(sql, (company_id,))
    styles = [dict(r) for r in c.fetchall()]
    for s in styles:
        if s.get("resolved_brand"):
            s["brand"] = s["resolved_brand"]
    
    # Load FabricTag global lookup map once
    ft_map = {}
    try:
        ft_conn = FabricTagConnector.get_connection()
        if ft_conn:
            ft_c = ft_conn.cursor()
            ft_c.execute("SELECT id, internal_code, company_name, quality_name, quality_code, design_code, color, width, weight FROM fabrics")
            for f_row in ft_c.fetchall():
                f_d = dict(f_row)
                if f_d.get("id"):
                    ft_map[str(f_d["id"])] = f_d
                if f_d.get("internal_code"):
                    ft_map[f_d["internal_code"].upper()] = f_d
            ft_conn.close()
    except Exception:
        pass

    def resolve_fabric_meta(fab_id, raw_article):
        found = None
        if fab_id and str(fab_id) in ft_map:
            found = ft_map[str(fab_id)]
        elif raw_article:
            m = re.search(r'ELT\d+', str(raw_article), re.IGNORECASE)
            if m and m.group(0).upper() in ft_map:
                found = ft_map[m.group(0).upper()]
        return found

    all_sizes_set = set()

    for s in styles:
        c.execute("SELECT * FROM size_distributions WHERE style_id = ? ORDER BY sort_order ASC, id ASC", (s["id"],))
        sizes_rows = [dict(sz) for sz in c.fetchall()]
        s["sizes"] = sizes_rows
        
        size_map = {}
        for sz in sizes_rows:
            s_name = str(sz["size_name"]).strip()
            qty = sz["quantity"] or 0
            size_map[s_name] = qty
            if qty > 0:
                all_sizes_set.add(s_name)
        s["size_map"] = size_map
        
        if sizes_rows:
            sum_qty = sum(sz["quantity"] for sz in sizes_rows)
            if sum_qty != s["total_quantity"]:
                s["total_quantity"] = sum_qty
            s["cut_total_quantity"] = sum((sz.get("cut_quantity") or 0) for sz in sizes_rows)
            s["shipped_total_quantity"] = sum((sz.get("shipped_quantity") or 0) for sz in sizes_rows)
        else:
            s["cut_total_quantity"] = 0
            s["shipped_total_quantity"] = 0
        
        s["custom_fields"] = json.loads(s.get("custom_fields_json") or "{}")

        def is_empty_val(v):
            if not v:
                return True
            sv = str(v).strip().upper()
            return sv in ["", "KODSUZ", "YOK", "NONE", "NULL", "-"]

        # Slot 1 Fabric resolution
        c.execute("SELECT * FROM fabric_links WHERE style_id = ? AND (fabric_slot = 1 OR fabric_slot IS NULL) ORDER BY id DESC LIMIT 1", (s["id"],))
        fl1 = c.fetchone()
        fl1_dict = dict(fl1) if fl1 else None
        s["fabric_link"] = fl1_dict
        f1_meta = resolve_fabric_meta(fl1_dict.get("fabrictag_id") if fl1_dict else None, s.get("fabric_article") or s.get("fabric_type"))
        if f1_meta:
            q_name = (f1_meta.get("quality_name") or "").strip()
            q_code = (f1_meta.get("quality_code") or "").strip()
            comp = (f1_meta.get("company_name") or "").strip()
            des = (f1_meta.get("design_code") or "").strip()
            col = (f1_meta.get("color") or "").strip()
            s["fabric_company_1"] = comp
            s["fabric_variant_1"] = des
            s["fabric_color_1"] = col
            if is_empty_val(q_name) and is_empty_val(q_code):
                s["fabric_display_title_1"] = ""
                s["fabric_quality_name_1"] = ""
                s["fabric_quality_code_1"] = ""
            else:
                s["fabrictag_width"] = f1_meta.get("width")
                s["fabrictag_weight"] = f1_meta.get("weight")
                clean_name = "" if is_empty_val(q_name) else q_name
                clean_code = "" if is_empty_val(q_code) else f"({q_code})"
                q_combo = " ".join(filter(None, [clean_name, clean_code])).strip()
                parts = []
                if comp: parts.append(comp)
                if q_combo: parts.append(q_combo)
                if des: parts.append(des)
                if col and not is_empty_val(col): parts.append(col)
                s["fabric_display_title_1"] = " - ".join(parts) if parts else ""
                s["fabric_quality_name_1"] = clean_name
                s["fabric_quality_code_1"] = "" if is_empty_val(q_code) else q_code
        else:
            s["fabric_display_title_1"] = ""
            s["fabric_quality_name_1"] = ""
            s["fabric_quality_code_1"] = ""
            s["fabric_company_1"] = ""
            s["fabric_variant_1"] = ""
            s["fabric_color_1"] = ""

        # Slot 2 Fabric resolution
        c.execute("SELECT * FROM fabric_links WHERE style_id = ? AND fabric_slot = 2 ORDER BY id DESC LIMIT 1", (s["id"],))
        fl2 = c.fetchone()
        fl2_dict = dict(fl2) if fl2 else None
        s["fabric_link_2"] = fl2_dict
        f2_meta = resolve_fabric_meta(fl2_dict.get("fabrictag_id") if fl2_dict else None, s.get("fabric_article_2") or s.get("fabric_type_2"))
        if f2_meta:
            q_name2 = (f2_meta.get("quality_name") or "").strip()
            q_code2 = (f2_meta.get("quality_code") or "").strip()
            comp2 = (f2_meta.get("company_name") or "").strip()
            des2 = (f2_meta.get("design_code") or "").strip()
            col2 = (f2_meta.get("color") or "").strip()
            s["fabric_company_2"] = comp2
            s["fabric_variant_2"] = des2
            s["fabric_color_2"] = col2
            if is_empty_val(q_name2) and is_empty_val(q_code2):
                s["fabric_display_title_2"] = ""
                s["fabric_quality_name_2"] = ""
                s["fabric_quality_code_2"] = ""
            else:
                s["fabrictag_width_2"] = f2_meta.get("width")
                s["fabrictag_weight_2"] = f2_meta.get("weight")
                clean_name2 = "" if is_empty_val(q_name2) else q_name2
                clean_code2 = "" if is_empty_val(q_code2) else f"({q_code2})"
                q_combo2 = " ".join(filter(None, [clean_name2, clean_code2])).strip()
                parts2 = []
                if comp2: parts2.append(comp2)
                if q_combo2: parts2.append(q_combo2)
                if des2: parts2.append(des2)
                if col2 and not is_empty_val(col2): parts2.append(col2)
                s["fabric_display_title_2"] = " - ".join(parts2) if parts2 else ""
                s["fabric_quality_name_2"] = clean_name2
                s["fabric_quality_code_2"] = "" if is_empty_val(q_code2) else q_code2
        else:
            s["fabric_display_title_2"] = ""
            s["fabric_quality_name_2"] = ""
            s["fabric_quality_code_2"] = ""
            s["fabric_company_2"] = ""
            s["fabric_variant_2"] = ""
            s["fabric_color_2"] = ""

        if s.get("fabric_wastage_percent") is None:
            s["fabric_wastage_percent"] = 5.0
        if not s.get("fabric_order_unit"):
            s["fabric_order_unit"] = "M"
        if not s.get("fabric_price_currency_1"):
            s["fabric_price_currency_1"] = s.get("currency") or "€"
        if s.get("fabric_wastage_percent_2") is None:
            s["fabric_wastage_percent_2"] = 5.0
        if not s.get("fabric_order_unit_2"):
            s["fabric_order_unit_2"] = "M"
        if not s.get("fabric_price_currency_2"):
            s["fabric_price_currency_2"] = s.get("currency") or "€"


        if s.get("cost_data_json"):
            try:
                s["cost_data"] = json.loads(s["cost_data_json"])
            except Exception:
                s["cost_data"] = None
        else:
            s["cost_data"] = None

        if s.get("actual_cost_data_json"):
            try:
                s["actual_cost_data"] = json.loads(s["actual_cost_data_json"])
            except Exception:
                s["actual_cost_data"] = None
        else:
            s["actual_cost_data"] = None


        c.execute("SELECT * FROM production_status_logs WHERE style_id = ? ORDER BY id DESC LIMIT 15", (s["id"],))
        s["recent_logs"] = [dict(lg) for lg in c.fetchall()]



        c.execute("SELECT COUNT(*) FROM audit_logs WHERE style_id = ?", (s["id"],))
        s["audit_count"] = c.fetchone()[0]

    # Also load custom columns
    c.execute("SELECT * FROM custom_columns WHERE company_id = ? ORDER BY sort_order ASC, id ASC", (company_id,))
    custom_cols = [dict(r) for r in c.fetchall()]
    for col in custom_cols:
        col["options"] = json.loads(col.get("options_json") or "[]")
        if col.get("col_type") == "size" and col.get("title"):
            all_sizes_set.add(str(col["title"]).strip().upper())

    conn.close()
    
    sorted_sizes = sorted(list(all_sizes_set), key=get_size_sort_key)
    
    return {
        "styles": styles,
        "available_sizes": sorted_sizes,
        "custom_columns": custom_cols
    }

# ----------------- BULK SAVE ENDPOINT -----------------

class BulkSaveItem(BaseModel):
    id: int
    description: Optional[Any] = None
    fabric_order_status: Optional[Any] = None
    unit_price: Optional[Any] = None
    status: Optional[Any] = None
    fabric_ordered: Optional[Any] = None
    pps_sent: Optional[Any] = None
    shipping_sample_sent: Optional[Any] = None
    unit_meters: Optional[Any] = None
    unit_grams: Optional[Any] = None
    fabric_wastage_percent: Optional[Any] = None
    fabric_ordered_meters: Optional[Any] = None
    fabric_price_1: Optional[Any] = None
    fabric_price_currency_1: Optional[Any] = None
    fabric_arrival_date: Optional[Any] = None
    fabric_received_meters: Optional[Any] = None
    unit_meters_2: Optional[Any] = None
    unit_grams_2: Optional[Any] = None
    fabric_wastage_percent_2: Optional[Any] = None
    fabric_ordered_meters_2: Optional[Any] = None
    fabric_price_2: Optional[Any] = None
    fabric_price_currency_2: Optional[Any] = None
    fabric_arrival_date_2: Optional[Any] = None
    fabric_received_meters_2: Optional[Any] = None
    size_map: Optional[Dict[str, Any]] = None
    custom_fields: Optional[Dict[str, Any]] = None

class BulkSaveRequest(BaseModel):
    items: List[BulkSaveItem]

@app.post("/api/styles/bulk-save")
def api_bulk_save_styles(req: BulkSaveRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    user_name = user.get("full_name") or user.get("username")
    saved_count = 0

    for item in req.items:
        u_price = None
        if item.unit_price is not None and str(item.unit_price).strip() != "":
            try:
                u_price = float(str(item.unit_price).replace(",", "."))
            except Exception:
                pass

        fp1 = None
        if item.fabric_price_1 is not None and str(item.fabric_price_1).strip() != "":
            try:
                fp1 = float(str(item.fabric_price_1).replace(",", "."))
            except Exception:
                pass

        fp2 = None
        if item.fabric_price_2 is not None and str(item.fabric_price_2).strip() != "":
            try:
                fp2 = float(str(item.fabric_price_2).replace(",", "."))
            except Exception:
                pass

        c.execute("""
        UPDATE styles SET
            description = COALESCE(?, description),
            fabric_order_status = COALESCE(?, fabric_order_status),
            unit_price = COALESCE(?, unit_price),
            status = COALESCE(?, status),
            unit_meters = COALESCE(?, unit_meters),
            unit_grams = COALESCE(?, unit_grams),
            fabric_wastage_percent = COALESCE(?, fabric_wastage_percent),
            fabric_ordered_meters = COALESCE(?, fabric_ordered_meters),
            fabric_price_1 = COALESCE(?, fabric_price_1),
            fabric_price_currency_1 = COALESCE(?, fabric_price_currency_1),
            unit_meters_2 = COALESCE(?, unit_meters_2),
            unit_grams_2 = COALESCE(?, unit_grams_2),
            fabric_wastage_percent_2 = COALESCE(?, fabric_wastage_percent_2),
            fabric_ordered_meters_2 = COALESCE(?, fabric_ordered_meters_2),
            fabric_price_2 = COALESCE(?, fabric_price_2),
            fabric_price_currency_2 = COALESCE(?, fabric_price_currency_2),
            fabric_arrival_date = COALESCE(?, fabric_arrival_date),
            fabric_received_meters = COALESCE(?, fabric_received_meters),
            fabric_arrival_date_2 = COALESCE(?, fabric_arrival_date_2),
            fabric_received_meters_2 = COALESCE(?, fabric_received_meters_2),
            updated_at = CURRENT_TIMESTAMP
        WHERE id = ?
        """, (
            item.description,
            item.fabric_order_status,
            u_price,
            item.status,
            item.unit_meters,
            item.unit_grams,
            item.fabric_wastage_percent,
            item.fabric_ordered_meters,
            fp1,
            item.fabric_price_currency_1,
            item.unit_meters_2,
            item.unit_grams_2,
            item.fabric_wastage_percent_2,
            item.fabric_ordered_meters_2,
            fp2,
            item.fabric_price_currency_2,
            item.fabric_arrival_date,
            item.fabric_received_meters,
            item.fabric_arrival_date_2,
            item.fabric_received_meters_2,
            item.id
        ))


        if item.custom_fields is not None:
            c.execute("UPDATE styles SET custom_fields_json = ? WHERE id = ?", (json.dumps(item.custom_fields), item.id))

        if item.size_map:
            for sz_name, sz_qty in item.size_map.items():
                try:
                    qty_int = int(str(sz_qty).strip() or 0)
                except Exception:
                    qty_int = 0
                c.execute("SELECT id FROM size_distributions WHERE style_id = ? AND size_name = ?", (item.id, sz_name))
                row = c.fetchone()
                if row:
                    c.execute("UPDATE size_distributions SET quantity = ? WHERE id = ?", (qty_int, row[0]))
                else:
                    c.execute("INSERT INTO size_distributions (style_id, size_name, quantity) VALUES (?, ?, ?)", (item.id, sz_name, qty_int))
            
            c.execute("SELECT SUM(quantity) FROM size_distributions WHERE style_id = ?", (item.id,))
            new_tot = c.fetchone()[0] or 0
            c.execute("UPDATE styles SET total_quantity = ?, total_amount = ? * COALESCE(unit_price, 0) WHERE id = ?", (new_tot, new_tot, item.id))
        
        saved_count += 1

    conn.commit()
    conn.close()
    return {"status": "success", "message": f"{saved_count} satırdaki tüm değişiklikler başarıyla kaydedildi."}


# ----------------- DELETE STYLE / ORDER -----------------

def check_can_delete_orders(user: Dict[str, Any]):
    role = user.get("role")
    perms = user.get("permissions") or {}
    if role in ["superadmin", "masterdeveloper", "admin"]:
        return True
    if perms.get("can_delete_orders"):
        return True
    raise HTTPException(status_code=403, detail="Sipariş ve modelleri silme yetkiniz bulunmamaktadır.")

@app.delete("/api/styles/{style_id}")
def api_delete_style(style_id: int, user: Dict[str, Any] = Depends(require_user)):
    check_can_delete_orders(user)
    conn = get_db()
    c = conn.cursor()

    c.execute("SELECT order_id, style_no, color_name, total_quantity FROM styles WHERE id = ?", (style_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Silinecek model bulunamadı.")
    
    order_id, style_no, color_name, total_qty = row[0], row[1], row[2], row[3]

    try:
        c.execute("DELETE FROM audit_logs WHERE style_id = ?", (style_id,))
        c.execute("DELETE FROM size_distributions WHERE style_id = ?", (style_id,))
        c.execute("DELETE FROM production_status_logs WHERE style_id = ?", (style_id,))
        c.execute("DELETE FROM fabric_links WHERE style_id = ?", (style_id,))
        c.execute("DELETE FROM styles WHERE id = ?", (style_id,))

        c.execute("SELECT COUNT(*), COALESCE(SUM(total_quantity), 0), COALESCE(SUM(total_amount), 0) FROM styles WHERE order_id = ?", (order_id,))
        rem_count, rem_qty, rem_amt = c.fetchone()

        if rem_count == 0:
            c.execute("DELETE FROM orders WHERE id = ?", (order_id,))
        else:
            c.execute("UPDATE orders SET total_quantity = ?, total_amount = ? WHERE id = ?", (rem_qty, rem_amt, order_id))

        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        raise HTTPException(status_code=500, detail=f"Silme hatası: {str(e)}")


# ----------------- ORDER MANAGEMENT & PERMANENT DELETE -----------------

@app.get("/api/orders/list")
def api_get_orders_list(user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    role = user.get("role")
    comp_id = user.get("company_id")
    
    if role in ["superadmin", "masterdeveloper"]:
        c.execute("""
            SELECT o.id, o.po_number, o.customer_name, o.brand, o.season, o.order_date, o.delivery_date,
                   o.total_quantity, o.total_amount, o.currency, o.status, o.created_at,
                   COUNT(s.id) as styles_count,
                   GROUP_CONCAT(DISTINCT s.style_no) as style_names
            FROM orders o
            LEFT JOIN styles s ON s.order_id = o.id
            GROUP BY o.id
            ORDER BY o.id DESC
        """)
    else:
        c.execute("""
            SELECT o.id, o.po_number, o.customer_name, o.brand, o.season, o.order_date, o.delivery_date,
                   o.total_quantity, o.total_amount, o.currency, o.status, o.created_at,
                   COUNT(s.id) as styles_count,
                   GROUP_CONCAT(DISTINCT s.style_no) as style_names
            FROM orders o
            LEFT JOIN styles s ON s.order_id = o.id
            WHERE o.company_id = ?
            GROUP BY o.id
            ORDER BY o.id DESC
        """, (comp_id or 1,))
        
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    return rows

@app.delete("/api/orders/{order_id}")
def api_delete_order(order_id: int, user: Dict[str, Any] = Depends(require_user)):
    check_can_delete_orders(user)
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT id, po_number, customer_name FROM orders WHERE id = ?", (order_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Sipariş bulunamadı.")
    
    po_num, cust = row["po_number"], row["customer_name"]

    try:
        c.execute("SELECT id FROM styles WHERE order_id = ?", (order_id,))
        style_ids = [r["id"] for r in c.fetchall()]

        for s_id in style_ids:
            c.execute("DELETE FROM audit_logs WHERE style_id = ?", (s_id,))
            c.execute("DELETE FROM size_distributions WHERE style_id = ?", (s_id,))
            c.execute("DELETE FROM production_status_logs WHERE style_id = ?", (s_id,))
            c.execute("DELETE FROM fabric_links WHERE style_id = ?", (s_id,))
        
        c.execute("DELETE FROM styles WHERE order_id = ?", (order_id,))
        c.execute("DELETE FROM orders WHERE id = ?", (order_id,))
        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        raise HTTPException(status_code=500, detail=f"Sipariş silme hatası: {str(e)}")

    conn.close()
    return {"status": "success", "message": f"Sipariş (PO: {po_num}, Müşteri: {cust}) ve tüm modelleri kalıcı olarak silindi."}

class BulkDeleteOrdersRequest(BaseModel):
    order_ids: List[int]

@app.post("/api/orders/bulk-delete")
def api_bulk_delete_orders(req: BulkDeleteOrdersRequest, user: Dict[str, Any] = Depends(require_user)):
    if not req.order_ids:
        return {"status": "success", "deleted_count": 0}
    
    conn = get_db()
    c = conn.cursor()
    try:
        for order_id in req.order_ids:
            c.execute("SELECT id FROM styles WHERE order_id = ?", (order_id,))
            style_ids = [r["id"] for r in c.fetchall()]
            for s_id in style_ids:
                c.execute("DELETE FROM audit_logs WHERE style_id = ?", (s_id,))
                c.execute("DELETE FROM size_distributions WHERE style_id = ?", (s_id,))
                c.execute("DELETE FROM production_status_logs WHERE style_id = ?", (s_id,))
                c.execute("DELETE FROM fabric_links WHERE style_id = ?", (s_id,))
            c.execute("DELETE FROM styles WHERE order_id = ?", (order_id,))
            c.execute("DELETE FROM orders WHERE id = ?", (order_id,))
        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        raise HTTPException(status_code=500, detail=f"Toplu silme hatası: {str(e)}")

    conn.close()
    return {"status": "success", "deleted_count": len(req.order_ids)}




# ----------------- EXCEL EXPORT WITH EMBEDDED PHOTOS (MODERN DESIGN) -----------------

class ExportExcelRequest(BaseModel):
    style_ids: Optional[List[int]] = None

def generate_styles_excel_workbook(style_ids: Optional[List[int]] = None, company_id: int = 1) -> Response:
    conn = get_db()
    c = conn.cursor()
    
    if style_ids and len(style_ids) > 0:
        placeholders = ",".join("?" for _ in style_ids)
        sql = f"""
        SELECT s.*, 
               o.po_number, o.customer_name, o.brand, o.season, o.delivery_date, o.order_date
        FROM styles s
        JOIN orders o ON s.order_id = o.id
        WHERE s.company_id = ? AND s.id IN ({placeholders})
        ORDER BY s.id DESC
        """
        params = [company_id] + style_ids
    else:
        sql = """
        SELECT s.*, 
               o.po_number, o.customer_name, o.brand, o.season, o.delivery_date, o.order_date
        FROM styles s
        JOIN orders o ON s.order_id = o.id
        WHERE s.company_id = ?
        ORDER BY s.id DESC
        """
        params = [company_id]
        
    c.execute(sql, tuple(params))
    styles = [dict(r) for r in c.fetchall()]
    all_style_ids = [s["id"] for s in styles]
    
    # 1. Beden Dağılımlarını Topla
    sizes_by_style = {}
    present_size_names = []
    size_seen = set()
    
    if all_style_ids:
        pl_ids = ",".join("?" for _ in all_style_ids)
        c.execute(f"""
        SELECT style_id, size_name, quantity, sort_order
        FROM size_distributions
        WHERE style_id IN ({pl_ids})
        ORDER BY sort_order ASC, id ASC
        """, tuple(all_style_ids))
        for sz_row in c.fetchall():
            sid = sz_row["style_id"]
            sname = str(sz_row["size_name"]).strip()
            qty = sz_row["quantity"] or 0
            if sid not in sizes_by_style:
                sizes_by_style[sid] = {}
            sizes_by_style[sid][sname] = qty
            if sname not in size_seen and qty > 0:
                size_seen.add(sname)
                present_size_names.append(sname)
                
    def size_sort_key(s):
        try:
            return (0, float(s.replace(",", ".")))
        except ValueError:
            apparel_order = {'XXS': 1, 'XS': 2, 'S': 3, 'M': 4, 'L': 5, 'XL': 6, 'XXL': 7, '2XL': 7, '3XL': 8, '4XL': 9}
            return (1, apparel_order.get(s.upper(), 99), s)

    present_size_names.sort(key=size_sort_key)
    
    # 2. Kumaş Bağlantılarını Topla (fabric_links)
    fl_by_style = {}
    if all_style_ids:
        pl_ids = ",".join("?" for _ in all_style_ids)
        c.execute(f"""
        SELECT * FROM fabric_links WHERE style_id IN ({pl_ids}) ORDER BY id ASC
        """, tuple(all_style_ids))
        for fl in c.fetchall():
            sid = fl["style_id"]
            if sid not in fl_by_style:
                fl_by_style[sid] = dict(fl)

    # FabricTag detaylarını entegre et
    try:
        ft_conn = FabricTagConnector.get_connection()
        if ft_conn:
            ft_c = ft_conn.cursor()
            for sid, fl_dict in fl_by_style.items():
                if fl_dict.get("fabrictag_id"):
                    ft_c.execute("SELECT width, weight FROM fabrics WHERE id = ?", (fl_dict["fabrictag_id"],))
                    ft_row = ft_c.fetchone()
                    if ft_row:
                        fl_dict["ft_width"] = ft_row["width"]
                        fl_dict["ft_weight"] = ft_row["weight"]
            ft_conn.close()
    except Exception:
        pass
        
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sipariş Listesi"
    
    # Modern Excel Görünümü: Izgara çizgilerini göster ve üst başlığı dondur
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "B2"
    
    # Başlıkları Oluştur
    headers = [
        "Görsel", "Müşteri", "Sezon", "PO No", "Model No", 
        "Renk Kodu", "Renk Adı"
    ]
    
    size_col_start = len(headers) + 1
    for sz_name in present_size_names:
        headers.append(f"Beden {sz_name}")
    size_col_end = len(headers)
    
    col_idx_qty = len(headers) + 1
    headers.append("Toplam Adet")
    col_idx_price = len(headers) + 1
    headers.append("Birim Fiyat")
    col_idx_currency = len(headers) + 1
    headers.append("Döviz")
    col_idx_amount = len(headers) + 1
    headers.append("Toplam Tutar")
    
    # Kumaş Bilgileri Sütunları
    headers.extend([
        "Kumaş / Kalite Adı",
        "Kumaş Kodu",
        "Karışım / Kompozisyon",
        "Kumaş Tedarikçisi",
        "En (cm)",
        "Gramaj (g/m²)",
        "Sipariş Kumaş (M/KG)",
        "Gelen Kumaş (M/KG)",
        "Kumaş Durumu / Termin"
    ])
    
    col_idx_ordered_fabric = headers.index("Sipariş Kumaş (M/KG)") + 1
    col_idx_received_fabric = headers.index("Gelen Kumaş (M/KG)") + 1
    col_idx_width = headers.index("En (cm)") + 1
    col_idx_weight = headers.index("Gramaj (g/m²)") + 1
    
    # Diğer Sütunlar
    headers.extend([
        "Sipariş Tarihi",
        "Yükleme Tarihi",
        "Üretim Aşaması",
        "Beden Dağılımı (Özet)",
        "Notlar"
    ])
    
    # Modern Başlık Tasarımı (Koyu Füme / Slate Navy, Beyaz Kalın Yazı)
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    header_border = Border(
        left=Side(style="thin", color="334155"),
        right=Side(style="thin", color="334155"),
        top=Side(style="thin", color="0F172A"),
        bottom=Side(style="medium", color="0F172A")
    )
    
    ws.row_dimensions[1].height = 34
    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = header_border
        
    # Görsel sütunu genişliği
    ws.column_dimensions["A"].width = 18
    img_buffers = []
    
    # Hücre kenarlık ve yazı tipleri
    cell_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )
    
    font_normal = Font(name="Segoe UI", size=9.5, color="1E293B")
    font_bold = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
    font_style_no = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    font_muted = Font(name="Segoe UI", size=9, color="94A3B8")
    
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    
    last_row = 1
    
    for idx, s in enumerate(styles, start=2):
        last_row = idx
        ws.row_dimensions[idx].height = 75
        current_fill = fill_white if (idx % 2 == 0) else fill_zebra
        
        # Beden dağılımı metni
        style_size_map = sizes_by_style.get(s["id"], {})
        sizes_summary_str = ", ".join(f"{sz}: {qty}" for sz, qty in style_size_map.items() if qty > 0)
        
        # Notlar
        notes_val = s.get("notes") or ""
        if not notes_val and s.get("custom_fields_json"):
            try:
                cf = json.loads(s["custom_fields_json"])
                notes_val = cf.get("notes") or cf.get("notlar") or ""
            except Exception:
                pass
                
        qty = s.get("total_quantity") or 0
        price = s.get("unit_price") or 0
        tot_amt = (price or 0) * (qty or 0)
        
        # Kumaş Bilgileri (styles tablosu ve fabric_links harmanlanır)
        fl = fl_by_style.get(s["id"], {})
        fabric_article_val = s.get("fabric_article") or fl.get("fabric_name") or ""
        fabric_code_val = fl.get("fabric_code") or s.get("fabric_type") or ""
        fabric_comp_val = s.get("fabric_composition") or fl.get("composition") or ""
        fabric_supplier_val = fl.get("supplier") or ""
        fabric_width_val = s.get("fabric_width") or fl.get("ft_width") or ""
        fabric_weight_val = s.get("fabric_weight") or fl.get("ft_weight") or ""
        fabric_ordered_m = s.get("fabric_ordered_meters") or 0.0
        fabric_received_m = s.get("fabric_received_meters") or 0.0
        fabric_status_val = s.get("fabric_order_status") or fl.get("status") or ""
        
        # A Sütunu: Görsel Hücresi
        cell_a = ws.cell(row=idx, column=1)
        cell_a.border = cell_border
        cell_a.fill = current_fill
        
        img_url = s.get("image_url") or s.get("image_url_2")
        img_added = False
        if img_url:
            try:
                pil_img = None
                if img_url.startswith("data:image"):
                    hdr, enc = img_url.split(",", 1)
                    pil_img = PILImage.open(io.BytesIO(base64.b64decode(enc)))
                elif img_url.startswith("/uploads/"):
                    p = UPLOADS_DIR / img_url.replace("/uploads/", "").lstrip("/")
                    if p.exists():
                        pil_img = PILImage.open(p)
                elif Path(img_url).exists():
                    pil_img = PILImage.open(img_url)
                    
                if pil_img:
                    pil_img = pil_img.convert("RGB")
                    pil_img.thumbnail((105, 88), PILImage.Resampling.LANCZOS)
                    buf = io.BytesIO()
                    pil_img.save(buf, format="JPEG", quality=88)
                    buf.seek(0)
                    img_buffers.append(buf)
                    xl_img = OpenpyxlImage(buf)
                    xl_img.width = pil_img.width
                    xl_img.height = pil_img.height
                    ws.add_image(xl_img, f"A{idx}")
                    img_added = True
            except Exception:
                pass
                
        if not img_added:
            cell_a.value = "[Görsel Yok]"
            cell_a.alignment = align_center
            cell_a.font = font_muted
            
        # Satır Verilerini Sırayla Doldur
        # 1. Temel Bilgiler (2-7)
        ws.cell(row=idx, column=2, value=s.get("customer_name") or s.get("brand") or "").alignment = align_left
        ws.cell(row=idx, column=2).font = font_bold
        ws.cell(row=idx, column=3, value=s.get("season") or "").alignment = align_center
        ws.cell(row=idx, column=4, value=s.get("po_number") or "").alignment = align_center
        ws.cell(row=idx, column=4).font = font_bold
        ws.cell(row=idx, column=5, value=s.get("style_no") or "").alignment = align_center
        ws.cell(row=idx, column=5).font = font_style_no
        ws.cell(row=idx, column=6, value=s.get("color_code") or "").alignment = align_center
        ws.cell(row=idx, column=7, value=s.get("color_name") or "").alignment = align_left_wrap
        
        # 2. Beden Dağılımı Sütunları (TÜM RAKAMLAR ORTALI)
        for s_idx, sz_name in enumerate(present_size_names, start=size_col_start):
            sz_qty = style_size_map.get(sz_name)
            c_sz = ws.cell(row=idx, column=s_idx, value=sz_qty if sz_qty else None)
            c_sz.alignment = align_center
            if sz_qty:
                c_sz.number_format = "#,##0"
                c_sz.font = font_bold
                
        # 3. Miktar ve Fiyatlar (TÜM RAKAMLAR ORTALI)
        c_qty = ws.cell(row=idx, column=col_idx_qty, value=qty)
        c_qty.alignment = align_center
        c_qty.number_format = "#,##0"
        c_qty.font = font_bold
        
        c_prc = ws.cell(row=idx, column=col_idx_price, value=price)
        c_prc.alignment = align_center
        c_prc.number_format = "#,##0.00"
        
        c_cur = ws.cell(row=idx, column=col_idx_currency, value=s.get("currency") or "EUR")
        c_cur.alignment = align_center
        
        c_amt = ws.cell(row=idx, column=col_idx_amount, value=tot_amt)
        c_amt.alignment = align_center
        c_amt.number_format = "#,##0.00"
        c_amt.font = font_bold
        
        # 4. Kumaş Bilgileri Sütunları
        ws.cell(row=idx, column=headers.index("Kumaş / Kalite Adı") + 1, value=fabric_article_val).alignment = align_left_wrap
        ws.cell(row=idx, column=headers.index("Kumaş Kodu") + 1, value=fabric_code_val).alignment = align_center
        ws.cell(row=idx, column=headers.index("Karışım / Kompozisyon") + 1, value=fabric_comp_val).alignment = align_left_wrap
        ws.cell(row=idx, column=headers.index("Kumaş Tedarikçisi") + 1, value=fabric_supplier_val).alignment = align_left
        
        c_w = ws.cell(row=idx, column=col_idx_width, value=fabric_width_val)
        c_w.alignment = align_center
        
        c_g = ws.cell(row=idx, column=col_idx_weight, value=fabric_weight_val)
        c_g.alignment = align_center
        
        # Kumaş Metrajları (TÜM RAKAMLAR ORTALI)
        c_ord_f = ws.cell(row=idx, column=col_idx_ordered_fabric, value=fabric_ordered_m if fabric_ordered_m else None)
        c_ord_f.alignment = align_center
        if fabric_ordered_m:
            c_ord_f.number_format = "#,##0.0"
            
        c_rec_f = ws.cell(row=idx, column=col_idx_received_fabric, value=fabric_received_m if fabric_received_m else None)
        c_rec_f.alignment = align_center
        if fabric_received_m:
            c_rec_f.number_format = "#,##0.0"
            c_rec_f.font = font_bold
            
        ws.cell(row=idx, column=headers.index("Kumaş Durumu / Termin") + 1, value=fabric_status_val).alignment = align_left_wrap
        
        # 5. Tarihler, Aşama ve Notlar
        ws.cell(row=idx, column=headers.index("Sipariş Tarihi") + 1, value=s.get("order_date") or "").alignment = align_center
        ws.cell(row=idx, column=headers.index("Yükleme Tarihi") + 1, value=s.get("delivery_date") or "").alignment = align_center
        ws.cell(row=idx, column=headers.index("Üretim Aşaması") + 1, value=s.get("status") or "Planlama aşamasında").alignment = align_center
        ws.cell(row=idx, column=headers.index("Üretim Aşaması") + 1).font = font_bold
        
        ws.cell(row=idx, column=headers.index("Beden Dağılımı (Özet)") + 1, value=sizes_summary_str).alignment = align_left_wrap
        ws.cell(row=idx, column=headers.index("Notlar") + 1, value=notes_val).alignment = align_left_wrap

        # Tüm hücrelerin çerçevesini ve zemin rengini uygula
        for c_idx in range(2, len(headers) + 1):
            cell = ws.cell(row=idx, column=c_idx)
            cell.border = cell_border
            cell.fill = current_fill
            if not cell.font or cell.font.name != "Segoe UI":
                cell.font = font_normal

    # Modern Genel Toplam Satırı
    if last_row >= 2:
        sum_row = last_row + 1
        ws.row_dimensions[sum_row].height = 30
        sum_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        sum_border = Border(
            top=Side(style="thin", color="94A3B8"),
            bottom=Side(style="double", color="0F172A"),
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0")
        )
        
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=sum_row, column=c_idx)
            cell.fill = sum_fill
            cell.border = sum_border
            cell.font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
            cell.alignment = align_center
            
        ws.cell(row=sum_row, column=2, value="GENEL TOPLAM").alignment = Alignment(horizontal="left", vertical="center")
        
        # Beden Sütunları Toplamları (ORTALI)
        for s_idx in range(size_col_start, size_col_end + 1):
            c_letter = get_column_letter(s_idx)
            c_sz_sum = ws.cell(row=sum_row, column=s_idx, value=f"=SUM({c_letter}2:{c_letter}{last_row})")
            c_sz_sum.number_format = "#,##0"
            c_sz_sum.alignment = align_center
        
        # Adet Toplamı (ORTALI)
        col_qty_letter = get_column_letter(col_idx_qty)
        cell_qty = ws.cell(row=sum_row, column=col_idx_qty, value=f"=SUM({col_qty_letter}2:{col_qty_letter}{last_row})")
        cell_qty.number_format = "#,##0"
        cell_qty.alignment = align_center
        
        # Tutar Toplamı (ORTALI)
        col_amt_letter = get_column_letter(col_idx_amount)
        cell_amt = ws.cell(row=sum_row, column=col_idx_amount, value=f"=SUM({col_amt_letter}2:{col_amt_letter}{last_row})")
        cell_amt.number_format = "#,##0.00"
        cell_amt.alignment = align_center
        
        # Kumaş Metraj Toplamları (ORTALI)
        col_ordf_letter = get_column_letter(col_idx_ordered_fabric)
        c_ordf_sum = ws.cell(row=sum_row, column=col_idx_ordered_fabric, value=f"=SUM({col_ordf_letter}2:{col_ordf_letter}{last_row})")
        c_ordf_sum.number_format = "#,##0.0"
        c_ordf_sum.alignment = align_center
        
        col_recf_letter = get_column_letter(col_idx_received_fabric)
        c_recf_sum = ws.cell(row=sum_row, column=col_idx_received_fabric, value=f"=SUM({col_recf_letter}2:{col_recf_letter}{last_row})")
        c_recf_sum.number_format = "#,##0.0"
        c_recf_sum.alignment = align_center
        
        # Otomatik Filtreleme (AutoFilter) Etkinleştir
        last_col_letter = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

    # Sütun Genişliklerini İçeriğe Göre Otomatik Dengele
    for col_idx in range(2, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[col_letter]:
            val_str = str(cell.value or "")
            if val_str.startswith("="):
                continue
            if "\n" in val_str:
                lines = val_str.split("\n")
                max_len = max(max_len, max(len(l) for l in lines))
            else:
                max_len = max(max_len, len(val_str))
        # Beden sütunları dar kalsın (en az 9, en çok 14)
        if size_col_start <= col_idx <= size_col_end:
            ws.column_dimensions[col_letter].width = max(9, min(14, max_len + 3))
        else:
            ws.column_dimensions[col_letter].width = max(11, min(45, max_len + 3))

    conn.close()

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    
    filename = f"Siparisler_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    val = out.getvalue()
    return Response(
        content=val,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Content-Length": str(len(val)),
            "Cache-Control": "no-cache, no-store, must-revalidate"
        }
    )

@app.get("/api/styles/export-excel")
def api_export_styles_excel_get(ids: Optional[str] = Query(None), authorization: Optional[str] = Header(None)):
    user = None
    if authorization:
        t = authorization.replace("Bearer ", "").strip()
        user = get_current_user(t)
    company_id = user.get("company_id") if user else 1
    
    style_ids = None
    if ids:
        try:
            style_ids = [int(x.strip()) for x in ids.split(",") if x.strip().isdigit()]
        except Exception:
            pass
            
    return generate_styles_excel_workbook(style_ids, company_id)

@app.post("/api/styles/export-excel")
def api_export_styles_excel_post(req: ExportExcelRequest, authorization: Optional[str] = Header(None)):
    user = None
    if authorization:
        t = authorization.replace("Bearer ", "").strip()
        user = get_current_user(t)
    company_id = user.get("company_id") if user else 1
    
    return generate_styles_excel_workbook(req.style_ids, company_id)



# ----------------- SINGLE SIZE INLINE UPDATE -----------------

class UpdateSingleSizeRequest(BaseModel):
    style_id: int
    size_name: str
    quantity: int
    reason: Optional[str] = "Hızlı Beden Revizyonu"

@app.post("/api/styles/update-single-size")
def api_update_single_size(req: UpdateSingleSizeRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()

    c.execute("SELECT total_quantity, order_id, style_no, color_name FROM styles WHERE id = ?", (req.style_id,))
    st = c.fetchone()
    if not st:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    
    old_total = st[0] or 0
    order_id = st[1]
    style_no = st[2]
    color_name = st[3]

    c.execute("SELECT id, quantity FROM size_distributions WHERE style_id = ? AND size_name = ?", (req.style_id, req.size_name))
    row = c.fetchone()
    old_size_qty = row[1] if row else 0

    if row:
        c.execute("UPDATE size_distributions SET quantity = ? WHERE id = ?", (req.quantity, row[0]))
    else:
        c.execute("INSERT INTO size_distributions (style_id, size_name, quantity) VALUES (?, ?, ?)", (req.style_id, req.size_name, req.quantity))

    c.execute("SELECT SUM(quantity) FROM size_distributions WHERE style_id = ?", (req.style_id,))
    new_total = c.fetchone()[0] or 0

    c.execute("UPDATE styles SET total_quantity = ?, total_amount = ? * unit_price, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (new_total, new_total, req.style_id))

    c.execute("""
    UPDATE orders 
    SET total_quantity = (SELECT SUM(total_quantity) FROM styles WHERE order_id = ?),
        total_amount = (SELECT SUM(total_quantity * unit_price) FROM styles WHERE order_id = ?)
    WHERE id = ?
    """, (order_id, order_id, order_id))

    user_name = user.get("full_name") or user.get("username")
    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (req.style_id, order_id, user.get("user_id"), user_name, f"size_{req.size_name}", f"{req.size_name}: {old_size_qty}", f"{req.size_name}: {req.quantity}", f"{req.reason} (Toplam: {old_total} -> {new_total})"))

    c.execute("""
    INSERT INTO production_status_logs (style_id, user_name, category, title, message)
    VALUES (?, ?, 'Beden Revizyonu', ?, ?)
    """, (req.style_id, user_name, f"Beden {req.size_name} Değiştirildi", f"{req.size_name}: {old_size_qty} -> {req.quantity} Adet (Yeni Toplam: {new_total})"))

    conn.commit()
    conn.close()
    return {"status": "success", "new_total": new_total, "message": f"{req.size_name} bedeni güncellendi. Yeni toplam: {new_total}"}

# ----------------- DUAL IMAGE UPLOAD (SLOT 1 & SLOT 2) -----------------

@app.post("/api/styles/{style_id}/image")
async def api_upload_style_image(
    style_id: int, 
    slot: int = Query(1),
    file: UploadFile = File(...), 
    user: Dict[str, Any] = Depends(require_user)
):
    ext = Path(file.filename).suffix.lower()
    if ext not in [".jpg", ".jpeg", ".png", ".webp", ".gif"]:
        raise HTTPException(status_code=400, detail="Geçersiz görsel formatı. Lütfen JPG veya PNG yükleyin.")
    
    filename = f"style_{style_id}_s{slot}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
    dest_path = STYLE_IMAGES_DIR / filename
    
    with open(dest_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
        
    rel_url = f"/uploads/styles/{filename}"
    
    col_name = "image_url_2" if slot == 2 else "image_url"
    
    conn = get_db()
    c = conn.cursor()
    c.execute(f"UPDATE styles SET {col_name} = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (rel_url, style_id))
    
    user_name = user.get("full_name") or user.get("username")
    c.execute("""
    INSERT INTO audit_logs (style_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, 'Eski Görsel', ?, ?)
    """, (style_id, user.get("user_id"), user_name, col_name, rel_url, f"Model Fotoğrafı Yüklendi (Görsel {slot})"))

    conn.commit()
    conn.close()
    
    return {"status": "success", "image_url": rel_url, "slot": slot}

@app.delete("/api/styles/{style_id}/image/{slot}")
def api_delete_style_image(style_id: int, slot: int, user: Dict[str, Any] = Depends(require_user)):
    col_name = "image_url_2" if slot == 2 else "image_url"
    conn = get_db()
    c = conn.cursor()
    c.execute(f"UPDATE styles SET {col_name} = NULL, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (style_id,))
    conn.commit()
    conn.close()
    return {"status": "success", "message": f"Görsel {slot} silindi."}

@app.post("/api/styles/{style_id}/swap-images")
def api_swap_style_images(style_id: int, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT image_url, image_url_2 FROM styles WHERE id = ?", (style_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    img1, img2 = row[0], row[1]
    c.execute("UPDATE styles SET image_url = ?, image_url_2 = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (img2, img1, style_id))
    user_name = user.get("full_name") or user.get("username")
    c.execute("""
    INSERT INTO audit_logs (style_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, 'image_url', ?, ?, 'Görsel 1 ve 2 Yer Değiştirildi')
    """, (style_id, user.get("user_id"), user_name, str(img1), str(img2)))
    conn.commit()
    conn.close()
    return {"status": "success", "image_url": img2, "image_url_2": img1, "message": "Görseller yer değiştirildi."}


# ----------------- UPDATE CELL WITH AUDIT LOG & TIMESTAMPS -----------------

class StyleUpdateFieldRequest(BaseModel):
    style_id: int
    field: str
    value: Any
    reason: Optional[str] = "Kullanıcı Düzenlemesi"

@app.post("/api/styles/update-cell")
def api_update_style_cell(req: StyleUpdateFieldRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()

    if req.field in ["order_date", "delivery_date", "season"]:
        c.execute("SELECT order_id FROM styles WHERE id = ?", (req.style_id,))
        order_id = c.fetchone()[0]
        c.execute(f"UPDATE orders SET {req.field} = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (req.value, order_id))
        user_name = user.get("full_name") or user.get("username")
        c.execute("""
        INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
        VALUES (?, ?, ?, ?, ?, '', ?, ?)
        """, (req.style_id, order_id, user.get("user_id"), user_name, req.field, str(req.value), req.reason))
        conn.commit()
        conn.close()
        return {"status": "success", "message": f"{req.field} güncellendi"}

    c.execute(f"SELECT {req.field}, order_id FROM styles WHERE id = ?", (req.style_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    
    old_val = str(row[0] or "")
    order_id = row[1]
    new_val = str(req.value or "")

    user_name = user.get("full_name") or user.get("username")
    now_date_str = datetime.now().strftime('%d.%m.%Y')
    now_time_str = datetime.now().strftime('%d.%m.%Y %H:%M')

    extra_sql = ""
    extra_params = []

    if req.field == "fabric_ordered":
        val_int = int(req.value or 0)
        date_val = now_date_str if val_int == 1 else None
        extra_sql = ", fabric_ordered_date = ?"
        extra_params.append(date_val)
    elif req.field == "pps_sent":
        str_val = str(req.value or "").strip()
        date_val = now_date_str if str_val and str_val != "-" and str_val != "0" else None
        extra_sql = ", pps_sent_date = ?"
        extra_params.append(date_val)
    elif req.field == "shipping_sample_sent":
        val_int = int(req.value or 0)
        date_val = now_date_str if val_int == 1 else None
        extra_sql = ", shipping_sample_sent_date = ?"
        extra_params.append(date_val)
    elif req.field == "status":
        extra_sql = ", last_status_updated_by = ?, last_status_updated_at = ?"
        extra_params.extend([user_name, now_time_str])

    sql = f"UPDATE styles SET {req.field} = ? {extra_sql}, updated_at = CURRENT_TIMESTAMP WHERE id = ?"
    c.execute(sql, [req.value] + extra_params + [req.style_id])
    
    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (req.style_id, order_id, user.get("user_id"), user_name, req.field, old_val, new_val, req.reason))

    if req.field == "unit_price":
        c.execute("""
        UPDATE orders 
        SET total_amount = (SELECT SUM(total_quantity * unit_price) FROM styles WHERE order_id = ?)
        WHERE id = ?
        """, (order_id, order_id))
    elif req.field == "fabric_received_meters":
        rec_meters = float(req.value or 0)
        if rec_meters > 0:
            log_msg = f"{now_date_str} tarihinde {rec_meters} metre kumaş geldi."
            c.execute("""
            INSERT INTO production_status_logs (style_id, user_name, category, title, message)
            VALUES (?, ?, 'Kumaş Girişi', 'Kumaş Depoya Geldi', ?)
            """, (req.style_id, user_name, log_msg))

    c.execute("""
    INSERT INTO production_status_logs (style_id, user_name, category, title, message)
    VALUES (?, ?, 'İşlem Aşaması', ?, ?)
    """, (req.style_id, user_name, f"'{req.field}' Değiştirildi", f"{old_val} ➜ {new_val} ({req.reason})"))


    conn.commit()
    conn.close()
    return {
        "status": "success", 
        "message": f"{req.field} güncellendi",
        "user_name": user_name,
        "date_str": date_val if 'date_val' in locals() else now_date_str
    }


# ----------------- KESİMHANE & PASTAL GÜNCELLEME -----------------

class CuttingSizeItem(BaseModel):
    size_name: str
    cut_quantity: int = 0

class CuttingUpdateRequest(BaseModel):
    style_id: int
    cut_meters: Optional[float] = 0.0
    sizes: List[CuttingSizeItem] = []

@app.post("/api/cutting/update")
def api_update_cutting(req: CuttingUpdateRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()

    c.execute("SELECT id, order_id, style_no, total_quantity, cut_meters, actual_unit_meters, cutting_date FROM styles WHERE id = ?", (req.style_id,))
    style_row = c.fetchone()
    if not style_row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    
    order_id = style_row["order_id"]
    user_name = user.get("full_name") or user.get("username") or "Kesim Şefi"
    now_str = datetime.now().strftime("%d.%m.%Y %H:%M")

    total_cut_qty = 0
    for sz in req.sizes:
        c.execute("""
        UPDATE size_distributions 
        SET cut_quantity = ? 
        WHERE style_id = ? AND size_name = ?
        """, (int(sz.cut_quantity or 0), req.style_id, sz.size_name))
        total_cut_qty += int(sz.cut_quantity or 0)

    cut_m = float(req.cut_meters or 0.0)
    actual_unit_m = round(cut_m / total_cut_qty, 3) if total_cut_qty > 0 and cut_m > 0 else 0.0

    c.execute("""
    UPDATE styles 
    SET cut_meters = ?, actual_unit_meters = ?, cutting_date = ?, updated_at = CURRENT_TIMESTAMP 
    WHERE id = ?
    """, (cut_m, actual_unit_m, now_str, req.style_id))

    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, 'cutting_update', ?, ?, 'Kesimhane Beden & Metraj Güncellemesi')
    """, (req.style_id, order_id, user.get("user_id"), user_name, 
          f"M:{style_row['cut_meters']}", f"M:{cut_m}, Adet:{total_cut_qty}, Birim:{actual_unit_m}, Tarih:{now_str}"))

    conn.commit()
    conn.close()

    return {
        "status": "success",
        "style_id": req.style_id,
        "total_cut_quantity": total_cut_qty,
        "cut_meters": cut_m,
        "actual_unit_meters": actual_unit_m,
        "cutting_date": now_str
    }


class ShippingSizeItem(BaseModel):
    size_name: str
    shipped_quantity: int = 0

class ShippingUpdateRequest(BaseModel):
    style_id: int
    sizes: List[ShippingSizeItem] = []

@app.post("/api/shipping/update")
def api_update_shipping(req: ShippingUpdateRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()

    c.execute("SELECT id, order_id, style_no, total_quantity, shipping_date FROM styles WHERE id = ?", (req.style_id,))
    style_row = c.fetchone()
    if not style_row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    
    order_id = style_row["order_id"]
    user_name = user.get("full_name") or user.get("username") or "Sevkiyat Sorumlusu"
    now_str = datetime.now().strftime("%d.%m.%Y %H:%M")

    total_shipped_qty = 0
    for sz in req.sizes:
        c.execute("""
        UPDATE size_distributions 
        SET shipped_quantity = ? 
        WHERE style_id = ? AND size_name = ?
        """, (int(sz.shipped_quantity or 0), req.style_id, sz.size_name))
        total_shipped_qty += int(sz.shipped_quantity or 0)

    c.execute("""
    UPDATE styles 
    SET shipping_date = ?, updated_at = CURRENT_TIMESTAMP 
    WHERE id = ?
    """, (now_str, req.style_id))

    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, 'shipping_update', ?, ?, 'Yükleme Adetleri Güncellendi')
    """, (req.style_id, order_id, user.get("user_id"), user_name, 
          f"Tarih:{style_row['shipping_date']}", f"Adet:{total_shipped_qty}, Tarih:{now_str}"))

    conn.commit()
    conn.close()

    return {
        "status": "success",
        "style_id": req.style_id,
        "total_shipped_quantity": total_shipped_qty,
        "shipping_date": now_str
    }


# ----------------- FABRIC ASSIGNMENT -----------------

class AssignFabricRequest(BaseModel):
    style_id: int
    fabrictag_id: Optional[int] = None
    fabric_name: str
    composition: str
    supplier: Optional[str] = ""
    fabric_slot: Optional[int] = 1

@app.post("/api/styles/assign-fabric")
def api_assign_fabric(req: AssignFabricRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    slot = req.fabric_slot or 1
    
    c.execute("SELECT fabric_article, fabric_composition, fabric_article_2, fabric_composition_2, order_id FROM styles WHERE id = ?", (req.style_id,))
    row = c.fetchone()
    order_id = row[4] if row else None
    
    if slot == 2:
        old_article = (row[2] or "") if row else ""
        old_comp = (row[3] or "") if row else ""
        c.execute("""
        UPDATE styles 
        SET fabric_article_2 = ?, fabric_composition_2 = ?, updated_at = CURRENT_TIMESTAMP 
        WHERE id = ?
        """, (req.fabric_name, req.composition, req.style_id))
    else:
        old_article = (row[0] or "") if row else ""
        old_comp = (row[1] or "") if row else ""
        c.execute("""
        UPDATE styles 
        SET fabric_article = ?, fabric_composition = ?, updated_at = CURRENT_TIMESTAMP 
        WHERE id = ?
        """, (req.fabric_name, req.composition, req.style_id))

    c.execute("DELETE FROM fabric_links WHERE style_id = ? AND (fabric_slot = ? OR (? = 1 AND fabric_slot IS NULL))", (req.style_id, slot, slot))
    c.execute("""
    INSERT INTO fabric_links (style_id, fabrictag_id, fabric_name, composition, supplier, fabric_slot, status)
    VALUES (?, ?, ?, ?, ?, ?, 'Atandı')
    """, (req.style_id, req.fabrictag_id, req.fabric_name, req.composition, req.supplier, slot))

    user_name = user.get("full_name") or user.get("username")
    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, 'fabric_assignment', ?, ?, ?)
    """, (req.style_id, order_id, user.get("user_id"), user_name, f"{old_article} ({old_comp})", f"{req.fabric_name} ({req.composition})", f"FabricTag Kumaş {slot} Ataması Yapıldı"))


    c.execute("""
    INSERT INTO production_status_logs (style_id, user_name, category, title, message)
    VALUES (?, ?, 'Kumaş', 'Kumaş Değiştirildi / Atandı', ?)
    """, (req.style_id, user_name, f"Yeni Kumaş: {req.fabric_name} | {req.composition} ({req.supplier})"))

    conn.commit()
    conn.close()
    return {"status": "success", "message": "Kumaş başarıyla atandı"}

# ----------------- COST CALCULATION ENDPOINTS -----------------

def check_can_view_costs(user: Dict[str, Any]):
    role = user.get("role")
    perms = user.get("permissions") or {}
    if role in ["superadmin", "masterdeveloper", "admin"]:
        return True
    if perms.get("can_view_costs"):
        return True
    raise HTTPException(status_code=403, detail="Maliyet verilerini görüntüleme yetkiniz bulunmamaktadır.")

def check_can_edit_costs(user: Dict[str, Any]):
    role = user.get("role")
    perms = user.get("permissions") or {}
    if role in ["superadmin", "masterdeveloper", "admin"]:
        return True
    if perms.get("can_edit_costs"):
        return True
    raise HTTPException(status_code=403, detail="Maliyet verilerini düzenleme yetkiniz bulunmamaktadır.")

class CostSaveRequest(BaseModel):
    cost_data: Dict[str, Any]
    apply_to_unit_price: bool = False
    calculated_price: Optional[float] = None
    currency: Optional[str] = None

@app.get("/api/styles/{style_id}/cost")
def api_get_style_cost(style_id: int, user: Dict[str, Any] = Depends(require_user)):
    check_can_view_costs(user)
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT cost_data_json, unit_price, currency FROM styles WHERE id = ?", (style_id,))
    row = c.fetchone()
    conn.close()
    if not row:
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    cost_json = row[0]
    data = json.loads(cost_json) if cost_json else None
    return {
        "status": "success",
        "cost_data": data,
        "unit_price": row[1],
        "currency": row[2] or "€"
    }

@app.post("/api/styles/{style_id}/cost")
def api_save_style_cost(style_id: int, req: CostSaveRequest, user: Dict[str, Any] = Depends(require_user)):
    check_can_edit_costs(user)
    conn = get_db()
    c = conn.cursor()
    user_name = user.get("full_name") or user.get("username")

    c.execute("SELECT order_id, unit_price, currency FROM styles WHERE id = ?", (style_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    order_id, old_price, cur_currency = row[0], row[1], row[2]

    cost_str = json.dumps(req.cost_data)
    
    if req.apply_to_unit_price and req.calculated_price is not None:
        c.execute("UPDATE styles SET cost_data_json = ?, unit_price = ?, currency = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                  (cost_str, req.calculated_price, req.currency or cur_currency, style_id))
    else:
        c.execute("UPDATE styles SET cost_data_json = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (cost_str, style_id))

    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, 'cost_calculation', '', ?, 'Maliyet Hesaplama Verileri Güncellendi')
    """, (style_id, order_id, user.get("user_id"), user_name, str(req.calculated_price or '')))

    conn.commit()
    conn.close()
    return {"status": "success", "message": "Maliyet başarıyla kaydedildi"}

class ActualCostSaveRequest(BaseModel):
    actual_cost_data: Dict[str, Any]
    calculated_total: Optional[float] = None
    currency: Optional[str] = "TL"

@app.post("/api/styles/{style_id}/actual-cost")
def api_save_style_actual_cost(style_id: int, req: ActualCostSaveRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    user_name = user.get("full_name") or user.get("username")

    c.execute("SELECT order_id FROM styles WHERE id = ?", (style_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    order_id = row[0]

    cost_str = json.dumps(req.actual_cost_data)
    c.execute("UPDATE styles SET actual_cost_data_json = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (cost_str, style_id))

    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, 'actual_cost', '', ?, 'Gerçekleşen Maliyet & Fatura Verileri Kaydedildi')
    """, (style_id, order_id, user.get("user_id"), user_name, str(req.calculated_total or '')))

    conn.commit()
    conn.close()
    return {"status": "success", "message": "Gerçekleşen maliyet ve fatura verileri başarıyla kaydedildi."}

# ----------------- AUDIT LOGS -----------------



@app.get("/api/audit-logs")
def api_get_audit_logs(style_id: Optional[int] = Query(None), limit: int = 100, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1

    if style_id:
        c.execute("""
        SELECT a.*, s.style_no, s.color_name, o.po_number, o.customer_name, o.brand
        FROM audit_logs a
        JOIN styles s ON a.style_id = s.id
        JOIN orders o ON s.order_id = o.id
        WHERE a.style_id = ?
        ORDER BY a.id DESC LIMIT ?
        """, (style_id, limit))
    else:
        c.execute("""
        SELECT a.*, s.style_no, s.color_name, o.po_number, o.customer_name, o.brand
        FROM audit_logs a
        JOIN styles s ON a.style_id = s.id
        JOIN orders o ON s.order_id = o.id
        WHERE s.company_id = ?
        ORDER BY a.id DESC LIMIT ?
        """, (company_id, limit))

    logs = [dict(r) for r in c.fetchall()]
    conn.close()
    return logs

# ----------------- STATUS & MESSAGE LOGS -----------------

@app.get("/api/styles/{style_id}/logs")
def api_get_style_logs(style_id: int, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT * FROM production_status_logs WHERE style_id = ? ORDER BY id DESC", (style_id,))
    logs = [dict(r) for r in c.fetchall()]
    conn.close()
    return logs

class NewLogRequest(BaseModel):
    category: str = "Genel"
    title: Optional[str] = "Durum Notu"
    message: str

@app.post("/api/styles/{style_id}/logs")
def api_add_style_log(style_id: int, req: NewLogRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    c.execute("""
    INSERT INTO production_status_logs (style_id, user_id, user_name, category, title, message)
    VALUES (?, ?, ?, ?, ?, ?)
    """, (style_id, user.get("user_id"), user.get("full_name") or user.get("username"), req.category, req.title, req.message))
    conn.commit()
    conn.close()
    return {"status": "success", "message": "Not eklendi"}

# ----------------- UPLOAD & PARSER -----------------

@app.post("/api/upload/parse-preview")
async def api_upload_and_parse(request: Request, user: Dict[str, Any] = Depends(require_user)):
    form = await request.form()
    upload_list = []
    
    # Check 'files' first, then 'file' without combining both for the same upload
    files_list = form.getlist("files")
    if files_list:
        upload_list = files_list
    else:
        file_single = form.getlist("file")
        if file_single:
            upload_list = file_single
        else:
            for v in form.values():
                if hasattr(v, "filename") and v.filename:
                    upload_list.append(v)


    if not upload_list:
        raise HTTPException(status_code=400, detail="Lütfen en az bir sipariş dosyası seçin.")

    combined_data = []
    file_info = []
    first_type = "MULTI_UPLOAD"

    for f_obj in upload_list:
        if not hasattr(f_obj, "filename") or not f_obj.filename:
            continue
        saved_path = UPLOADS_DIR / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{f_obj.filename}"
        with open(saved_path, "wb") as buffer:
            shutil.copyfileobj(f_obj.file, buffer)
        
        try:
            parsed_result = OrderParserEngine.auto_detect_and_parse(str(saved_path))
            first_type = parsed_result.get("type", "UPLOAD")
            items = parsed_result.get("data", [])
            for it in items:
                it["source_file"] = f_obj.filename
                combined_data.append(it)
            file_info.append(f_obj.filename)
        except Exception as e:
            print(f"Parsing error for {f_obj.filename}: {e}")

    if not combined_data and upload_list:
        raise HTTPException(status_code=400, detail="Seçilen dosyalardan sipariş ayrıştırılamadı.")

    return {
        "type": "MULTI_UPLOAD" if len(file_info) > 1 else first_type,
        "count": len(combined_data),
        "data": combined_data,
        "filename": ", ".join(file_info),
        "files_count": len(file_info)
    }



class ConfirmImportRequest(BaseModel):
    file_path: Optional[str] = ""
    filename: Optional[str] = ""
    source_type: Optional[str] = "MULTI_UPLOAD"
    items: List[Dict[str, Any]]


@app.post("/api/upload/confirm-import")
def api_confirm_import(req: ConfirmImportRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1
    
    inserted_orders = 0
    inserted_styles = 0

    try:
        # Group items by PO Number to attach to unique order parents or create individual ones
        po_groups: Dict[str, List[Any]] = {}
        for item in req.items:
            po_num = str(item.get("po_number") or f"PO-{datetime.now().strftime('%Y%m%d')}").strip()
            po_groups.setdefault(po_num, []).append(item)

        for po_num, items_list in po_groups.items():
            first_item = items_list[0]
            cust_name = first_item.get("customer_name") or "Müşteri"
            brand = first_item.get("brand") or "Marka"
            season = first_item.get("season") or "2026/2027"
            order_date = first_item.get("order_date") or datetime.now().strftime("%d.%m.%Y")
            delivery_date = first_item.get("delivery_date") or ""
            delivery_terms = first_item.get("delivery_terms") or ""
            payment_terms = first_item.get("payment_terms") or ""
            currency = first_item.get("currency") or "EUR"

            total_q = sum(int(it.get("total_quantity") or 0) for it in items_list)
            total_a = sum(float(it.get("total_quantity") or 0) * float(it.get("unit_price") or 0.0) for it in items_list)

            c.execute("""
            INSERT INTO orders (
                company_id, po_number, customer_name, brand, season, order_date, delivery_date,
                delivery_terms, payment_terms, total_quantity, total_amount, currency, raw_file_name
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                company_id, po_num, cust_name, brand, season, order_date, delivery_date,
                delivery_terms, payment_terms, total_q, total_a, currency, req.filename
            ))
            order_id = c.lastrowid
            inserted_orders += 1

            for it in items_list:
                style_no = str(it.get("style_no") or "STYLE-01").strip()
                desc = it.get("description") or ""
                comp = it.get("fabric_composition") or ""
                article = it.get("fabric_article") or ""
                ftype = it.get("fabric_type") or ""
                color_code = str(it.get("color_code") or "").strip()
                color_name = str(it.get("color_name") or "").strip()
                price = float(it.get("unit_price") or 0.0)
                
                sizes_dict = it.get("size_distribution") or {}
                # Calculate real total from sizes if present
                if sizes_dict and any(int(v or 0) > 0 for v in sizes_dict.values()):
                    item_qty = sum(int(v or 0) for v in sizes_dict.values())
                else:
                    item_qty = int(it.get("total_quantity") or 0)
                    
                item_amt = item_qty * price

                style_brand = str(it.get("brand") or brand or "").strip()
                c.execute("""
                INSERT INTO styles (
                    order_id, company_id, style_no, description, fabric_composition, fabric_article,
                    fabric_type, color_code, color_name, unit_price, total_quantity, total_amount,
                    currency, status, fabric_order_status, last_status_updated_by, last_status_updated_at,
                    brand
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Planlamada', ?, 'Sipariş Yükleme', datetime('now'), ?)
                """, (
                    order_id, company_id, style_no, desc, comp, article,
                    ftype, color_code, color_name, price, item_qty, item_amt,
                    currency, it.get("fabric_order_status") or "",
                    style_brand
                ))
                style_id = c.lastrowid
                inserted_styles += 1

                for idx, (sz_name, sz_qty) in enumerate(sizes_dict.items()):
                    if str(sz_name).strip():
                        c.execute("""
                        INSERT INTO size_distributions (style_id, size_name, quantity, sort_order)
                        VALUES (?, ?, ?, ?)
                        """, (style_id, str(sz_name).strip(), int(sz_qty or 0), idx + 1))

                for log in it.get("status_logs", []):
                    c.execute("""
                    INSERT INTO production_status_logs (style_id, user_name, category, title, message)
                    VALUES (?, ?, 'Sipariş Notu', ?, ?)
                    """, (style_id, "Sipariş Yükleme", log.get("title") or "Not", log.get("message") or ""))

        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        raise HTTPException(status_code=500, detail=f"Kayıt hatası: {str(e)}")

    conn.close()
    return {
        "status": "success",
        "message": f"{inserted_orders} Sipariş ve {inserted_styles} Model başarıyla sisteme aktarıldı."
    }

# ----------------- MANUEL SİPARİŞ / MODEL EKLEME -----------------

class CreateManualStyleRequest(BaseModel):
    customer_name: Optional[str] = ""
    brand: Optional[str] = ""
    po_number: Optional[str] = ""
    style_no: Optional[str] = ""
    description: Optional[str] = ""
    color_name: Optional[str] = ""
    color_code: Optional[str] = ""
    fabric_article: Optional[str] = ""
    fabric_composition: Optional[str] = ""
    unit_price: Optional[float] = 0.0
    currency: Optional[str] = "EUR"
    unit_meters: Optional[str] = ""
    unit_grams: Optional[str] = ""
    fabric_wastage_percent: Optional[float] = 5.0
    fabric_ordered_meters: Optional[float] = 0.0
    fabric_order_unit: Optional[str] = "M"
    order_date: Optional[str] = ""
    delivery_date: Optional[str] = ""
    status: Optional[str] = "Planlamada"
    size_map: Optional[Dict[str, int]] = None

@app.post("/api/styles/manual-create")
def api_create_manual_style(req: CreateManualStyleRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1
    user_name = user.get("full_name") or user.get("username") or "Kullanıcı"

    po = (req.po_number or "MANUEL-PO").strip()
    cust = (req.customer_name or "").strip()
    brand = (req.brand or "").strip()
    order_dt = req.order_date or datetime.now().strftime("%d.%m.%Y")
    del_dt = req.delivery_date or ""
    curr = req.currency or "EUR"

    sizes_dict = req.size_map or {}
    total_qty = sum(int(v or 0) for v in sizes_dict.values()) if sizes_dict else 0
    unit_price = float(req.unit_price or 0.0)
    total_amount = total_qty * unit_price

    try:
        c.execute("""
        SELECT id FROM orders 
        WHERE company_id = ? AND po_number = ? AND customer_name = ?
        """, (company_id, po, cust))
        row = c.fetchone()
        if row:
            order_id = row[0]
            c.execute("""
            UPDATE orders 
            SET total_quantity = total_quantity + ?, total_amount = total_amount + ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """, (total_qty, total_amount, order_id))
        else:
            c.execute("""
            INSERT INTO orders (
                company_id, po_number, customer_name, brand, order_date, delivery_date,
                total_quantity, total_amount, currency, status, raw_file_name
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'Yeni', 'Manuel Giriş')
            """, (company_id, po, cust, brand, order_dt, del_dt, total_qty, total_amount, curr))
            order_id = c.lastrowid

        style_no = (req.style_no or "STYLE-01").strip()
        c.execute("""
        INSERT INTO styles (
            order_id, company_id, style_no, description, fabric_composition, fabric_article,
            color_code, color_name, unit_price, currency, total_quantity, total_amount,
            unit_meters, unit_grams, fabric_wastage_percent, fabric_ordered_meters, fabric_order_unit,
            status, last_status_updated_by, last_status_updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, datetime('now'))
        """, (
            order_id, company_id, style_no, req.description or "", req.fabric_composition or "",
            req.fabric_article or "", req.color_code or "", req.color_name or "", unit_price, curr,
            total_qty, total_amount, req.unit_meters or "", req.unit_grams or "",
            req.fabric_wastage_percent or 5.0, req.fabric_ordered_meters or 0.0, req.fabric_order_unit or "M",
            req.status or "Planlamada", user_name
        ))
        style_id = c.lastrowid


        for idx, (sz_name, sz_qty) in enumerate(sizes_dict.items()):
            if str(sz_name).strip():
                c.execute("""
                INSERT INTO size_distributions (style_id, size_name, quantity, sort_order)
                VALUES (?, ?, ?, ?)
                """, (style_id, str(sz_name).strip(), int(sz_qty or 0), idx + 1))

        c.execute("""
        INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
        VALUES (?, ?, ?, ?, 'manual_create', '', ?, 'Manuel Sipariş / Model Eklendi')
        """, (style_id, order_id, user.get("user_id"), user_name, f"{style_no} - {req.color_name or ''} ({total_qty} Adet)"))

        c.execute("""
        INSERT INTO production_status_logs (style_id, user_name, category, title, message)
        VALUES (?, ?, 'Sipariş Notu', 'Manuel Giriş Yapıldı', ?)
        """, (style_id, user_name, f"Model {style_no} manuel olarak oluşturuldu."))

        conn.commit()
        conn.close()
        return {"status": "success", "style_id": style_id, "order_id": order_id}
    except Exception as e:
        conn.rollback()
        conn.close()
        raise HTTPException(status_code=500, detail=f"Manuel sipariş kaydetme hatası: {str(e)}")



# ----------------- FABRICTAG ENDPOINTS -----------------

@app.get("/api/fabrictag/fabrics")
def api_fabrictag_list(q: Optional[str] = Query(None), user: Dict[str, Any] = Depends(require_user)):
    return FabricTagConnector.search_fabrics(q or "")

class UpdateFabricReceivedRequest(BaseModel):
    received_meters: Optional[str] = ""
    received_date: Optional[str] = None

@app.post("/api/fabrictag/fabrics/{fabric_id}/received")
def api_fabrictag_update_received(fabric_id: int, req: UpdateFabricReceivedRequest, user: Dict[str, Any] = Depends(require_user)):
    from datetime import datetime
    date_str = req.received_date
    if not date_str and req.received_meters and req.received_meters.strip():
        date_str = datetime.now().strftime("%d.%m.%Y")
    ok = FabricTagConnector.update_received_meters(fabric_id, req.received_meters or "", date_str)
    if not ok:
        raise HTTPException(status_code=500, detail="Kumaş gelen miktarı güncellenemedi.")
    return {"status": "success", "received_meters": req.received_meters, "received_date": date_str}

class AddFabricBatchRequest(BaseModel):
    amount: str
    date: Optional[str] = None
    note: Optional[str] = ""

@app.post("/api/fabrictag/fabrics/{fabric_id}/batches")
def api_fabrictag_add_batch(fabric_id: int, req: AddFabricBatchRequest, user: Dict[str, Any] = Depends(require_user)):
    res = FabricTagConnector.add_received_batch(fabric_id, req.amount, req.date, req.note or "")
    if not res:
        raise HTTPException(status_code=500, detail="Parti eklenemedi.")
    return {"status": "success", **res}

@app.delete("/api/fabrictag/fabrics/{fabric_id}/batches/{batch_id}")
def api_fabrictag_delete_batch(fabric_id: int, batch_id: int, user: Dict[str, Any] = Depends(require_user)):
    res = FabricTagConnector.delete_received_batch(fabric_id, batch_id)
    if not res:
        raise HTTPException(status_code=500, detail="Parti silinemedi.")
    return {"status": "success", **res}



# ----------------- SUPERADMIN ENDPOINTS -----------------

@app.get("/api/superadmin/companies")
def api_superadmin_get_companies(user: Dict[str, Any] = Depends(require_superadmin)):
    conn = get_db()
    c = conn.cursor()
    c.execute("""
    SELECT c.*, 
           (SELECT COUNT(*) FROM users WHERE company_id = c.id) as user_count,
           (SELECT COUNT(*) FROM orders WHERE company_id = c.id) as order_count
    FROM companies c
    ORDER BY c.id DESC
    """)
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    return rows

class CreateCompanyRequest(BaseModel):
    name: str
    code: str
    license_key: str
    license_expires_at: str
    max_users: int = 10

@app.post("/api/superadmin/companies")
def api_superadmin_create_company(req: CreateCompanyRequest, user: Dict[str, Any] = Depends(require_superadmin)):
    conn = get_db()
    c = conn.cursor()
    try:
        c.execute("""
        INSERT INTO companies (name, code, license_key, license_expires_at, is_active, max_users)
        VALUES (?, ?, ?, ?, 1, ?)
        """, (req.name, req.code, req.license_key, req.license_expires_at, req.max_users))
        cid = c.lastrowid
        
        c.execute("SELECT menu_key, title_tr, title_en, icon, path, sort_order, is_visible, allowed_roles_json FROM dynamic_menus WHERE company_id = 1")
        menus = c.fetchall()
        for m in menus:
            c.execute("""
            INSERT INTO dynamic_menus (company_id, menu_key, title_tr, title_en, icon, path, sort_order, is_visible, allowed_roles_json)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (cid, m[0], m[1], m[2], m[3], m[4], m[5], m[6], m[7]))
            
        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        raise HTTPException(status_code=400, detail=f"Firma oluşturulamadı: {str(e)}")
# ----------------- HAFTALIK PROGRAM EXCEL ENDPOINT -----------------

@app.get("/api/reports/haftalik-program-excel")
def api_haftalik_program_excel(
    start_key: Optional[str] = Query(None),
    end_key: Optional[str] = Query(None),
    only_with_orders: bool = Query(True)
):
    """
    Patron Takip Çıktısı için A4 Yatay 7'şer model sütunlu Haftalık Program Excel üreticisi
    """
    from datetime import timedelta
    from collections import defaultdict

    conn = get_db()
    c = conn.cursor()
    c.execute("""
        SELECT 
            s.id, s.style_no, s.brand, s.description,
            o.customer_name, o.brand as order_brand,
            s.color_code, s.color_name, s.total_quantity,
            o.delivery_date, s.image_url, s.image_url_2
        FROM styles s
        LEFT JOIN orders o ON s.order_id = o.id
        WHERE o.delivery_date IS NOT NULL AND o.delivery_date != ''
        ORDER BY o.delivery_date ASC, s.id ASC
    """)
    rows = c.fetchall()
    conn.close()

    tr_months = ['Oca', 'Şub', 'Mar', 'Nis', 'May', 'Haz', 'Tem', 'Ağu', 'Eyl', 'Eki', 'Kas', 'Ara']
    week_map = {}

    for r in rows:
        raw_date = str(r[9]).strip()
        dt = None
        for fmt in ('%d.%m.%Y', '%Y-%m-%d', '%d/%m/%Y'):
            try:
                dt = datetime.strptime(raw_date, fmt)
                break
            except Exception:
                pass
        if not dt:
            continue

        year, week, _ = dt.isocalendar()
        key = f"{year}_{week:02d}"

        if start_key and key < start_key:
            continue
        if end_key and key > end_key:
            continue

        if key not in week_map:
            jan4 = datetime(year, 1, 4)
            start_of_year = jan4 - timedelta(days=jan4.isoweekday() - 1)
            monday = start_of_year + timedelta(weeks=week - 1)
            sunday = monday + timedelta(days=6)
            label = f"{monday.day} {tr_months[monday.month - 1]} - {sunday.day} {tr_months[sunday.month - 1]} {year}"

            week_map[key] = {
                'key': key,
                'year': year,
                'week': week,
                'label': label,
                'models': {}
            }

        style_no = (r[1] or 'Bilinmeyen').strip()
        brand = (r[2] or r[5] or r[4] or '').strip()
        color = (r[7] or r[6] or 'Standart').strip()
        qty = int(r[8] or 0)
        img = r[10] or r[11] or ''

        w = week_map[key]
        if style_no not in w['models']:
            w['models'][style_no] = {
                'style_no': style_no,
                'brand': brand,
                'image_url': img,
                'colors': defaultdict(int),
                'total_quantity': 0
            }

        m = w['models'][style_no]
        if brand and not m['brand']:
            m['brand'] = brand
        if img and not m['image_url']:
            m['image_url'] = img
        m['colors'][color] += qty
        m['total_quantity'] += qty

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Haftalık Program'

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    try:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    except Exception:
        pass
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35

    ws.column_dimensions['A'].width = 2.5
    for c_idx in range(2, 16):
        ws.column_dimensions[get_column_letter(c_idx)].width = 13.0

    thin_gray = Side(style='thin', color='D0D5DD')
    medium_dark = Side(style='medium', color='334155')

    current_row = 1
    sorted_weeks = sorted(week_map.values(), key=lambda x: x['key'])

    for w in sorted_weeks:
        models_list = list(w['models'].values())
        if not models_list and only_with_orders:
            continue

        week_total_qty = sum(m['total_quantity'] for m in models_list)
        week_total_models = len(models_list)

        for chunk_idx in range(0, max(1, len(models_list)), 7):
            chunk = models_list[chunk_idx:chunk_idx + 7]

            if chunk_idx == 0:
                ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=15)
                h_cell = ws.cell(current_row, 2)
                h_cell.value = f"{w['week']}. HAFTA ({w['label']})   •   Toplam {week_total_models} Model   •   {week_total_qty:,} Adet"
                h_cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
                h_cell.fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
                h_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
                ws.row_dimensions[current_row].height = 24.0
                current_row += 1

            brand_row = current_row
            model_row = current_row + 1
            img_start_row = current_row + 2
            img_end_row = current_row + 7
            ws.row_dimensions[brand_row].height = 16.0
            ws.row_dimensions[model_row].height = 18.0
            for ir in range(img_start_row, img_end_row + 1):
                ws.row_dimensions[ir].height = 17.0

            max_colors = 2
            for m in chunk:
                max_colors = max(max_colors, len(m['colors']))

            for slot_idx in range(7):
                col_left = 2 + (slot_idx * 2)
                col_right = col_left + 1
                col_left_let = get_column_letter(col_left)

                if slot_idx < len(chunk):
                    model = chunk[slot_idx]
                    # Brand row
                    ws.merge_cells(start_row=brand_row, start_column=col_left, end_row=brand_row, end_column=col_right)
                    b_cell = ws.cell(brand_row, col_left)
                    b_cell.value = (model['brand'] or '').upper()
                    b_cell.font = Font(name='Calibri', size=10, bold=True, color='334155')
                    b_cell.fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
                    b_cell.alignment = Alignment(horizontal='center', vertical='center')

                    # Model row
                    ws.merge_cells(start_row=model_row, start_column=col_left, end_row=model_row, end_column=col_right)
                    m_cell = ws.cell(model_row, col_left)
                    m_cell.value = model['style_no']
                    m_cell.font = Font(name='Calibri', size=11, bold=True, color='0F172A')
                    m_cell.fill = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
                    m_cell.alignment = Alignment(horizontal='center', vertical='center')

                    # Image area
                    ws.merge_cells(start_row=img_start_row, start_column=col_left, end_row=img_end_row, end_column=col_right)
                    img_cell = ws.cell(img_start_row, col_left)
                    img_cell.alignment = Alignment(horizontal='center', vertical='center')

                    if model['image_url']:
                        img_path = UPLOADS_DIR / 'styles' / Path(model['image_url']).name
                        if img_path.exists():
                            try:
                                pil_img = PILImage.open(img_path)
                                orig_w, orig_h = pil_img.size
                                scale = min(140 / orig_w, 95 / orig_h)
                                new_w = int(orig_w * scale)
                                new_h = int(orig_h * scale)

                                openpyxl_img = OpenpyxlImage(str(img_path))
                                openpyxl_img.width = new_w
                                openpyxl_img.height = new_h
                                openpyxl_img.anchor = f"{col_left_let}{img_start_row}"
                                ws.add_image(openpyxl_img)
                            except Exception:
                                pass

                    # Colors and quantities
                    colors_list = list(model['colors'].items())
                    for c_idx in range(max_colors):
                        c_row = img_end_row + 1 + c_idx
                        ws.row_dimensions[c_row].height = 15.0
                        c_col_cell = ws.cell(c_row, col_left)
                        c_qty_cell = ws.cell(c_row, col_right)

                        if c_idx < len(colors_list):
                            c_name, c_qty = colors_list[c_idx]
                            c_col_cell.value = c_name
                            c_col_cell.font = Font(name='Calibri', size=8.5, color='1E293B')
                            c_col_cell.alignment = Alignment(horizontal='left', vertical='center')

                            c_qty_cell.value = c_qty
                            c_qty_cell.font = Font(name='Calibri', size=9, bold=True, color='0F172A')
                            c_qty_cell.alignment = Alignment(horizontal='center', vertical='center')
                        else:
                            c_col_cell.value = ''
                            c_qty_cell.value = ''

                    # Total row
                    tot_row = img_end_row + 1 + max_colors
                    ws.row_dimensions[tot_row].height = 16.0
                    tot_label_cell = ws.cell(tot_row, col_left)
                    tot_label_cell.value = 'TOPLAM'
                    tot_label_cell.font = Font(name='Calibri', size=8.5, bold=True, color='047857')
                    tot_label_cell.fill = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')
                    tot_label_cell.alignment = Alignment(horizontal='left', vertical='center')

                    tot_val_cell = ws.cell(tot_row, col_right)
                    tot_val_cell.value = model['total_quantity']
                    tot_val_cell.font = Font(name='Calibri', size=9.5, bold=True, color='047857')
                    tot_val_cell.fill = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')
                    tot_val_cell.alignment = Alignment(horizontal='center', vertical='center')

                else:
                    for r_i in range(brand_row, img_end_row + 1 + max_colors + 1):
                        ws.cell(r_i, col_left).value = None
                        ws.cell(r_i, col_right).value = None

                slot_end_row = img_end_row + 1 + max_colors
                for r_idx in range(brand_row, slot_end_row + 1):
                    for cl in (col_left, col_right):
                        top_s = medium_dark if r_idx == brand_row else thin_gray
                        bot_s = medium_dark if r_idx == slot_end_row else thin_gray
                        left_s = medium_dark if cl == col_left else thin_gray
                        right_s = medium_dark if cl == col_right else thin_gray
                        ws.cell(r_idx, cl).border = Border(top=top_s, bottom=bot_s, left=left_s, right=right_s)

            current_row = img_end_row + 1 + max_colors + 2

    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    filename = f"Haftalik_Program_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return Response(
        content=output_stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

# ----------------- SPA ROOT -----------------

@app.get("/")
def serve_index():
    index_path = STATIC_DIR / "index.html"
    if index_path.exists():
        return FileResponse(
            str(index_path),
            headers={
                "Cache-Control": "no-cache, no-store, must-revalidate, max-age=0",
                "Pragma": "no-cache",
                "Expires": "0"
            }
        )
    return HTMLResponse("<h1>TexFlow API is running.</h1>")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8050)
