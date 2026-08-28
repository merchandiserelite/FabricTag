import sqlite3
import csv
import os
import re
import shutil
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(__file__)
DB_FILE = os.path.join(BASE_DIR, "database.db")
BACKUPS_DIR = os.path.join(BASE_DIR, "backups")

# Detect OneDrive for automatic real-time cloud syncing
USER_HOME = os.path.expanduser("~")
ONEDRIVE_DIR = os.path.join(USER_HOME, "OneDrive")
ONEDRIVE_SOUL_WAYS = os.path.join(ONEDRIVE_DIR, "OneDrive - Soul ways")

if os.path.exists(ONEDRIVE_SOUL_WAYS):
    ONEDRIVE_BACKUP_DIR = os.path.join(ONEDRIVE_SOUL_WAYS, "FabricTag_Cloud_Backups")
elif os.path.exists(ONEDRIVE_DIR):
    ONEDRIVE_BACKUP_DIR = os.path.join(ONEDRIVE_DIR, "FabricTag_Cloud_Backups")
else:
    ONEDRIVE_BACKUP_DIR = None

os.makedirs(BACKUPS_DIR, exist_ok=True)
if ONEDRIVE_BACKUP_DIR:
    try:
        os.makedirs(ONEDRIVE_BACKUP_DIR, exist_ok=True)
    except Exception:
        pass

def get_db_connection():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS fabrics (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            internal_code TEXT UNIQUE NOT NULL,
            company_name TEXT,
            quality_code TEXT,
            quality_name TEXT,
            design_code TEXT,
            width TEXT,
            weight TEXT,
            composition TEXT,
            barcode_or_qr TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.commit()
    conn.close()

def load_code_settings():
    try:
        import json
        settings_path = os.path.join(BASE_DIR, "settings.json")
        if os.path.exists(settings_path):
            with open(settings_path, "r", encoding="utf-8") as f:
                s = json.load(f)
                return {
                    "prefix_type": s.get("code_prefix_type", "letters"),
                    "prefix_text": s.get("code_prefix_text", "ELT"),
                    "separator": s.get("code_separator", ""),
                    "digits": int(s.get("code_digits", 7)),
                    "start_num": int(s.get("code_start_number", 1))
                }
    except Exception:
        pass
    return {
        "prefix_type": "letters",
        "prefix_text": "ELT",
        "separator": "",
        "digits": 7,
        "start_num": 1
    }

def get_next_internal_code(conn, code_settings=None):
    if code_settings is None:
        code_settings = load_code_settings()
        
    prefix_type = code_settings.get("prefix_type", "letters")
    prefix_text = str(code_settings.get("prefix_text", "ELT")).strip().upper()
    separator = str(code_settings.get("separator", ""))
    digits = int(code_settings.get("digits", 7))
    start_num = int(code_settings.get("start_num", 1))
    
    clean_prefix = prefix_text if prefix_type == "letters" else ""
    full_prefix = f"{clean_prefix}{separator}" if clean_prefix else ""
    
    cursor = conn.cursor()
    cursor.execute("SELECT internal_code FROM fabrics")
    rows = cursor.fetchall()
    
    max_num = 0
    for row in rows:
        code = str(row["internal_code"]).strip().upper()
        if full_prefix and code.startswith(full_prefix):
            rest = code[len(full_prefix):]
            if rest.isdigit():
                num = int(rest)
                if num > max_num:
                    max_num = num
        elif not full_prefix and code.isdigit():
            num = int(code)
            if num > max_num:
                max_num = num
        else:
            m = re.search(r'(\d+)$', code)
            if m:
                num = int(m.group(1))
                if num > max_num:
                    max_num = num
                    
    if max_num == 0:
        next_num = start_num
    else:
        next_num = max(max_num + 1, start_num)
        
    if digits > 0:
        formatted_num = f"{next_num:0{digits}d}"
    else:
        formatted_num = str(next_num)
        
    return f"{full_prefix}{formatted_num}"

def add_fabric(company_name, quality_code, quality_name, design_code, width, weight, composition, barcode_or_qr, internal_code=None):
    conn = get_db_connection()
    cursor = conn.cursor()
    
    company_name = company_name.strip().upper() if company_name else "GENEL"
    quality_code = quality_code.strip().upper() if quality_code else "KODSUZ"
    quality_name = quality_name.strip().upper() if quality_name else ""
    design_code = design_code.strip().upper() if design_code else ""
    width = width.strip().upper() if width else ""
    weight = weight.strip().upper() if weight else ""
    barcode_or_qr = barcode_or_qr.strip().upper() if barcode_or_qr else ""
    
    if composition:
        comp_clean = re.sub(r'[/,;+]+', ' ', composition)
        composition = re.sub(r'\s+', ' ', comp_clean).strip().upper()
    else:
        composition = ""

    # 1. If internal_code is provided and exists in DB, UPDATE existing record
    if internal_code:
        cursor.execute("SELECT * FROM fabrics WHERE internal_code = ?", (internal_code.strip().upper(),))
        existing_by_code = cursor.fetchone()
        if existing_by_code:
            cursor.execute("""
                UPDATE fabrics SET
                    company_name = ?,
                    quality_code = ?,
                    quality_name = ?,
                    design_code = ?,
                    width = ?,
                    weight = ?,
                    composition = ?,
                    barcode_or_qr = ?
                WHERE internal_code = ?
            """, (company_name, quality_code, quality_name, design_code,
                  width, weight, composition, barcode_or_qr, internal_code.strip().upper()))
            conn.commit()
            cursor.execute("SELECT * FROM fabrics WHERE internal_code = ?", (internal_code.strip().upper(),))
            updated_row = dict(cursor.fetchone())
            updated_row["is_duplicate"] = False
            updated_row["is_updated"] = True
            conn.close()
            return updated_row

    # 2. Otherwise check if (company_name, quality_code) already exists
    if company_name != "GENEL" or quality_code != "KODSUZ":
        cursor.execute(
            "SELECT * FROM fabrics WHERE LOWER(company_name) = LOWER(?) AND LOWER(quality_code) = LOWER(?)",
            (company_name, quality_code)
        )
        existing = cursor.fetchone()
        if existing:
            dup_dict = dict(existing)
            dup_dict["is_duplicate"] = True
            dup_dict["is_updated"] = False
            conn.close()
            return dup_dict

    # 3. New Record Insertion
    new_code = get_next_internal_code(conn)
    cursor.execute("""
        INSERT INTO fabrics (
            internal_code, company_name, quality_code, quality_name, 
            design_code, width, weight, composition, barcode_or_qr
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (new_code, company_name, quality_code, quality_name, 
          design_code, width, weight, composition, barcode_or_qr))
    
    conn.commit()
    cursor.execute("SELECT * FROM fabrics WHERE id = ?", (cursor.lastrowid,))
    new_row = dict(cursor.fetchone())
    new_row["is_duplicate"] = False
    new_row["is_updated"] = False
    conn.close()
    return new_row

def get_all_fabrics(search_query=None, start_date=None, end_date=None):
    conn = get_db_connection()
    cursor = conn.cursor()
    
    conditions = []
    params = []
    
    if search_query:
        query = f"%{search_query}%"
        conditions.append("(internal_code LIKE ? OR company_name LIKE ? OR quality_code LIKE ? OR quality_name LIKE ? OR design_code LIKE ? OR composition LIKE ?)")
        params.extend([query, query, query, query, query, query])
        
    if start_date:
        conditions.append("created_at >= ?")
        params.append(start_date)
        
    if end_date:
        conditions.append("created_at <= ?")
        params.append(end_date)
        
    where_clause = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    sql = f"SELECT * FROM fabrics {where_clause} ORDER BY created_at DESC"
    
    cursor.execute(sql, params)
    rows = cursor.fetchall()
    conn.close()
    return [dict(row) for row in rows]

def export_to_xlsx(filepath, start_date=None, end_date=None):
    rows = get_all_fabrics(start_date=start_date, end_date=end_date)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kumaş Kartela Listesi"
    
    # Enable grid lines in Excel view
    ws.views.sheetView[0].showGridLines = True
    
    headers = [
        "Sıra No",
        "Ic Kod (Internal Code)",
        "Firma Adı (Company)",
        "Kalite Kodu (Article)",
        "Kalite Adı (Name)",
        "Desen Kodu (Design)",
        "Karışım / Yapı (Composition)",
        "En (Width)",
        "Gramaj (Weight)",
        "Barkod / QR",
        "Kayıt Tarihi & Saati"
    ]
    
    # Styling definitions
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    
    # Write headers
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        
    # Write data rows
    for idx, row in enumerate(rows, start=1):
        row_data = [
            idx,
            row["internal_code"],
            row["company_name"],
            row["quality_code"],
            row["quality_name"] or "-",
            row["design_code"] or "-",
            row["composition"] or "-",
            row["width"] or "-",
            row["weight"] or "-",
            row["barcode_or_qr"] or "-",
            row["created_at"] or "-"
        ]
        ws.append(row_data)
        current_row = idx + 1
        ws.row_dimensions[current_row].height = 22
        
        row_fill = zebra_fill if idx % 2 == 0 else white_fill
        
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.fill = row_fill
            cell.border = thin_border
            cell.font = Font(name="Segoe UI", size=10)
            
            # Alignments: codes and dates centered, names left-aligned
            if col_num in [1, 2, 4, 6, 8, 10, 11]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
    wb.save(filepath)
    return len(rows)

def export_to_csv(filepath):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM fabrics ORDER BY created_at DESC")
    rows = cursor.fetchall()
    conn.close()
    
    headers = [
        "İç Kod (Internal Code)", 
        "Firma (Company)", 
        "Kalite Kodu (Article)", 
        "Kalite Adı (Name)", 
        "Desen Kodu (Design)", 
        "En (Width)", 
        "Gramaj (Weight)", 
        "Karışım (Composition)", 
        "Barkod / QR",
        "Kayıt Tarihi (Created At)"
    ]
    
    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(headers)
        for row in rows:
            writer.writerow([
                row["internal_code"],
                row["company_name"],
                row["quality_code"],
                row["quality_name"],
                row["design_code"],
                row["width"],
                row["weight"],
                row["composition"],
                row["barcode_or_qr"],
                row["created_at"]
            ])

def create_timestamped_backup(start_date=None, end_date=None, label_prefix="MANUEL_YEDEK"):
    now_str = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    filename = f"ELITE_Kumas_{label_prefix}_{now_str}.xlsx"
    filepath = os.path.join(BACKUPS_DIR, filename)
    count = export_to_xlsx(filepath, start_date=start_date, end_date=end_date)
    
    # Auto-mirror to OneDrive Cloud Backup if active
    if ONEDRIVE_BACKUP_DIR and os.path.exists(filepath):
        try:
            shutil.copy2(filepath, os.path.join(ONEDRIVE_BACKUP_DIR, filename))
            # Also backup database file
            if os.path.exists(DB_FILE):
                shutil.copy2(DB_FILE, os.path.join(ONEDRIVE_BACKUP_DIR, "database_backup.db"))
        except Exception:
            pass
            
    return {
        "filename": filename,
        "filepath": filepath,
        "record_count": count,
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "size_kb": round(os.path.getsize(filepath) / 1024, 1)
    }

def list_backups():
    if not os.path.exists(BACKUPS_DIR):
        return []
    
    files = []
    for f in os.listdir(BACKUPS_DIR):
        if f.lower().endswith(('.xlsx', '.csv', '.db')):
            f_path = os.path.join(BACKUPS_DIR, f)
            stat = os.stat(f_path)
            files.append({
                "filename": f,
                "created_at": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
                "size_kb": round(stat.st_size / 1024, 1),
                "download_url": f"/api/backups/download/{f}"
            })
            
    files.sort(key=lambda x: x["created_at"], reverse=True)
    return files

def generate_empty_excel_template(filepath):
    """
    Generates a beautifully styled empty Excel template for fabric importing.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kumas_Sablonu"
    
    headers = [
        "Ic Kod (Internal Code)",
        "Firma (Supplier)",
        "Kalite Kodu (Article)",
        "Kalite Adi (Quality Name)",
        "Desen Kodu (Design/Color)",
        "Karisim (Composition)",
        "En (Width - cm)",
        "Gramaj (Weight - g/m2)",
        "Barkod (Barcode)"
    ]
    
    # Header styling
    header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid") # Deep indigo
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    # Write Title Note
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = "FabricTag - Kumas Veri Aktarim Sablonu (Ic kod alanini bos birakirsaniz sistem siradaki otomatik ic kodu atar)"
    title_cell.font = Font(name="Calibri", size=10, italic=True, color="4338CA", bold=True)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    title_cell.fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    ws.row_dimensions[1].height = 25
    
    # Write Headers on row 2
    ws.row_dimensions[2].height = 28
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        
    # Write 2 Sample Rows for Guidance
    sample_rows = [
        ["ELT0000001", "ELITE TEKSTIL", "ART-1024", "IPEK SATEN", "01 LACIVERT", "%95 PAMUK %5 ELASTAN", "145", "220", "ELT0000001"],
        ["", "DOKUMA TEKSTIL", "DK-550", "KETEN KETEN", "12 BEJ", "%100 KETEN", "150", "280", ""]
    ]
    
    sample_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    data_font = Font(name="Calibri", size=10)
    
    for row_idx, row_data in enumerate(sample_rows, 3):
        ws.row_dimensions[row_idx].height = 22
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 7, 8, 9] else "left", vertical="center")
            if row_idx == 3:
                cell.fill = sample_fill

    # Set Column Widths
    col_widths = [24, 26, 22, 24, 22, 30, 16, 22, 22]
    for i, w in enumerate(col_widths, 1):
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = w
        
    wb.save(filepath)
    return filepath

def normalize_header_text(text):
    if not text:
        return ""
    t = str(text).strip().lower()
    t = t.replace("ı", "i").replace("ğ", "g").replace("ü", "u").replace("ş", "s").replace("ö", "o").replace("ç", "c")
    t = re.sub(r'[^a-z0-9]', '', t)
    return t

def detect_column_field(header_name):
    """
    Intelligently determines which standard fabric database field a column header belongs to.
    """
    norm = normalize_header_text(header_name)
    if not norm:
        return None
        
    if any(k in norm for k in ["ickod", "internalcode", "stokkodu", "stokno", "urunkodu", "itemcode"]):
        return "internal_code"
    if any(k in norm for k in ["firma", "tedarikci", "supplier", "company", "vendor", "uretimci", "musteri", "fabrika"]):
        return "company_name"
    if any(k in norm for k in ["kalitekodu", "kaliteno", "article", "artno", "kumascinsi", "kumaskodu", "qualitycode", "desenno"]):
        return "quality_code"
    if any(k in norm for k in ["kaliteadi", "kaliteismi", "qualityname", "kumasadi", "fabricname", "cins", "aciklama", "description"]):
        return "quality_name"
    if any(k in norm for k in ["desenkodu", "desen", "renkkodu", "renkno", "renk", "varyant", "color", "colour", "design", "designcode", "variant"]):
        return "design_code"
    if any(k in norm for k in ["karisim", "kompozisyon", "composition", "icerik", "elyaf", "content", "fabriccontent", "material"]):
        return "composition"
    if any(k in norm for k in ["en", "genislik", "width", "laize", "ebat"]):
        return "width"
    if any(k in norm for k in ["gramaj", "agirlik", "weight", "gm2", "gsm", "grm2", "gr"]):
        return "weight"
    if any(k in norm for k in ["barkod", "barcode", "qr", "qrcode"]):
        return "barcode_or_qr"
        
    return None

def analyze_excel_columns(filepath):
    """
    Reads an uploaded Excel/CSV file, returns header names, sample rows, total rows, and suggested field mappings.
    """
    ext = os.path.splitext(filepath)[1].lower()
    headers = []
    sample_rows = []
    total_rows = 0
    
    if ext in [".xlsx", ".xlsm", ".xltx"]:
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        ws = wb.active
        
        row_iter = ws.iter_rows(values_only=True)
        header_row = None
        
        for row in row_iter:
            # Skip completely empty rows or title note rows
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue
            cleaned = [str(c).strip() if c is not None else "" for c in row]
            # Detect header row (at least 2 non-empty columns)
            if sum(1 for c in cleaned if c) >= 2 and header_row is None:
                header_row = cleaned
                break
                
        if not header_row:
            return {"error": "Excel dosyasında geçerli bir başlık satırı bulunamadı."}
            
        headers = header_row
        
        # Read sample data rows
        for row in row_iter:
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue
            total_rows += 1
            if len(sample_rows) < 4:
                sample_rows.append([str(c).strip() if c is not None else "" for c in row])
                
        wb.close()
    elif ext == ".csv":
        with open(filepath, "r", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.reader(f)
            for row in reader:
                if not row or all(str(v).strip() == "" for v in row):
                    continue
                if not headers:
                    headers = [str(c).strip() for c in row]
                else:
                    total_rows += 1
                    if len(sample_rows) < 4:
                        sample_rows.append([str(c).strip() for c in row])
    else:
        return {"error": "Desteklenmeyen dosya formatı. Lütfen .xlsx veya .csv yükleyin."}

    # Generate suggested mappings
    mappings = {}
    used_fields = set()
    for col_name in headers:
        field = detect_column_field(col_name)
        if field and field not in used_fields:
            mappings[col_name] = field
            used_fields.add(field)
        else:
            mappings[col_name] = ""
            
    return {
        "headers": headers,
        "sample_rows": sample_rows,
        "total_rows": total_rows,
        "suggested_mappings": mappings
    }

def clean_excel_text(val):
    if val is None:
        return ""
    v = str(val).strip()
    return v.upper()

def import_excel_fabrics(filepath, column_mappings):
    """
    Imports fabrics from Excel/CSV using provided column_mappings dict (header -> field_name).
    Maintains existing internal codes, fills missing internal codes using the auto-counter,
    and updates database seamlessly.
    """
    ext = os.path.splitext(filepath)[1].lower()
    rows_to_process = []
    
    if ext in [".xlsx", ".xlsm", ".xltx"]:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active
        
        all_rows = list(ws.iter_rows(values_only=True))
        wb.close()
        
        header_idx = None
        for idx, row in enumerate(all_rows):
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue
            cleaned = [str(c).strip() if c is not None else "" for c in row]
            if sum(1 for c in cleaned if c) >= 2:
                header_idx = idx
                headers = cleaned
                break
                
        if header_idx is None:
            raise ValueError("Geçerli başlık satırı bulunamadı.")
            
        for row in all_rows[header_idx + 1:]:
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue
            row_dict = {}
            for col_idx, col_name in enumerate(headers):
                if col_idx < len(row):
                    row_dict[col_name] = row[col_idx]
            rows_to_process.append(row_dict)
            
    elif ext == ".csv":
        with open(filepath, "r", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.DictReader(f)
            for row in reader:
                if not any(row.values()):
                    continue
                rows_to_process.append(row)
    else:
        raise ValueError("Desteklenmeyen dosya formatı.")

    conn = get_db_connection()
    cursor = conn.cursor()
    code_settings = load_code_settings()
    
    imported_count = 0
    updated_count = 0
    errors = []
    
    # Process rows inside transaction
    try:
        for idx, row in enumerate(rows_to_process, 1):
            record = {
                "internal_code": "",
                "company_name": "",
                "quality_code": "",
                "quality_name": "",
                "design_code": "",
                "width": "",
                "weight": "",
                "composition": "",
                "barcode_or_qr": ""
            }
            
            for header, field_name in column_mappings.items():
                if field_name and field_name in record and header in row:
                    val = row[header]
                    if val is not None:
                        record[field_name] = str(val).strip()
            
            # Text cleaning
            company = clean_excel_text(record["company_name"]) or "GENEL"
            q_code = clean_excel_text(record["quality_code"]) or "KODSUZ"
            q_name = clean_excel_text(record["quality_name"])
            d_code = clean_excel_text(record["design_code"])
            width = clean_excel_text(record["width"])
            weight = clean_excel_text(record["weight"])
            comp = clean_excel_text(record["composition"])
            barcode = clean_excel_text(record["barcode_or_qr"])
            
            # Internal code determination
            int_code = str(record["internal_code"]).strip().upper() if record["internal_code"] else ""
            
            if not int_code:
                int_code = get_next_internal_code(conn, code_settings)
            
            if not barcode:
                barcode = int_code
                
            # Check if internal_code exists in DB
            cursor.execute("SELECT id FROM fabrics WHERE internal_code = ?", (int_code,))
            existing = cursor.fetchone()
            
            if existing:
                cursor.execute("""
                    UPDATE fabrics SET 
                        company_name = ?, quality_code = ?, quality_name = ?, 
                        design_code = ?, width = ?, weight = ?, composition = ?, barcode_or_qr = ?
                    WHERE id = ?
                """, (company, q_code, q_name, d_code, width, weight, comp, barcode, existing["id"]))
                updated_count += 1
            else:
                cursor.execute("""
                    INSERT INTO fabrics (
                        internal_code, company_name, quality_code, quality_name,
                        design_code, width, weight, composition, barcode_or_qr
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (int_code, company, q_code, q_name, d_code, width, weight, comp, barcode))
                imported_count += 1
                
        conn.commit()
    except Exception as e:
        conn.rollback()
        raise e
    finally:
        conn.close()
        
    return {
        "success": True,
        "imported_count": imported_count,
        "updated_count": updated_count,
        "total_processed": len(rows_to_process)
    }

