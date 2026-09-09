import sqlite3
import openpyxl
from database import DB_PATH, init_db

init_db()

conn = sqlite3.connect(DB_PATH)
c = conn.cursor()

# Clear existing
c.execute('DELETE FROM size_distributions')
c.execute('DELETE FROM production_status_logs')
c.execute('DELETE FROM audit_logs')
c.execute('DELETE FROM fabric_links')
c.execute('DELETE FROM styles')
c.execute('DELETE FROM orders')
conn.commit()

# 1. PARSE & SEED ROBERTO SARTO ÇARŞAF EXCEL
wb = openpyxl.load_workbook(r'C:\Users\ASLI CELIK\Desktop\ROBERTO SARTO FALL26 çarşaf - Kopya.XLSX', data_only=True)
ws = wb.active

carsaf_rows = []
for r in range(2, ws.max_row + 1):
    po = ws.cell(r, 1).value
    style = ws.cell(r, 3).value
    if not po or not style:
        continue
    
    color = str(ws.cell(r, 4).value or '')
    fabric = str(ws.cell(r, 5).value or '')
    pcs = int(ws.cell(r, 6).value or ws.cell(r, 21).value or 0)
    price = float(ws.cell(r, 7).value or ws.cell(r, 22).value or 0.0)
    fabric_status = str(ws.cell(r, 8).value or '')
    sms_meters = str(ws.cell(r, 9).value or '')
    pps_meters = str(ws.cell(r, 10).value or '')
    fabric_ordered = 1 if ws.cell(r, 11).value else 0
    pps_sent = 1 if ws.cell(r, 12).value else 0
    shipping_sent = 1 if ws.cell(r, 13).value else 0
    fabric_type = str(ws.cell(r, 14).value or '')
    
    sizes = {
        '38': int(ws.cell(r, 15).value or 0),
        '40': int(ws.cell(r, 16).value or 0),
        '42': int(ws.cell(r, 17).value or 0),
        '44': int(ws.cell(r, 18).value or 0),
        '46': int(ws.cell(r, 19).value or 0),
        '48': int(ws.cell(r, 20).value or 0)
    }
    desc = str(ws.cell(r, 31).value or '')
    comp = str(ws.cell(r, 32).value or '')
    rec_meters = ws.cell(r, 33).value
    arr_date = str(ws.cell(r, 34).value or '')
    
    aks1 = str(ws.cell(r, 35).value or '')
    aks2 = str(ws.cell(r, 36).value or '')
    aks3 = str(ws.cell(r, 37).value or '')
    aks4 = str(ws.cell(r, 38).value or '')
    
    logs = []
    for log_col in range(39, 49):
        lv = ws.cell(r, log_col).value
        if lv:
            header_title = str(ws.cell(1, log_col).value or 'Durum')
            logs.append({'title': header_title, 'message': str(lv)})

    carsaf_rows.append({
        'po_number': str(po).strip(),
        'style_no': str(style).strip(),
        'color_name': color,
        'fabric_article': fabric,
        'total_quantity': pcs,
        'unit_price': price,
        'fabric_order_status': fabric_status,
        'sms_unit_meters': sms_meters,
        'pps_unit_meters': pps_meters,
        'fabric_ordered': fabric_ordered,
        'pps_sent': pps_sent,
        'shipping_sample_sent': shipping_sent,
        'fabric_type': fabric_type,
        'sizes': sizes,
        'description': desc,
        'fabric_composition': comp,
        'fabric_received_meters': rec_meters,
        'fabric_arrival_date': arr_date,
        'accessory_1': aks1,
        'accessory_2': aks2,
        'accessory_3': aks3,
        'accessory_4': aks4,
        'logs': logs
    })

# Group Roberto Sarto POs
po_groups = {}
for item in carsaf_rows:
    po_groups.setdefault(item['po_number'], []).append(item)

for po_num, styles_list in po_groups.items():
    tot_q = sum(st['total_quantity'] for st in styles_list)
    tot_amt = sum(st['total_quantity'] * st['unit_price'] for st in styles_list)
    
    c.execute('''
    INSERT INTO orders (company_id, po_number, customer_name, brand, season, total_quantity, total_amount, currency, raw_file_name)
    VALUES (1, ?, 'Sarto Fashion BV', 'Roberto Sarto', 'Fall 26', ?, ?, 'EUR', 'ROBERTO SARTO FALL26 carsaf.xlsx')
    ''', (po_num, tot_q, tot_amt))
    order_id = c.lastrowid
    
    for st in styles_list:
        c.execute('''
        INSERT INTO styles (
            order_id, company_id, style_no, description, fabric_composition, fabric_article,
            fabric_type, color_name, unit_price, total_quantity, total_amount, currency,
            sms_unit_meters, pps_unit_meters, fabric_order_status, fabric_ordered,
            pps_sent, shipping_sample_sent, fabric_received_meters, fabric_arrival_date,
            accessory_1, accessory_2, accessory_3, accessory_4, status
        )
        VALUES (?, 1, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'EUR', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Üretimde')
        ''', (
            order_id, st['style_no'], st['description'], st['fabric_composition'], st['fabric_article'],
            st['fabric_type'], st['color_name'], st['unit_price'], st['total_quantity'],
            st['total_quantity'] * st['unit_price'], st['sms_unit_meters'], st['pps_unit_meters'],
            st['fabric_order_status'], st['fabric_ordered'], st['pps_sent'], st['shipping_sample_sent'],
            st['fabric_received_meters'], st['fabric_arrival_date'], st['accessory_1'], st['accessory_2'],
            st['accessory_3'], st['accessory_4']
        ))
        style_id = c.lastrowid
        
        for idx, (sz_name, sz_qty) in enumerate(st['sizes'].items()):
            c.execute('''
            INSERT INTO size_distributions (style_id, size_name, quantity, sort_order)
            VALUES (?, ?, ?, ?)
            ''', (style_id, sz_name, sz_qty, idx + 1))
            
        for lg in st['logs']:
            c.execute('''
            INSERT INTO production_status_logs (style_id, user_name, category, title, message)
            VALUES (?, 'Excel Aktarımı', 'Müşteri Kritiği', ?, ?)
            ''', (style_id, lg['title'], lg['message']))

# 2. SEED ANOTHERWOMAN ORDERS (FROM 11-PAGE PDF)
anotherwoman_orders = [
    {'po': 'PO26-0214', 'style': '632099', 'desc': 'T-shirt turtle print', 'color': 'New ornamental Teal (T2991)', 'comp': '95% CV 5% EL', 'price': 16.97, 'sizes': {'36': 14, '38': 23, '40': 24, '42': 24, '44': 21, '46': 13}, 'del_date': '2026-07-15'},
    {'po': 'PO26-0215', 'style': '632103', 'desc': 'Blouse scarve', 'color': 'Dotprint Berry-Stone (T2030)', 'comp': '100% CV', 'price': 25.40, 'sizes': {'36': 14, '38': 23, '40': 25, '42': 25, '44': 23, '46': 13}, 'del_date': '2026-07-15'},
    {'po': 'PO26-0216', 'style': '632104', 'desc': 'Blouse standup with ruffles', 'color': 'Smoke blue (T729)', 'comp': '83% CV 17% PA', 'price': 21.96, 'sizes': {'36': 11, '38': 18, '40': 19, '42': 23, '44': 17, '46': 12}, 'del_date': '2026-07-15'},
    {'po': 'PO26-0217', 'style': '632105', 'desc': 'Blouse ruffles', 'color': 'Check XL Taupe-Berry (T2050)', 'comp': '100% CV', 'price': 20.66, 'sizes': {'36': 10, '38': 15, '40': 24, '42': 22, '44': 17, '46': 12}, 'del_date': '2026-07-15'},
    {'po': 'PO26-0218', 'style': '632114', 'desc': 'Skirt check', 'color': 'Check XL Taupe-Berry (T2050)', 'comp': '100% CV', 'price': 23.45, 'sizes': {'36': 9, '38': 19, '40': 22, '42': 22, '44': 17, '46': 11}, 'del_date': '2026-07-15'},
    {'po': 'PO26-0219', 'style': '632118', 'desc': 'Jacket check jacquard', 'color': 'Tweedcheck Taupe (T2118)', 'comp': '60% PAN 30% PES 10% CO', 'price': 25.80, 'sizes': {'36': 17, '38': 27, '40': 28, '42': 31, '44': 18, '46': 9}, 'del_date': '2026-07-15'},
    {'po': 'PO26-0220', 'style': '632121', 'desc': 'T-shirt fancy collar', 'color': 'Check XL Taupe-Berry (T2050)', 'comp': '96% CV 4% EL', 'price': 13.85, 'sizes': {'36': 9, '38': 19, '40': 21, '42': 22, '44': 18, '46': 11}, 'del_date': '2026-07-15'},
    {'po': 'PO26-0221', 'style': '632304', 'desc': 'Skirt print shimmer', 'color': 'Oriental flower Midnight (T2304)', 'comp': '97% CV 3% PES', 'price': 22.60, 'sizes': {'36': 12, '38': 18, '40': 20, '42': 20, '44': 18, '46': 12}, 'del_date': '2026-07-15'},
    {'po': 'PO26-0222', 'style': '632341', 'desc': 'Jacket', 'color': 'Ornamental jacq Midnight (T2345)', 'comp': '73% PES 24% CO 3% ME', 'price': 21.78, 'sizes': {'36': 12, '38': 20, '40': 21, '42': 25, '44': 15, '46': 7}, 'del_date': '2026-07-15'},
    {'po': 'PO26-0223', 'style': '632344', 'desc': 'Blouse print', 'color': 'Oriental flower Midnight (T2304)', 'comp': '97% CV 3% PES', 'price': 21.24, 'sizes': {'36': 10, '38': 17, '40': 23, '42': 24, '44': 18, '46': 8}, 'del_date': '2026-07-15'},
    {'po': 'PO26-0224', 'style': '632359', 'desc': 'Blouse standup collar', 'color': 'New ornamental Teal (T2359)', 'comp': '65% CV 35% PA', 'price': 24.78, 'sizes': {'36': 10, '38': 20, '40': 22, '42': 22, '44': 17, '46': 9}, 'del_date': '2026-07-15'}
]

for aw in anotherwoman_orders:
    tot_q = sum(aw['sizes'].values())
    tot_amt = tot_q * aw['price']
    c.execute('''
    INSERT INTO orders (company_id, po_number, customer_name, brand, season, delivery_date, total_quantity, total_amount, currency, raw_file_name)
    VALUES (1, ?, 'Sarto Fashion BV', 'Anotherwoman', 'Fall 26', ?, ?, ?, 'EUR', 'Sarto_Fashion_Anotherwoman.pdf')
    ''', (aw['po'], aw['del_date'], tot_q, tot_amt))
    order_id = c.lastrowid
    
    c.execute('''
    INSERT INTO styles (order_id, company_id, style_no, description, fabric_composition, color_name, unit_price, total_quantity, total_amount, currency, status)
    VALUES (?, 1, ?, ?, ?, ?, ?, ?, ?, 'EUR', 'Onaylandı')
    ''', (order_id, aw['style'], aw['desc'], aw['comp'], aw['color'], aw['price'], tot_q, tot_amt))
    style_id = c.lastrowid
    
    for idx, (sz, q) in enumerate(aw['sizes'].items()):
        c.execute('''
        INSERT INTO size_distributions (style_id, size_name, quantity, sort_order)
        VALUES (?, ?, ?, ?)
        ''', (style_id, sz, q, idx + 1))

# 3. SEED POOOLS ORDER (PO26-0246)
poools_sizes = {'34': 11, '36': 39, '38': 43, '40': 48, '42': 44, '44': 10, '46': 35}
poools_tot_q = sum(poools_sizes.values())
poools_tot_amt = poools_tot_q * 21.94

c.execute('''
INSERT INTO orders (company_id, po_number, customer_name, brand, season, delivery_date, total_quantity, total_amount, currency, raw_file_name)
VALUES (1, 'PO26-0246', 'Sarto Fashion BV', 'Poools', 'Fall 26', '2026-09-04', ?, ?, 'EUR', 'Sarto_Fashion_Poools.pdf')
''', (poools_tot_q, poools_tot_amt))
order_id = c.lastrowid

c.execute('''
INSERT INTO styles (order_id, company_id, style_no, description, fabric_composition, color_name, color_code, unit_price, total_quantity, total_amount, currency, status)
VALUES (?, 1, '643124', 'Blouse', '35% RA 30% ST 27% CV 8% PES', 'Rouge brun 35', 'U299', 21.94, ?, ?, 'EUR', 'Onaylandı')
''', (order_id, poools_tot_q, poools_tot_amt))
style_id = c.lastrowid

for idx, (sz, q) in enumerate(poools_sizes.items()):
    c.execute('''
    INSERT INTO size_distributions (style_id, size_name, quantity, sort_order)
    VALUES (?, ?, ?, ?)
    ''', (style_id, sz, q, idx + 1))

# 4. SEED TRICOT COMPANY ORDER (IVY & ONESTO)
c.execute('''
INSERT INTO orders (company_id, po_number, customer_name, brand, season, delivery_date, total_quantity, total_amount, currency, raw_file_name)
VALUES (1, 'TCEL 004', 'Tricot Company B.V.', 'IVY', 'Fall 2026', '2026-08-30', 60, 840.0, 'EUR', 'Tricot_Company_TCEL004.xlsx')
''')
order_id_ivy = c.lastrowid

c.execute('''
INSERT INTO styles (order_id, company_id, style_no, description, fabric_composition, color_name, unit_price, total_quantity, total_amount, currency, status)
VALUES (?, 1, '26153', 'Blouse oversized boxy fit', '82% viscose 18% nylon', 'AOP Floral Blow up (IVY)', 14.0, 60, 840.0, 'EUR', 'Üretimde')
''', (order_id_ivy,))
st_id_ivy = c.lastrowid

for idx, (sz, q) in enumerate({'S': 10, 'M': 20, 'L': 20, 'XL': 10, 'XXL': 0}.items()):
    c.execute('INSERT INTO size_distributions (style_id, label_brand, size_name, quantity, sort_order) VALUES (?, ?, ?, ?, ?)', (st_id_ivy, 'IVY', sz, q, idx + 1))

c.execute('''
INSERT INTO orders (company_id, po_number, customer_name, brand, season, delivery_date, total_quantity, total_amount, currency, raw_file_name)
VALUES (1, 'TCEL 004', 'Tricot Company B.V.', 'ONESTO', 'Fall 2026', '2026-08-30', 977, 13678.0, 'EUR', 'Tricot_Company_TCEL004.xlsx')
''')
order_id_onesto = c.lastrowid

c.execute('''
INSERT INTO styles (order_id, company_id, style_no, description, fabric_composition, color_name, unit_price, total_quantity, total_amount, currency, status)
VALUES (?, 1, '26153', 'Blouse oversized boxy fit', '82% viscose 18% nylon', 'AOP Floral Blow up (ONESTO)', 14.0, 977, 13678.0, 'EUR', 'Üretimde')
''', (order_id_onesto,))
st_id_onesto = c.lastrowid

for idx, (sz, q) in enumerate({'S': 150, 'M': 320, 'L': 310, 'XL': 147, 'XXL': 50}.items()):
    c.execute('INSERT INTO size_distributions (style_id, label_brand, size_name, quantity, sort_order) VALUES (?, ?, ?, ?, ?)', (st_id_onesto, 'ONESTO', sz, q, idx + 1))

conn.commit()

print('=== UPDATED DATABASE SUMMARY ===')
print('Orders count:', c.execute('SELECT count(*) FROM orders').fetchone()[0])
print('Styles count:', c.execute('SELECT count(*) FROM styles').fetchone()[0])
print('Sizes count:', c.execute('SELECT count(*) FROM size_distributions').fetchone()[0])

print('\nDistinct Customer & Brands in DB:')
for r in c.execute('SELECT DISTINCT customer_name, brand FROM orders').fetchall():
    print(f'Customer: {r[0]:<22} | Brand: {r[1]}')

conn.close()
