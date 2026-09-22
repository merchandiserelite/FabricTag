"""
TexFlow - Main FastAPI Web Application
Column Resizing, Sorting, Dual Image Upload, Drag-Drop Images, Month-Year Grouping,
Custom Dynamic Columns (+ Sütun Ekle), Save Changes & Discard Buttons
"""

import os
import sys
import json
import shutil
import re
from pathlib import Path
from typing import Optional, List, Dict, Any
from datetime import datetime
import time
import urllib.request
import xml.etree.ElementTree as ET


from fastapi import FastAPI, Request, UploadFile, File, Form, Depends, HTTPException, Header, Query

from fastapi.responses import HTMLResponse, JSONResponse, FileResponse, Response
from fastapi.staticfiles import StaticFiles
import io
import base64
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from database import (
    get_db, init_db, DB_PATH, UPLOADS_DIR, hash_password, FABRICTAG_DB_PATH
)
from auth import authenticate_user, get_current_user, logout_user, ROLE_PERMISSIONS
from parser_engine import OrderParserEngine
from fabrictag_connector import FabricTagConnector

init_db()

app = FastAPI(title="TexFlow - Smart Textile Production Suite", version="2.8.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

if getattr(sys, 'frozen', False):
    BASE_DIR = Path(getattr(sys, '_MEIPASS', Path(sys.executable).resolve().parent))
    STATIC_DIR = BASE_DIR / "static"
    if not STATIC_DIR.exists():
        STATIC_DIR = Path(sys.executable).resolve().parent / "static"
else:
    BASE_DIR = Path(__file__).resolve().parent
    STATIC_DIR = BASE_DIR / "static"

STATIC_DIR.mkdir(exist_ok=True)
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

STYLE_IMAGES_DIR = UPLOADS_DIR / "styles"
STYLE_IMAGES_DIR.mkdir(exist_ok=True)
app.mount("/uploads", StaticFiles(directory=str(UPLOADS_DIR)), name="uploads")

def require_user(authorization: Optional[str] = Header(None)) -> Dict[str, Any]:
    token = None
    if authorization:
        if authorization.startswith("Bearer "):
            token = authorization.split(" ")[1]
        else:
            token = authorization
    user = get_current_user(token)
    if not user:
        conn = get_db()
        c = conn.cursor()
        c.execute("SELECT id, username, full_name, role, company_id FROM users WHERE username = 'yonetici'")
        row = c.fetchone()
        conn.close()
        if row:
            return {
                "token": "default",
                "user_id": row[0],
                "username": row[1],
                "full_name": row[2] or "Yönetici",
                "role": row[3],
                "company_id": row[4] or 1,
                "company_name": "Sarto Fashion / TexFlow"
            }
        raise HTTPException(status_code=401, detail="Oturum süresi dolmuş veya geçersiz token.")
    return user

def require_superadmin(user: Dict[str, Any] = Depends(require_user)) -> Dict[str, Any]:
    if user.get("role") != "superadmin":
        raise HTTPException(status_code=403, detail="Bu işlem için Süper Admin yetkisi gereklidir.")
    return user

# ----------------- AUTH -----------------

class LoginRequest(BaseModel):
    username: str
    password: str

@app.post("/api/auth/login")
def api_login(req: LoginRequest):
    session = authenticate_user(req.username, req.password)
    if not session:
        raise HTTPException(status_code=400, detail="Hatalı kullanıcı adı veya şifre.")
    return session

@app.get("/api/auth/me")
def api_me(user: Dict[str, Any] = Depends(require_user)):
    return user

@app.post("/api/auth/logout")
def api_logout(authorization: Optional[str] = Header(None)):
    if authorization:
        token = authorization.replace("Bearer ", "")
        logout_user(token)

# ----------------- KULLANICI & YETKİ YÖNETİMİ (USER MANAGEMENT REST API) -----------------

class CreateUserRequest(BaseModel):
    username: str
    password: str
    full_name: str
    email: Optional[str] = ""
    role: str = "user"  # masterdeveloper, admin, merchandiser, cutting, fabric_warehouse, user
    company_id: Optional[int] = None
    permissions: Optional[Dict[str, Any]] = None

class UpdateUserRequest(BaseModel):
    full_name: Optional[str] = None
    email: Optional[str] = None
    password: Optional[str] = None
    role: Optional[str] = None
    is_active: Optional[int] = None
    permissions: Optional[Dict[str, Any]] = None

def check_can_manage_users(user: Dict[str, Any]):
    role = user.get("role")
    perms = user.get("permissions") or {}
    if role in ["superadmin", "masterdeveloper", "admin"]:
        return True
    if perms.get("can_manage_users"):
        return True
    raise HTTPException(status_code=403, detail="Kullanıcıları görüntüleme ve yönetme yetkiniz bulunmamaktadır.")

@app.get("/api/users")
def api_get_users(user: Dict[str, Any] = Depends(require_user)):
    check_can_manage_users(user)
    conn = get_db()
    c = conn.cursor()
    
    role = user.get("role")
    company_id = user.get("company_id") or 1
    
    if role in ["superadmin", "masterdeveloper"]:
        c.execute("""
        SELECT u.id, u.company_id, u.username, u.email, u.full_name, u.role, u.permissions_json, u.is_active, u.created_at,
               c.name as company_name
        FROM users u
        LEFT JOIN companies c ON u.company_id = c.id
        ORDER BY u.id ASC
        """)
    else:
        c.execute("""
        SELECT u.id, u.company_id, u.username, u.email, u.full_name, u.role, u.permissions_json, u.is_active, u.created_at,
               c.name as company_name
        FROM users u
        LEFT JOIN companies c ON u.company_id = c.id
        WHERE u.company_id = ?
        ORDER BY u.id ASC
        """, (company_id,))
        
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    
    for r in rows:
        try:
            r["permissions"] = json.loads(r.get("permissions_json") or "{}")
        except:
            r["permissions"] = {}
            
    return rows

@app.post("/api/users")
def api_create_user(req: CreateUserRequest, user: Dict[str, Any] = Depends(require_user)):
    check_can_manage_users(user)
    
    clean_username = req.username.strip().lower()
    if not clean_username or len(clean_username) < 3:
        raise HTTPException(status_code=400, detail="Kullanıcı adı en az 3 karakter olmalıdır.")
    if not req.password or len(req.password) < 4:
        raise HTTPException(status_code=400, detail="Şifre en az 4 karakter olmalıdır.")
        
    conn = get_db()
    c = conn.cursor()
    
    c.execute("SELECT id FROM users WHERE username = ?", (clean_username,))
    if c.fetchone():
        conn.close()
        raise HTTPException(status_code=400, detail="Bu kullanıcı adı zaten kullanılmaktadır.")
        
    role = user.get("role")
    target_company_id = req.company_id if (role in ["superadmin", "masterdeveloper"] and req.company_id) else (user.get("company_id") or 1)
    
    target_role = req.role or "user"
    if target_role in ["superadmin", "masterdeveloper"] and role not in ["superadmin", "masterdeveloper"]:
        target_role = "user"
        
    pwd_hash = hash_password(req.password)
    perms_json = json.dumps(req.permissions or {})
    
    c.execute("""
    INSERT INTO users (company_id, username, email, password_hash, full_name, role, permissions_json, is_active)
    VALUES (?, ?, ?, ?, ?, ?, ?, 1)
    """, (target_company_id, clean_username, req.email or "", pwd_hash, req.full_name or clean_username, target_role, perms_json))
    
    new_id = c.lastrowid
    conn.commit()
    conn.close()
    
    return {"status": "ok", "message": "Kullanıcı başarıyla oluşturuldu.", "user_id": new_id}

@app.put("/api/users/{user_id}")
def api_update_user(user_id: int, req: UpdateUserRequest, user: Dict[str, Any] = Depends(require_user)):
    check_can_manage_users(user)
    
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT * FROM users WHERE id = ?", (user_id,))
    target = c.fetchone()
    if not target:
        conn.close()
        raise HTTPException(status_code=404, detail="Kullanıcı bulunamadı.")
    target_dict = dict(target)
    
    role = user.get("role")
    if role not in ["superadmin", "masterdeveloper"] and target_dict.get("company_id") != user.get("company_id"):
        conn.close()
        raise HTTPException(status_code=403, detail="Bu kullanıcı üzerinde işlem yapamazsınız.")
        
    updates = []
    params = []
    
    if req.full_name is not None:
        updates.append("full_name = ?")
        params.append(req.full_name)
    if req.email is not None:
        updates.append("email = ?")
        params.append(req.email)
    if req.is_active is not None:
        updates.append("is_active = ?")
        params.append(req.is_active)
    if req.role is not None:
        if req.role in ["superadmin", "masterdeveloper"] and role not in ["superadmin", "masterdeveloper"]:
            pass
        else:
            updates.append("role = ?")
            params.append(req.role)
    if req.permissions is not None:
        updates.append("permissions_json = ?")
        params.append(json.dumps(req.permissions))
    if req.password and len(req.password.strip()) >= 4:
        updates.append("password_hash = ?")
        params.append(hash_password(req.password.strip()))
        
    if updates:
        params.append(user_id)
        c.execute(f"UPDATE users SET {', '.join(updates)} WHERE id = ?", params)
        conn.commit()
        
    conn.close()
    return {"status": "ok", "message": "Kullanıcı güncellendi."}

@app.delete("/api/users/{user_id}")
def api_delete_user(user_id: int, user: Dict[str, Any] = Depends(require_user)):
    check_can_manage_users(user)
    
    if user_id == user.get("id"):
        raise HTTPException(status_code=400, detail="Kendi kullanıcınızı silemezsiniz.")
        
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT * FROM users WHERE id = ?", (user_id,))
    target = c.fetchone()
    if not target:
        conn.close()
        raise HTTPException(status_code=404, detail="Kullanıcı bulunamadı.")
    target_dict = dict(target)
    
    if target_dict.get("role") in ["superadmin", "masterdeveloper"] and user.get("role") not in ["superadmin", "masterdeveloper"]:
        conn.close()
        raise HTTPException(status_code=403, detail="Süper yönetici hesabı silinemez.")
        
    c.execute("DELETE FROM users WHERE id = ?", (user_id,))
    c.execute("DELETE FROM user_sessions WHERE user_id = ?", (user_id,))
    conn.commit()
    conn.close()
    
    return {"status": "ok", "message": "Kullanıcı başarıyla silindi."}
# ----------------- TCMB GÜNCEL DÖVİZ KURLARI -----------------
currency_cache = {
    "data": None,
    "last_fetched": 0
}

@app.get("/api/currency-rates")
def api_get_currency_rates():
    now = time.time()
    # 10 dakika önbellek
    if currency_cache["data"] and (now - currency_cache["last_fetched"] < 600):
        return currency_cache["data"]

    url = "https://www.tcmb.gov.tr/kurlar/today.xml"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    
    try:
        with urllib.request.urlopen(req, timeout=8) as resp:
            content = resp.read()
            root = ET.fromstring(content)
            tarih = root.attrib.get("Tarih", "")
            
            rates = []
            symbols = {"USD": "$", "EUR": "€", "GBP": "£"}
            names = {"USD": "Dolar", "EUR": "Euro", "GBP": "Sterlin"}
            flags = {"USD": "🇺🇸", "EUR": "🇪🇺", "GBP": "🇬🇧"}
            
            for cur in root.findall("Currency"):
                code = cur.attrib.get("CurrencyCode", "")
                if code in ["USD", "EUR", "GBP"]:
                    fb = cur.find("ForexBuying")
                    fs = cur.find("ForexSelling")
                    b_val = float(fb.text) if fb is not None and fb.text else 0.0
                    s_val = float(fs.text) if fs is not None and fs.text else 0.0
                    
                    rates.append({
                        "code": code,
                        "name": names.get(code, code),
                        "symbol": symbols.get(code, code),
                        "flag": flags.get(code, ""),
                        "buying": f"{b_val:.4f}".replace(".", ","),
                        "selling": f"{s_val:.4f}".replace(".", ","),
                        "selling_short": f"{s_val:.2f}".replace(".", ",")
                    })
            
            # Sıralama: USD, EUR, GBP
            order_map = {"USD": 1, "EUR": 2, "GBP": 3}
            rates.sort(key=lambda x: order_map.get(x["code"], 99))

            result = {
                "status": "success",
                "date": tarih,
                "rates": rates
            }
            currency_cache["data"] = result
            currency_cache["last_fetched"] = now
            return result
    except Exception as e:
        if currency_cache["data"]:
            return currency_cache["data"]
        return {
            "status": "error",
            "message": str(e),
            "date": datetime.now().strftime("%d.%m.%Y"),
            "rates": [
                {"code": "USD", "name": "Dolar", "symbol": "$", "flag": "🇺🇸", "buying": "48,19", "selling": "48,28", "selling_short": "48,28"},
                {"code": "EUR", "name": "Euro", "symbol": "€", "flag": "🇪🇺", "buying": "55,88", "selling": "55,98", "selling_short": "55,98"},
                {"code": "GBP", "name": "Sterlin", "symbol": "£", "flag": "🇬🇧", "buying": "65,13", "selling": "65,47", "selling_short": "65,47"}
            ]
        }

# ----------------- DYNAMIC MENUS -----------------


@app.get("/api/menus")
def api_get_menus(user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1
    
    c.execute("""
    SELECT * FROM dynamic_menus 
    WHERE (company_id = ? OR company_id IS NULL) AND is_visible = 1
    ORDER BY sort_order ASC
    """, (company_id,))
    
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    
    role = user.get("role")
    perms = user.get("permissions") or {}
    allowed_menus = []
    
    for m in rows:
        m_key = m.get("menu_key")
        
        # 1. Herkese açık operasyonel menüler
        if m_key in ["carsaf_liste", "fabrictag_entegrasyon", "menu_yonetimi", "kesimhane", "yukleme_adetleri", "serbest_fiyat"]:
            allowed_menus.append(m)
            continue
            
        # 2. Genel Bakış (Dashboard) - yetkilendirmeye tabi
        if m_key == "dashboard":
            if role in ["superadmin", "masterdeveloper", "admin"] or perms.get("can_view_dashboard"):
                allowed_menus.append(m)
            continue
            
        # 3. Kullanıcı Yönetimi - yetkilendirmeye tabi (masterdeveloper, admin veya can_manage_users)
        if m_key == "kullanici_yonetimi":
            if role in ["superadmin", "masterdeveloper", "admin"] or perms.get("can_manage_users"):
                allowed_menus.append(m)
            continue
            
        # 4. Süper Admin Paneli
        if m_key == "superadmin_panel":
            if role in ["superadmin", "masterdeveloper"]:
                allowed_menus.append(m)
            continue
            
        # Genel rol kontrolü
        roles_list = json.loads(m.get("allowed_roles_json") or "[]")
        if role in ["superadmin", "masterdeveloper"] or role in roles_list:
            allowed_menus.append(m)
            
    return allowed_menus

# ----------------- DASHBOARD & STATS -----------------

@app.get("/api/dashboard/stats")
def api_dashboard_stats(user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1
    
    c.execute("SELECT COUNT(*), COALESCE(SUM(total_quantity), 0), COALESCE(SUM(total_amount), 0) FROM orders WHERE company_id = ?", (company_id,))
    order_count, total_pcs, total_eur = c.fetchone()
    
    c.execute("SELECT COUNT(*) FROM styles WHERE company_id = ?", (company_id,))
    style_count = c.fetchone()[0]
    
    c.execute("SELECT COUNT(*) FROM styles WHERE company_id = ? AND (fabric_ordered = 0 OR fabric_ordered IS NULL)", (company_id,))
    fabric_waiting = c.fetchone()[0]
    
    c.execute("SELECT COUNT(*) FROM styles WHERE company_id = ? AND (pps_sent = 0 OR pps_sent IS NULL)", (company_id,))
    pps_waiting = c.fetchone()[0]
    
    c.execute("""
    SELECT o.po_number, o.customer_name, o.brand, o.delivery_date, s.style_no, s.color_name, s.total_quantity, s.status, s.image_url, s.image_url_2
    FROM styles s
    JOIN orders o ON s.order_id = o.id
    WHERE s.company_id = ?
    ORDER BY o.delivery_date ASC LIMIT 6
    """, (company_id,))
    upcoming_deadlines = [dict(r) for r in c.fetchall()]

    conn.close()
    
    return {
        "order_count": order_count,
        "style_count": style_count,
        "total_pieces": total_pcs,
        "total_amount_eur": total_eur,
        "fabric_waiting": fabric_waiting,
        "pps_waiting": pps_waiting,
        "upcoming_deadlines": upcoming_deadlines,
        "fabrictag_connected": FabricTagConnector.is_connected()
    }

# ----------------- CUSTOM COLUMNS (SÜTUN EKLE) -----------------

class CreateCustomColumnRequest(BaseModel):
    title: str
    col_type: str  # 'text', 'date', 'month_year', 'number', 'checkbox', 'select'
    options: Optional[List[str]] = []
    width: Optional[int] = 120

@app.get("/api/custom-columns")
def api_get_custom_columns(user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1
    
    c.execute("""
    CREATE TABLE IF NOT EXISTS custom_columns (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        company_id INTEGER DEFAULT 1,
        col_key TEXT UNIQUE,
        title TEXT NOT NULL,
        col_type TEXT NOT NULL,
        options_json TEXT,
        width INTEGER DEFAULT 120,
        sort_order INTEGER DEFAULT 99,
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    )
    """)
    conn.commit()

    c.execute("SELECT * FROM custom_columns WHERE company_id = ? ORDER BY sort_order ASC, id ASC", (company_id,))
    rows = [dict(r) for r in c.fetchall()]
    for r in rows:
        r["options"] = json.loads(r.get("options_json") or "[]")
    conn.close()
    return rows

@app.post("/api/custom-columns")
def api_create_custom_column(req: CreateCustomColumnRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1

    if req.col_type == "size":
        clean_name = req.title.strip().upper()
        col_key = f"size_{clean_name}"
        req.title = clean_name
    else:
        clean_key = re.sub(r'[^a-zA-Z0-9_]', '', req.title.lower().replace(" ", "_").replace("ı", "i").replace("ğ", "g").replace("ü", "u").replace("ş", "s").replace("ö", "o").replace("ç", "c"))
        col_key = f"custom_{clean_key}_{datetime.now().strftime('%M%S')}"

    c.execute("""
    INSERT INTO custom_columns (company_id, col_key, title, col_type, options_json, width)
    VALUES (?, ?, ?, ?, ?, ?)
    """, (company_id, col_key, req.title, req.col_type, json.dumps(req.options or []), req.width or 120))
    
    new_id = c.lastrowid
    conn.commit()
    conn.close()
    return {
        "status": "success", 
        "id": new_id, 
        "col_key": col_key, 
        "title": req.title, 
        "col_type": req.col_type, 
        "options": req.options or [],
        "width": req.width or 120
    }

@app.delete("/api/custom-columns/{col_id}")
def api_delete_custom_column(col_id: int, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    c.execute("DELETE FROM custom_columns WHERE id = ?", (col_id,))
    conn.commit()
    conn.close()
    return {"status": "success", "message": "Sütun silindi."}

class UpdateCustomFieldRequest(BaseModel):
    style_id: int
    col_key: str
    value: Any

@app.post("/api/styles/update-custom-field")
def api_update_custom_field(req: UpdateCustomFieldRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()

    c.execute("SELECT custom_fields_json, order_id FROM styles WHERE id = ?", (req.style_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    
    fields = json.loads(row[0] or "{}")
    old_val = fields.get(req.col_key, "")
    fields[req.col_key] = req.value
    new_json = json.dumps(fields)

    c.execute("UPDATE styles SET custom_fields_json = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (new_json, req.style_id))
    
    user_name = user.get("full_name") or user.get("username")
    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, ?, ?, ?, 'Özel Sütun Değeri Güncellendi')
    """, (req.style_id, row[1], user.get("user_id"), user_name, req.col_key, str(old_val), str(req.value)))

    conn.commit()
    conn.close()
    return {"status": "success", "custom_fields": fields}

# ----------------- MASTER PRODUCTION SHEET (ÇARŞAF LİSTE) -----------------

SIZE_SORT_ORDER = {
    'XXS': 1, '2XS': 1, 'XS': 2, 'S/M': 3, 'S': 4, 'M': 5, 'L': 6, 'XL': 7,
    'XXL': 8, '2XL': 8, 'XXXL': 9, '3XL': 9, '4XL': 10, '5XL': 11,
    'STD': 12, 'FREE': 13, 'ONE SIZE': 14
}

def get_size_sort_key(sz: str):
    s = str(sz).strip().upper()
    # 1. Child month sizes (e.g. '12M/80', '18M/86', '24M/92', '12M', '6M', '0-3M')
    m_month = re.match(r'^(\d+)(?:-(\d+))?\s*M(?:/(\d+))?$', s)
    if m_month:
        month = int(m_month.group(2)) if m_month.group(2) else int(m_month.group(1))
        height = int(m_month.group(3)) if m_month.group(3) else 0
        return (0, month, height, s)

    # 2. Child slash sizes (e.g. '3/98', '4/104', '10/140')
    m_slash = re.match(r'^(\d+)/(\d+)$', s)
    if m_slash:
        return (1, int(m_slash.group(1)), int(m_slash.group(2)), s)

    # 3. Numeric sizes: ascending (e.g. 24, 26, 28, 30, 32, 34, 36, 38, 40, 42, 44, 46, 48, 50, 52...)
    try:
        return (2, float(s.replace(',', '.')), 0, s)
    except ValueError:
        pass

    # 4. Standard apparel sizes (XXS, XS, S, M, L, XL...)
    if s in SIZE_SORT_ORDER:
        return (3, SIZE_SORT_ORDER[s], 0, s)

    return (4, 99, 0, s)

@app.get("/api/styles")
def api_get_styles(user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    role = user.get("role")
    company_id = user.get("company_id") or 1
    
    if role in ["superadmin", "masterdeveloper"]:
        sql = """
        SELECT s.*, 
               o.po_number, o.customer_name, COALESCE(NULLIF(s.brand, ''), o.brand) as resolved_brand, o.season, o.delivery_date, o.order_date
        FROM styles s
        JOIN orders o ON s.order_id = o.id
        ORDER BY o.id DESC, s.id ASC
        """
        c.execute(sql)
    else:
        sql = """
        SELECT s.*, 
               o.po_number, o.customer_name, COALESCE(NULLIF(s.brand, ''), o.brand) as resolved_brand, o.season, o.delivery_date, o.order_date
        FROM styles s
        JOIN orders o ON s.order_id = o.id
        WHERE s.company_id = ?
        ORDER BY o.id DESC, s.id ASC
        """
        c.execute(sql, (company_id,))
    styles = [dict(r) for r in c.fetchall()]
    for s in styles:
        if s.get("resolved_brand"):
            s["brand"] = s["resolved_brand"]
    
    # Load FabricTag global lookup map once
    ft_map = {}
    try:
        ft_conn = FabricTagConnector.get_connection()
        if ft_conn:
            ft_c = ft_conn.cursor()
            ft_c.execute("SELECT id, internal_code, company_name, quality_name, quality_code, design_code, color, width, weight FROM fabrics")
            for f_row in ft_c.fetchall():
                f_d = dict(f_row)
                if f_d.get("id"):
                    ft_map[str(f_d["id"])] = f_d
                if f_d.get("internal_code"):
                    ft_map[f_d["internal_code"].upper()] = f_d
            ft_conn.close()
    except Exception:
        pass

    def resolve_fabric_meta(fab_id, raw_article):
        found = None
        if fab_id and str(fab_id) in ft_map:
            found = ft_map[str(fab_id)]
        elif raw_article:
            m = re.search(r'ELT\d+', str(raw_article), re.IGNORECASE)
            if m and m.group(0).upper() in ft_map:
                found = ft_map[m.group(0).upper()]
        return found

    all_sizes_set = set()

    for s in styles:
        c.execute("SELECT * FROM size_distributions WHERE style_id = ? ORDER BY sort_order ASC, id ASC", (s["id"],))
        sizes_rows = [dict(sz) for sz in c.fetchall()]
        s["sizes"] = sizes_rows
        
        size_map = {}
        for sz in sizes_rows:
            s_name = str(sz["size_name"]).strip()
            qty = sz["quantity"] or 0
            size_map[s_name] = qty
            if qty > 0:
                all_sizes_set.add(s_name)
        s["size_map"] = size_map
        
        if sizes_rows:
            sum_qty = sum(sz["quantity"] for sz in sizes_rows)
            if sum_qty != s["total_quantity"]:
                s["total_quantity"] = sum_qty
            s["cut_total_quantity"] = sum((sz.get("cut_quantity") or 0) for sz in sizes_rows)
            s["shipped_total_quantity"] = sum((sz.get("shipped_quantity") or 0) for sz in sizes_rows)
        else:
            s["cut_total_quantity"] = 0
            s["shipped_total_quantity"] = 0
        
        s["custom_fields"] = json.loads(s.get("custom_fields_json") or "{}")

        def is_empty_val(v):
            if not v:
                return True
            sv = str(v).strip().upper()
            return sv in ["", "KODSUZ", "YOK", "NONE", "NULL", "-"]

        # Slot 1 Fabric resolution
        c.execute("SELECT * FROM fabric_links WHERE style_id = ? AND (fabric_slot = 1 OR fabric_slot IS NULL) ORDER BY id DESC LIMIT 1", (s["id"],))
        fl1 = c.fetchone()
        fl1_dict = dict(fl1) if fl1 else None
        s["fabric_link"] = fl1_dict
        f1_meta = resolve_fabric_meta(fl1_dict.get("fabrictag_id") if fl1_dict else None, s.get("fabric_article") or s.get("fabric_type"))
        if f1_meta:
            q_name = (f1_meta.get("quality_name") or "").strip()
            q_code = (f1_meta.get("quality_code") or "").strip()
            comp = (f1_meta.get("company_name") or "").strip()
            des = (f1_meta.get("design_code") or "").strip()
            col = (f1_meta.get("color") or "").strip()
            s["fabric_company_1"] = comp
            s["fabric_variant_1"] = des
            s["fabric_color_1"] = col
            if is_empty_val(q_name) and is_empty_val(q_code):
                s["fabric_display_title_1"] = ""
                s["fabric_quality_name_1"] = ""
                s["fabric_quality_code_1"] = ""
            else:
                s["fabrictag_width"] = f1_meta.get("width")
                s["fabrictag_weight"] = f1_meta.get("weight")
                clean_name = "" if is_empty_val(q_name) else q_name
                clean_code = "" if is_empty_val(q_code) else f"({q_code})"
                q_combo = " ".join(filter(None, [clean_name, clean_code])).strip()
                parts = []
                if comp: parts.append(comp)
                if q_combo: parts.append(q_combo)
                if des: parts.append(des)
                if col and not is_empty_val(col): parts.append(col)
                s["fabric_display_title_1"] = " - ".join(parts) if parts else ""
                s["fabric_quality_name_1"] = clean_name
                s["fabric_quality_code_1"] = "" if is_empty_val(q_code) else q_code
        else:
            s["fabric_display_title_1"] = ""
            s["fabric_quality_name_1"] = ""
            s["fabric_quality_code_1"] = ""
            s["fabric_company_1"] = ""
            s["fabric_variant_1"] = ""
            s["fabric_color_1"] = ""

        # Slot 2 Fabric resolution
        c.execute("SELECT * FROM fabric_links WHERE style_id = ? AND fabric_slot = 2 ORDER BY id DESC LIMIT 1", (s["id"],))
        fl2 = c.fetchone()
        fl2_dict = dict(fl2) if fl2 else None
        s["fabric_link_2"] = fl2_dict
        f2_meta = resolve_fabric_meta(fl2_dict.get("fabrictag_id") if fl2_dict else None, s.get("fabric_article_2") or s.get("fabric_type_2"))
        if f2_meta:
            q_name2 = (f2_meta.get("quality_name") or "").strip()
            q_code2 = (f2_meta.get("quality_code") or "").strip()
            comp2 = (f2_meta.get("company_name") or "").strip()
            des2 = (f2_meta.get("design_code") or "").strip()
            col2 = (f2_meta.get("color") or "").strip()
            s["fabric_company_2"] = comp2
            s["fabric_variant_2"] = des2
            s["fabric_color_2"] = col2
            if is_empty_val(q_name2) and is_empty_val(q_code2):
                s["fabric_display_title_2"] = ""
                s["fabric_quality_name_2"] = ""
                s["fabric_quality_code_2"] = ""
            else:
                s["fabrictag_width_2"] = f2_meta.get("width")
                s["fabrictag_weight_2"] = f2_meta.get("weight")
                clean_name2 = "" if is_empty_val(q_name2) else q_name2
                clean_code2 = "" if is_empty_val(q_code2) else f"({q_code2})"
                q_combo2 = " ".join(filter(None, [clean_name2, clean_code2])).strip()
                parts2 = []
                if comp2: parts2.append(comp2)
                if q_combo2: parts2.append(q_combo2)
                if des2: parts2.append(des2)
                if col2 and not is_empty_val(col2): parts2.append(col2)
                s["fabric_display_title_2"] = " - ".join(parts2) if parts2 else ""
                s["fabric_quality_name_2"] = clean_name2
                s["fabric_quality_code_2"] = "" if is_empty_val(q_code2) else q_code2
        else:
            s["fabric_display_title_2"] = ""
            s["fabric_quality_name_2"] = ""
            s["fabric_quality_code_2"] = ""
            s["fabric_company_2"] = ""
            s["fabric_variant_2"] = ""
            s["fabric_color_2"] = ""

        if s.get("fabric_wastage_percent") is None:
            s["fabric_wastage_percent"] = 5.0
        if not s.get("fabric_order_unit"):
            s["fabric_order_unit"] = "M"
        if not s.get("fabric_price_currency_1"):
            s["fabric_price_currency_1"] = s.get("currency") or "€"
        if s.get("fabric_wastage_percent_2") is None:
            s["fabric_wastage_percent_2"] = 5.0
        if not s.get("fabric_order_unit_2"):
            s["fabric_order_unit_2"] = "M"
        if not s.get("fabric_price_currency_2"):
            s["fabric_price_currency_2"] = s.get("currency") or "€"


        if s.get("cost_data_json"):
            try:
                s["cost_data"] = json.loads(s["cost_data_json"])
            except Exception:
                s["cost_data"] = None
        else:
            s["cost_data"] = None

        if s.get("actual_cost_data_json"):
            try:
                s["actual_cost_data"] = json.loads(s["actual_cost_data_json"])
            except Exception:
                s["actual_cost_data"] = None
        else:
            s["actual_cost_data"] = None


        c.execute("SELECT * FROM production_status_logs WHERE style_id = ? ORDER BY id DESC LIMIT 15", (s["id"],))
        s["recent_logs"] = [dict(lg) for lg in c.fetchall()]



        c.execute("SELECT COUNT(*) FROM audit_logs WHERE style_id = ?", (s["id"],))
        s["audit_count"] = c.fetchone()[0]

    # Also load custom columns
    c.execute("SELECT * FROM custom_columns WHERE company_id = ? ORDER BY sort_order ASC, id ASC", (company_id,))
    custom_cols = [dict(r) for r in c.fetchall()]
    for col in custom_cols:
        col["options"] = json.loads(col.get("options_json") or "[]")
        if col.get("col_type") == "size" and col.get("title"):
            all_sizes_set.add(str(col["title"]).strip().upper())

    conn.close()
    
    sorted_sizes = sorted(list(all_sizes_set), key=get_size_sort_key)
    
    return {
        "styles": styles,
        "available_sizes": sorted_sizes,
        "custom_columns": custom_cols
    }

# ----------------- BULK SAVE ENDPOINT -----------------

class BulkSaveItem(BaseModel):
    id: int
    description: Optional[Any] = None
    fabric_order_status: Optional[Any] = None
    unit_price: Optional[Any] = None
    status: Optional[Any] = None
    fabric_ordered: Optional[Any] = None
    pps_sent: Optional[Any] = None
    shipping_sample_sent: Optional[Any] = None
    unit_meters: Optional[Any] = None
    unit_grams: Optional[Any] = None
    fabric_wastage_percent: Optional[Any] = None
    fabric_ordered_meters: Optional[Any] = None
    fabric_price_1: Optional[Any] = None
    fabric_price_currency_1: Optional[Any] = None
    fabric_arrival_date: Optional[Any] = None
    fabric_received_meters: Optional[Any] = None
    unit_meters_2: Optional[Any] = None
    unit_grams_2: Optional[Any] = None
    fabric_wastage_percent_2: Optional[Any] = None
    fabric_ordered_meters_2: Optional[Any] = None
    fabric_price_2: Optional[Any] = None
    fabric_price_currency_2: Optional[Any] = None
    fabric_arrival_date_2: Optional[Any] = None
    fabric_received_meters_2: Optional[Any] = None
    channel: Optional[Any] = None
    size_map: Optional[Dict[str, Any]] = None
    custom_fields: Optional[Dict[str, Any]] = None

class BulkSaveRequest(BaseModel):
    items: List[BulkSaveItem]

@app.post("/api/styles/bulk-save")
def api_bulk_save_styles(req: BulkSaveRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    user_name = user.get("full_name") or user.get("username")
    saved_count = 0

    for item in req.items:
        u_price = None
        if item.unit_price is not None and str(item.unit_price).strip() != "":
            try:
                u_price = float(str(item.unit_price).replace(",", "."))
            except Exception:
                pass

        fp1 = None
        if item.fabric_price_1 is not None and str(item.fabric_price_1).strip() != "":
            try:
                fp1 = float(str(item.fabric_price_1).replace(",", "."))
            except Exception:
                pass

        fp2 = None
        if item.fabric_price_2 is not None and str(item.fabric_price_2).strip() != "":
            try:
                fp2 = float(str(item.fabric_price_2).replace(",", "."))
            except Exception:
                pass

        c.execute("""
        UPDATE styles SET
            description = COALESCE(?, description),
            fabric_order_status = COALESCE(?, fabric_order_status),
            unit_price = COALESCE(?, unit_price),
            status = COALESCE(?, status),
            unit_meters = COALESCE(?, unit_meters),
            unit_grams = COALESCE(?, unit_grams),
            fabric_wastage_percent = COALESCE(?, fabric_wastage_percent),
            fabric_ordered_meters = COALESCE(?, fabric_ordered_meters),
            fabric_price_1 = COALESCE(?, fabric_price_1),
            fabric_price_currency_1 = COALESCE(?, fabric_price_currency_1),
            unit_meters_2 = COALESCE(?, unit_meters_2),
            unit_grams_2 = COALESCE(?, unit_grams_2),
            fabric_wastage_percent_2 = COALESCE(?, fabric_wastage_percent_2),
            fabric_ordered_meters_2 = COALESCE(?, fabric_ordered_meters_2),
            fabric_price_2 = COALESCE(?, fabric_price_2),
            fabric_price_currency_2 = COALESCE(?, fabric_price_currency_2),
            fabric_arrival_date = COALESCE(?, fabric_arrival_date),
            fabric_received_meters = COALESCE(?, fabric_received_meters),
            fabric_arrival_date_2 = COALESCE(?, fabric_arrival_date_2),
            fabric_received_meters_2 = COALESCE(?, fabric_received_meters_2),
            channel = COALESCE(?, channel),
            updated_at = CURRENT_TIMESTAMP
        WHERE id = ?
        """, (
            item.description,
            item.fabric_order_status,
            u_price,
            item.status,
            item.unit_meters,
            item.unit_grams,
            item.fabric_wastage_percent,
            item.fabric_ordered_meters,
            fp1,
            item.fabric_price_currency_1,
            item.unit_meters_2,
            item.unit_grams_2,
            item.fabric_wastage_percent_2,
            item.fabric_ordered_meters_2,
            fp2,
            item.fabric_price_currency_2,
            item.fabric_arrival_date,
            item.fabric_received_meters,
            item.fabric_arrival_date_2,
            item.fabric_received_meters_2,
            item.channel,
            item.id
        ))


        if item.custom_fields is not None:
            c.execute("UPDATE styles SET custom_fields_json = ? WHERE id = ?", (json.dumps(item.custom_fields), item.id))

        if item.size_map:
            for sz_name, sz_qty in item.size_map.items():
                try:
                    qty_int = int(str(sz_qty).strip() or 0)
                except Exception:
                    qty_int = 0
                c.execute("SELECT id FROM size_distributions WHERE style_id = ? AND size_name = ?", (item.id, sz_name))
                row = c.fetchone()
                if row:
                    c.execute("UPDATE size_distributions SET quantity = ? WHERE id = ?", (qty_int, row[0]))
                else:
                    c.execute("INSERT INTO size_distributions (style_id, size_name, quantity) VALUES (?, ?, ?)", (item.id, sz_name, qty_int))
            
            c.execute("SELECT SUM(quantity) FROM size_distributions WHERE style_id = ?", (item.id,))
            new_tot = c.fetchone()[0] or 0
            c.execute("UPDATE styles SET total_quantity = ?, total_amount = ? * COALESCE(unit_price, 0) WHERE id = ?", (new_tot, new_tot, item.id))
        
        saved_count += 1

    conn.commit()
    conn.close()
    return {"status": "success", "message": f"{saved_count} satırdaki tüm değişiklikler başarıyla kaydedildi."}


# ----------------- DELETE STYLE / ORDER -----------------

def check_can_delete_orders(user: Dict[str, Any]):
    role = user.get("role")
    perms = user.get("permissions") or {}
    if role in ["superadmin", "masterdeveloper", "admin"]:
        return True
    if perms.get("can_delete_orders"):
        return True
    raise HTTPException(status_code=403, detail="Sipariş ve modelleri silme yetkiniz bulunmamaktadır.")

@app.delete("/api/styles/{style_id}")
def api_delete_style(style_id: int, user: Dict[str, Any] = Depends(require_user)):
    check_can_delete_orders(user)
    conn = get_db()
    c = conn.cursor()

    c.execute("SELECT order_id, style_no, color_name, total_quantity FROM styles WHERE id = ?", (style_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Silinecek model bulunamadı.")
    
    order_id, style_no, color_name, total_qty = row[0], row[1], row[2], row[3]

    try:
        c.execute("DELETE FROM audit_logs WHERE style_id = ?", (style_id,))
        c.execute("DELETE FROM size_distributions WHERE style_id = ?", (style_id,))
        c.execute("DELETE FROM production_status_logs WHERE style_id = ?", (style_id,))
        c.execute("DELETE FROM fabric_links WHERE style_id = ?", (style_id,))
        c.execute("DELETE FROM styles WHERE id = ?", (style_id,))

        c.execute("SELECT COUNT(*), COALESCE(SUM(total_quantity), 0), COALESCE(SUM(total_amount), 0) FROM styles WHERE order_id = ?", (order_id,))
        rem_count, rem_qty, rem_amt = c.fetchone()

        if rem_count == 0:
            c.execute("DELETE FROM orders WHERE id = ?", (order_id,))
        else:
            c.execute("UPDATE orders SET total_quantity = ?, total_amount = ? WHERE id = ?", (rem_qty, rem_amt, order_id))

        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        raise HTTPException(status_code=500, detail=f"Silme hatası: {str(e)}")


# ----------------- ORDER MANAGEMENT & PERMANENT DELETE -----------------

@app.get("/api/orders/list")
def api_get_orders_list(user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    role = user.get("role")
    comp_id = user.get("company_id")
    
    if role in ["superadmin", "masterdeveloper"]:
        c.execute("""
            SELECT o.id, o.po_number, o.customer_name, o.brand, o.season, o.order_date, o.delivery_date,
                   o.total_quantity, o.total_amount, o.currency, o.status, o.created_at,
                   COUNT(s.id) as styles_count,
                   GROUP_CONCAT(DISTINCT s.style_no) as style_names
            FROM orders o
            LEFT JOIN styles s ON s.order_id = o.id
            GROUP BY o.id
            ORDER BY o.id DESC
        """)
    else:
        c.execute("""
            SELECT o.id, o.po_number, o.customer_name, o.brand, o.season, o.order_date, o.delivery_date,
                   o.total_quantity, o.total_amount, o.currency, o.status, o.created_at,
                   COUNT(s.id) as styles_count,
                   GROUP_CONCAT(DISTINCT s.style_no) as style_names
            FROM orders o
            LEFT JOIN styles s ON s.order_id = o.id
            WHERE o.company_id = ?
            GROUP BY o.id
            ORDER BY o.id DESC
        """, (comp_id or 1,))
        
    rows = [dict(r) for r in c.fetchall()]
    
    if rows:
        order_ids = [r["id"] for r in rows]
        placeholders = ",".join("?" for _ in order_ids)
        c.execute(f"""
            SELECT id, order_id, style_no, description, color_code, color_name, total_quantity, status, image_url
            FROM styles
            WHERE order_id IN ({placeholders})
            ORDER BY id ASC
        """, tuple(order_ids))
        all_styles = [dict(s) for s in c.fetchall()]
        styles_map = {}
        for s in all_styles:
            styles_map.setdefault(s["order_id"], []).append(s)
        for r in rows:
            r["styles"] = styles_map.get(r["id"], [])
            # Also build a comprehensive search string for colors and styles
            color_parts = []
            for s in r["styles"]:
                if s.get("color_code"):
                    color_parts.append(str(s["color_code"]).strip())
                if s.get("color_name"):
                    color_parts.append(str(s["color_name"]).strip())
            r["color_names"] = ", ".join(dict.fromkeys(color_parts))

    conn.close()
    return rows

@app.delete("/api/orders/{order_id}")
def api_delete_order(order_id: int, user: Dict[str, Any] = Depends(require_user)):
    check_can_delete_orders(user)
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT id, po_number, customer_name FROM orders WHERE id = ?", (order_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Sipariş bulunamadı.")
    
    po_num, cust = row["po_number"], row["customer_name"]

    try:
        c.execute("SELECT id FROM styles WHERE order_id = ?", (order_id,))
        style_ids = [r["id"] for r in c.fetchall()]

        for s_id in style_ids:
            c.execute("DELETE FROM audit_logs WHERE style_id = ?", (s_id,))
            c.execute("DELETE FROM size_distributions WHERE style_id = ?", (s_id,))
            c.execute("DELETE FROM production_status_logs WHERE style_id = ?", (s_id,))
            c.execute("DELETE FROM fabric_links WHERE style_id = ?", (s_id,))
        
        c.execute("DELETE FROM styles WHERE order_id = ?", (order_id,))
        c.execute("DELETE FROM orders WHERE id = ?", (order_id,))
        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        raise HTTPException(status_code=500, detail=f"Sipariş silme hatası: {str(e)}")

    conn.close()
    return {"status": "success", "message": f"Sipariş (PO: {po_num}, Müşteri: {cust}) ve tüm modelleri kalıcı olarak silindi."}

class BulkDeleteOrdersRequest(BaseModel):
    order_ids: List[int]

@app.post("/api/orders/bulk-delete")
def api_bulk_delete_orders(req: BulkDeleteOrdersRequest, user: Dict[str, Any] = Depends(require_user)):
    if not req.order_ids:
        return {"status": "success", "deleted_count": 0}
    
    conn = get_db()
    c = conn.cursor()
    try:
        for order_id in req.order_ids:
            c.execute("SELECT id FROM styles WHERE order_id = ?", (order_id,))
            style_ids = [r["id"] for r in c.fetchall()]
            for s_id in style_ids:
                c.execute("DELETE FROM audit_logs WHERE style_id = ?", (s_id,))
                c.execute("DELETE FROM size_distributions WHERE style_id = ?", (s_id,))
                c.execute("DELETE FROM production_status_logs WHERE style_id = ?", (s_id,))
                c.execute("DELETE FROM fabric_links WHERE style_id = ?", (s_id,))
            c.execute("DELETE FROM styles WHERE order_id = ?", (order_id,))
            c.execute("DELETE FROM orders WHERE id = ?", (order_id,))
        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        raise HTTPException(status_code=500, detail=f"Toplu silme hatası: {str(e)}")

    conn.close()
    return {"status": "success", "deleted_count": len(req.order_ids)}




def get_excel_currency_format(cur_code: str = "EUR") -> tuple[str, str]:
    """
    Excel'e aktarılan tüm para birimlerinde sembolü rakamın önüne alır.
    Türkçe Excel'de '#,##0.00' formatı küsüratı virgül ',', binlikleri nokta '.' olarak görüntüler.
    Hücreye gerçek sayı (float) yazıldığında Excel bunu tam olarak sayı/para birimi olarak algılar.
    Örnek:
      EUR: symbol='€', format='"€" #,##0.00' -> '€ 26,25'
      TRY: symbol='₺', format='"₺" #,##0.00' -> '₺ 1.250,50'
      USD: symbol='$', format='"$" #,##0.00' -> '$ 45,50'
      GBP: symbol='£', format='"£" #,##0.00' -> '£ 30,00'
    """
    c = str(cur_code or "EUR").upper().strip()
    if c in ["USD", "$"]:
        return "$", '"$" #,##0.00'
    elif c in ["TL", "TRY", "₺"]:
        return "₺", '"₺" #,##0.00'
    elif c in ["GBP", "£"]:
        return "£", '"£" #,##0.00'
    else:
        return "€", '"€" #,##0.00'

def safe_price_float(val):
    if val is None or str(val).strip() == "":
        return None
    if isinstance(val, (int, float)):
        return float(val)
    try:
        s_clean = str(val).strip().replace("€", "").replace("₺", "").replace("$", "").replace("£", "").replace("TL", "").replace("EUR", "").replace("USD", "").strip()
        s_clean = s_clean.replace(",", ".")
        return float(s_clean)
    except Exception:
        return None

class ExportExcelRequest(BaseModel):
    style_ids: Optional[List[int]] = None

def generate_styles_excel_workbook(style_ids: Optional[List[int]] = None, company_id: int = 1) -> Response:
    conn = get_db()
    c = conn.cursor()
    
    if style_ids and len(style_ids) > 0:
        placeholders = ",".join("?" for _ in style_ids)
        order_by_clause = "CASE s.id " + " ".join(f"WHEN {sid} THEN {i}" for i, sid in enumerate(style_ids)) + " END"
        sql = f"""
        SELECT s.*, 
               o.po_number, o.customer_name, o.brand, o.season, o.delivery_date, o.order_date
        FROM styles s
        JOIN orders o ON s.order_id = o.id
        WHERE s.company_id = ? AND s.id IN ({placeholders})
        ORDER BY {order_by_clause}
        """
        params = [company_id] + style_ids
    else:
        sql = """
        SELECT s.*, 
               o.po_number, o.customer_name, o.brand, o.season, o.delivery_date, o.order_date
        FROM styles s
        JOIN orders o ON s.order_id = o.id
        WHERE s.company_id = ?
        ORDER BY o.id DESC, s.id ASC
        """
        params = [company_id]
        
    c.execute(sql, tuple(params))
    styles = [dict(r) for r in c.fetchall()]
    all_style_ids = [s["id"] for s in styles]
    
    # 1. Beden Dağılımlarını Topla
    sizes_by_style = {}
    present_size_names = []
    size_seen = set()
    
    if all_style_ids:
        pl_ids = ",".join("?" for _ in all_style_ids)
        c.execute(f"""
        SELECT style_id, size_name, quantity, sort_order
        FROM size_distributions
        WHERE style_id IN ({pl_ids})
        ORDER BY sort_order ASC, id ASC
        """, tuple(all_style_ids))
        for sz_row in c.fetchall():
            sid = sz_row["style_id"]
            sname = str(sz_row["size_name"]).strip()
            qty = sz_row["quantity"] or 0
            if sid not in sizes_by_style:
                sizes_by_style[sid] = {}
            sizes_by_style[sid][sname] = qty
            if sname not in size_seen and qty > 0:
                size_seen.add(sname)
                present_size_names.append(sname)
                
    present_size_names.sort(key=get_size_sort_key)
    
    # 2. Kumaş Bağlantılarını Topla (fabric_links)
    fl_by_style = {}
    if all_style_ids:
        pl_ids = ",".join("?" for _ in all_style_ids)
        c.execute(f"""
        SELECT * FROM fabric_links WHERE style_id IN ({pl_ids}) ORDER BY id ASC
        """, tuple(all_style_ids))
        for fl in c.fetchall():
            sid = fl["style_id"]
            if sid not in fl_by_style:
                fl_by_style[sid] = dict(fl)

    # FabricTag detaylarını entegre et
    ft_map = {}
    try:
        ft_conn = FabricTagConnector.get_connection()
        if ft_conn:
            ft_c = ft_conn.cursor()
            ft_c.execute("SELECT id, internal_code, company_name, quality_name, quality_code, design_code, color, width, weight, composition FROM fabrics")
            for f_row in ft_c.fetchall():
                f_d = dict(f_row)
                if f_d.get("id"):
                    ft_map[str(f_d["id"])] = f_d
                if f_d.get("internal_code"):
                    ft_map[f_d["internal_code"].upper()] = f_d
                if f_d.get("quality_code"):
                    ft_map[f_d["quality_code"].upper()] = f_d
            ft_conn.close()
    except Exception:
        pass
        
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sipariş Listesi"
    
    # Modern Excel Görünümü: Izgara çizgilerini göster ve üst başlığı dondur
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "B2"
    
    # Başlıkları Oluştur
    headers = [
        "Görsel", "Müşteri", "Marka", "Kanal", "Sezon", "PO No", "Model No", 
        "Renk Kodu", "Renk Adı"
    ]
    
    size_col_start = len(headers) + 1
    for sz_name in present_size_names:
        headers.append(f"Beden {sz_name}")
    size_col_end = len(headers)
    
    col_idx_qty = len(headers) + 1
    headers.append("Toplam Adet")
    col_idx_price = len(headers) + 1
    headers.append("Birim Fiyat")
    col_idx_currency = len(headers) + 1
    headers.append("Döviz")
    col_idx_amount = len(headers) + 1
    headers.append("Toplam Tutar")
    
    # Kumaş Bilgileri Sütunları (FabricTag Kumaş Deposu)
    headers.extend([
        "Kumaşçı",
        "Kalite Kodu",
        "Kalite Adı",
        "Varyant",
        "Renk",
        "Karışım",
        "Ağırlık",
        "En",
        "Sipariş Kumaş (M/KG)",
        "Gelen Kumaş (M/KG)",
        "Kumaş Durumu / Termin"
    ])
    
    col_idx_ordered_fabric = headers.index("Sipariş Kumaş (M/KG)") + 1
    col_idx_received_fabric = headers.index("Gelen Kumaş (M/KG)") + 1
    
    # Diğer Sütunlar
    headers.extend([
        "Sipariş Tarihi",
        "Yükleme Tarihi",
        "Üretim Aşaması",
        "Beden Dağılımı (Özet)",
        "Notlar"
    ])
    
    # Modern Başlık Tasarımı (Koyu Füme / Slate Navy, Beyaz Kalın Yazı)
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    header_border = Border(
        left=Side(style="thin", color="334155"),
        right=Side(style="thin", color="334155"),
        top=Side(style="thin", color="0F172A"),
        bottom=Side(style="medium", color="0F172A")
    )
    
    ws.row_dimensions[1].height = 34
    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = header_border
        
    # Görsel sütunu genişliği
    ws.column_dimensions["A"].width = 18
    img_buffers = []
    
    # Hücre kenarlık ve yazı tipleri
    cell_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )
    
    font_normal = Font(name="Segoe UI", size=9.5, color="1E293B")
    font_bold = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
    font_style_no = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    font_muted = Font(name="Segoe UI", size=9, color="94A3B8")
    
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    
    last_row = 1
    
    for idx, s in enumerate(styles, start=2):
        last_row = idx
        ws.row_dimensions[idx].height = 75
        current_fill = fill_white if (idx % 2 == 0) else fill_zebra
        
        # Beden dağılımı metni
        style_size_map = sizes_by_style.get(s["id"], {})
        sizes_summary_str = ", ".join(f"{sz}: {qty}" for sz, qty in style_size_map.items() if qty > 0)
        
        # Notlar
        notes_val = s.get("notes") or ""
        if not notes_val and s.get("custom_fields_json"):
            try:
                cf = json.loads(s["custom_fields_json"])
                notes_val = cf.get("notes") or cf.get("notlar") or ""
            except Exception:
                pass
                
        qty = s.get("total_quantity") or 0
        price = s.get("unit_price") or 0
        tot_amt = (price or 0) * (qty or 0)
        
        # Kumaş Bilgileri (FabricTag ve stiller tablosu harmanlanır)
        fl = fl_by_style.get(s["id"], {})
        f_meta = None
        if fl.get("fabrictag_id") and str(fl["fabrictag_id"]) in ft_map:
            f_meta = ft_map[str(fl["fabrictag_id"])]
        if not f_meta:
            raw_art = str(s.get("fabric_article") or s.get("fabric_type") or "")
            m_elt = re.search(r'ELT\d+', raw_art, re.IGNORECASE)
            if m_elt and m_elt.group(0).upper() in ft_map:
                f_meta = ft_map[m_elt.group(0).upper()]
        if not f_meta and fl.get("fabric_code") and fl["fabric_code"].upper() in ft_map:
            f_meta = ft_map[fl["fabric_code"].upper()]

        def clean_field(val):
            if not val:
                return ""
            sv = str(val).strip()
            if sv.upper() in ["KODSUZ", "YOK", "NONE", "NULL", "-", "—", "İSİMSİZ", "ISIMSIZ", "TANIMSIZ", "BİLGİSİ YOK", "BILGISI YOK", "BİLGİ YOK", "BILGI YOK", "BELİRTİLMEDİ", "BELIRTILMEDI"]:
                return ""
            return sv

        def clean_weight(val):
            v = clean_field(val)
            if not v:
                return ""
            v_num = re.sub(r'^(gr|gramaj|weight)[:\s]*', '', v, flags=re.I).strip()
            if v_num.isdigit():
                return f"{v_num} gr"
            return v

        def clean_width(val):
            v = clean_field(val)
            if not v:
                return ""
            v_num = re.sub(r'^(en|genişlik|width)[:\s]*', '', v, flags=re.I).strip()
            if v_num.isdigit():
                return f"{v_num} cm"
            return v

        fabric_supplier_val = clean_field(f_meta.get("company_name") if f_meta else (fl.get("supplier") or s.get("fabric_company_1") or ""))
        fabric_qcode_val = clean_field(f_meta.get("quality_code") if f_meta else (fl.get("fabric_code") or s.get("fabric_quality_code_1") or ""))
        fabric_qname_val = clean_field(f_meta.get("quality_name") if f_meta else (fl.get("fabric_name") or s.get("fabric_quality_name_1") or s.get("fabric_article") or ""))
        fabric_variant_val = clean_field(f_meta.get("design_code") if f_meta else (s.get("fabric_variant_1") or ""))
        fabric_color_val = clean_field(f_meta.get("color") if f_meta else (s.get("fabric_color_1") or ""))
        fabric_comp_val = clean_field(f_meta.get("composition") if f_meta else (fl.get("composition") or s.get("fabric_composition") or ""))
        fabric_weight_val = clean_weight(f_meta.get("weight") if f_meta else (s.get("fabric_weight") or fl.get("ft_weight") or ""))
        fabric_width_val = clean_width(f_meta.get("width") if f_meta else (s.get("fabric_width") or fl.get("ft_width") or ""))

        fabric_ordered_m = s.get("fabric_ordered_meters") or 0.0
        fabric_received_m = s.get("fabric_received_meters") or 0.0
        fabric_status_val = s.get("fabric_order_status") or fl.get("status") or ""
        
        # A Sütunu: Görsel Hücresi
        cell_a = ws.cell(row=idx, column=1)
        cell_a.border = cell_border
        cell_a.fill = current_fill
        
        img_url = s.get("image_url") or s.get("image_url_2")
        img_added = False
        if img_url:
            try:
                pil_img = None
                if img_url.startswith("data:image"):
                    hdr, enc = img_url.split(",", 1)
                    pil_img = PILImage.open(io.BytesIO(base64.b64decode(enc)))
                elif img_url.startswith("/uploads/"):
                    p = UPLOADS_DIR / img_url.replace("/uploads/", "").lstrip("/")
                    if p.exists():
                        pil_img = PILImage.open(p)
                elif Path(img_url).exists():
                    pil_img = PILImage.open(img_url)
                    
                if pil_img:
                    pil_img = pil_img.convert("RGB")
                    pil_img.thumbnail((105, 88), PILImage.Resampling.LANCZOS)
                    buf = io.BytesIO()
                    pil_img.save(buf, format="JPEG", quality=88)
                    buf.seek(0)
                    img_buffers.append(buf)
                    xl_img = OpenpyxlImage(buf)
                    img_w = pil_img.width
                    img_h = pil_img.height

                    # Hücrede En ve Boydan Ortalama (Horizontal & Vertical Center)
                    # Sütun A genişliği: 18 karakter (~140px), Satır yüksekliği: 75pt (~100px)
                    col_w_px = 140
                    row_h_px = 100
                    x_off_px = max(0, (col_w_px - img_w) // 2)
                    y_off_px = max(0, (row_h_px - img_h) // 2)

                    marker = AnchorMarker(col=0, colOff=pixels_to_EMU(x_off_px), row=idx - 1, rowOff=pixels_to_EMU(y_off_px))
                    size = XDRPositiveSize2D(pixels_to_EMU(img_w), pixels_to_EMU(img_h))
                    xl_img.anchor = OneCellAnchor(_from=marker, ext=size)
                    ws.add_image(xl_img)
                    img_added = True
            except Exception:
                pass
                
        if not img_added:
            cell_a.value = "[Görsel Yok]"
            cell_a.alignment = align_center
            cell_a.font = font_muted
            
        # Satır Verilerini Sırayla Doldur
        # 1. Temel Bilgiler (2-9)
        ws.cell(row=idx, column=2, value=s.get("customer_name") or "").alignment = align_left
        ws.cell(row=idx, column=2).font = font_bold
        ws.cell(row=idx, column=3, value=s.get("brand") or "").alignment = align_left
        ws.cell(row=idx, column=4, value=s.get("channel") or "").alignment = align_left
        ws.cell(row=idx, column=5, value=s.get("season") or "").alignment = align_center
        ws.cell(row=idx, column=6, value=s.get("po_number") or "").alignment = align_center
        ws.cell(row=idx, column=6).font = font_bold
        ws.cell(row=idx, column=7, value=s.get("style_no") or "").alignment = align_center
        ws.cell(row=idx, column=7).font = font_style_no
        ws.cell(row=idx, column=8, value=s.get("color_code") or "").alignment = align_center
        ws.cell(row=idx, column=9, value=s.get("color_name") or "").alignment = align_left_wrap
        
        # 2. Beden Dağılımı Sütunları (TÜM RAKAMLAR ORTALI)
        for s_idx, sz_name in enumerate(present_size_names, start=size_col_start):
            sz_qty = style_size_map.get(sz_name)
            c_sz = ws.cell(row=idx, column=s_idx, value=sz_qty if sz_qty else None)
            c_sz.alignment = align_center
            if sz_qty:
                c_sz.number_format = "#,##0"
                c_sz.font = font_bold
                
        # 3. Miktar ve Fiyatlar (TÜM RAKAMLAR ORTALI, PARA BİRİMİ ÖNDE, VİRGÜLLÜ)
        c_qty = ws.cell(row=idx, column=col_idx_qty, value=qty)
        c_qty.alignment = align_center
        c_qty.number_format = "#,##0"
        c_qty.font = font_bold
        
        cur_sym, cur_fmt = get_excel_currency_format(s.get("currency"))

        c_prc = ws.cell(row=idx, column=col_idx_price)
        c_prc.alignment = align_center
        prc_num = safe_price_float(price)
        c_prc.value = prc_num if prc_num is not None else 0.0
        c_prc.number_format = cur_fmt
        
        c_cur = ws.cell(row=idx, column=col_idx_currency, value=cur_sym)
        c_cur.alignment = align_center
        
        c_amt = ws.cell(row=idx, column=col_idx_amount)
        c_amt.alignment = align_center
        amt_num = safe_price_float(tot_amt)
        c_amt.value = amt_num if amt_num is not None else 0.0
        c_amt.number_format = cur_fmt
        c_amt.font = font_bold
        
        # 4. Kumaş Bilgileri Sütunları (Kumaşçı, Kalite Kodu, Kalite Adı, Varyant, Renk, Karışım, Ağırlık, En)
        ws.cell(row=idx, column=headers.index("Kumaşçı") + 1, value=fabric_supplier_val).alignment = align_left
        ws.cell(row=idx, column=headers.index("Kalite Kodu") + 1, value=fabric_qcode_val).alignment = align_center
        ws.cell(row=idx, column=headers.index("Kalite Adı") + 1, value=fabric_qname_val).alignment = align_left_wrap
        ws.cell(row=idx, column=headers.index("Varyant") + 1, value=fabric_variant_val).alignment = align_center
        ws.cell(row=idx, column=headers.index("Renk") + 1, value=fabric_color_val).alignment = align_center
        ws.cell(row=idx, column=headers.index("Karışım") + 1, value=fabric_comp_val).alignment = align_left_wrap
        ws.cell(row=idx, column=headers.index("Ağırlık") + 1, value=fabric_weight_val).alignment = align_center
        ws.cell(row=idx, column=headers.index("En") + 1, value=fabric_width_val).alignment = align_center
        
        # Kumaş Metrajları (TÜM RAKAMLAR ORTALI)
        c_ord_f = ws.cell(row=idx, column=col_idx_ordered_fabric, value=fabric_ordered_m if fabric_ordered_m else None)
        c_ord_f.alignment = align_center
        if fabric_ordered_m:
            c_ord_f.number_format = "#,##0.0"
            
        c_rec_f = ws.cell(row=idx, column=col_idx_received_fabric, value=fabric_received_m if fabric_received_m else None)
        c_rec_f.alignment = align_center
        if fabric_received_m:
            c_rec_f.number_format = "#,##0.0"
            c_rec_f.font = font_bold
            
        ws.cell(row=idx, column=headers.index("Kumaş Durumu / Termin") + 1, value=fabric_status_val).alignment = align_left_wrap
        
        # 5. Tarihler, Aşama ve Notlar
        ws.cell(row=idx, column=headers.index("Sipariş Tarihi") + 1, value=s.get("order_date") or "").alignment = align_center
        ws.cell(row=idx, column=headers.index("Yükleme Tarihi") + 1, value=s.get("delivery_date") or "").alignment = align_center
        ws.cell(row=idx, column=headers.index("Üretim Aşaması") + 1, value=s.get("status") or "Planlama aşamasında").alignment = align_center
        ws.cell(row=idx, column=headers.index("Üretim Aşaması") + 1).font = font_bold
        
        ws.cell(row=idx, column=headers.index("Beden Dağılımı (Özet)") + 1, value=sizes_summary_str).alignment = align_left_wrap
        ws.cell(row=idx, column=headers.index("Notlar") + 1, value=notes_val).alignment = align_left_wrap

        # Tüm hücrelerin çerçevesini ve zemin rengini uygula
        for c_idx in range(2, len(headers) + 1):
            cell = ws.cell(row=idx, column=c_idx)
            cell.border = cell_border
            cell.fill = current_fill
            if not cell.font or cell.font.name != "Segoe UI":
                cell.font = font_normal

    # Modern Genel Toplam Satırı
    if last_row >= 2:
        sum_row = last_row + 1
        ws.row_dimensions[sum_row].height = 30
        sum_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        sum_border = Border(
            top=Side(style="thin", color="94A3B8"),
            bottom=Side(style="double", color="0F172A"),
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0")
        )
        
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=sum_row, column=c_idx)
            cell.fill = sum_fill
            cell.border = sum_border
            cell.font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
            cell.alignment = align_center
            
        ws.cell(row=sum_row, column=2, value="GENEL TOPLAM").alignment = Alignment(horizontal="left", vertical="center")
        
        # Beden Sütunları Toplamları (ORTALI)
        for s_idx in range(size_col_start, size_col_end + 1):
            c_letter = get_column_letter(s_idx)
            c_sz_sum = ws.cell(row=sum_row, column=s_idx, value=f"=SUM({c_letter}2:{c_letter}{last_row})")
            c_sz_sum.number_format = "#,##0"
            c_sz_sum.alignment = align_center
        
        # Adet Toplamı (ORTALI)
        col_qty_letter = get_column_letter(col_idx_qty)
        cell_qty = ws.cell(row=sum_row, column=col_idx_qty, value=f"=SUM({col_qty_letter}2:{col_qty_letter}{last_row})")
        cell_qty.number_format = "#,##0"
        cell_qty.alignment = align_center
        
        # Tutar Toplamı (ORTALI, PARA BİRİMİ ÖNDE, VİRGÜLLÜ)
        col_amt_letter = get_column_letter(col_idx_amount)
        cell_amt = ws.cell(row=sum_row, column=col_idx_amount, value=f"=SUM({col_amt_letter}2:{col_amt_letter}{last_row})")
        first_cur = styles[0].get("currency") if styles else "EUR"
        _, tot_cur_fmt = get_excel_currency_format(first_cur)
        cell_amt.number_format = tot_cur_fmt
        cell_amt.alignment = align_center
        
        # Kumaş Metraj Toplamları (ORTALI)
        col_ordf_letter = get_column_letter(col_idx_ordered_fabric)
        c_ordf_sum = ws.cell(row=sum_row, column=col_idx_ordered_fabric, value=f"=SUM({col_ordf_letter}2:{col_ordf_letter}{last_row})")
        c_ordf_sum.number_format = "#,##0.0"
        c_ordf_sum.alignment = align_center
        
        col_recf_letter = get_column_letter(col_idx_received_fabric)
        c_recf_sum = ws.cell(row=sum_row, column=col_idx_received_fabric, value=f"=SUM({col_recf_letter}2:{col_recf_letter}{last_row})")
        c_recf_sum.number_format = "#,##0.0"
        c_recf_sum.alignment = align_center
        
        # Otomatik Filtreleme (AutoFilter) Etkinleştir
        last_col_letter = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

    # Sütun Genişliklerini İçeriğe Göre Otomatik Dengele
    for col_idx in range(2, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[col_letter]:
            val_str = str(cell.value or "")
            if val_str.startswith("="):
                continue
            if "\n" in val_str:
                lines = val_str.split("\n")
                max_len = max(max_len, max(len(l) for l in lines))
            else:
                max_len = max(max_len, len(val_str))
        # Beden sütunları dar kalsın (en az 9, en çok 14)
        if size_col_start <= col_idx <= size_col_end:
            ws.column_dimensions[col_letter].width = max(9, min(14, max_len + 3))
        else:
            ws.column_dimensions[col_letter].width = max(11, min(45, max_len + 3))

    conn.close()

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    
    filename = f"Siparisler_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    val = out.getvalue()
    return Response(
        content=val,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Content-Length": str(len(val)),
            "Cache-Control": "no-cache, no-store, must-revalidate"
        }
    )

@app.get("/api/styles/export-excel")
def api_export_styles_excel_get(ids: Optional[str] = Query(None), authorization: Optional[str] = Header(None)):
    user = None
    if authorization:
        t = authorization.replace("Bearer ", "").strip()
        user = get_current_user(t)
    company_id = (user.get("company_id") if user and user.get("company_id") is not None else 1)
    
    style_ids = None
    if ids:
        try:
            style_ids = [int(x.strip()) for x in ids.split(",") if x.strip().isdigit()]
        except Exception:
            pass
            
    return generate_styles_excel_workbook(style_ids, company_id)

@app.post("/api/styles/export-excel")
def api_export_styles_excel_post(req: ExportExcelRequest, authorization: Optional[str] = Header(None)):
    user = None
    if authorization:
        t = authorization.replace("Bearer ", "").strip()
        user = get_current_user(t)
    company_id = (user.get("company_id") if user and user.get("company_id") is not None else 1)
    
    return generate_styles_excel_workbook(req.style_ids, company_id)

# ----------------- ÖZET YAZDIR İMALAT ÇARŞAF EXCEL EXPORT -----------------

def generate_carsaf_print_excel_workbook(style_ids: Optional[List[int]] = None, title: Optional[str] = None, company_id: int = 1) -> Response:
    conn = get_db()
    c = conn.cursor()
    
    if style_ids and len(style_ids) > 0:
        placeholders = ",".join("?" for _ in style_ids)
        order_by_clause = "CASE s.id " + " ".join(f"WHEN {sid} THEN {i}" for i, sid in enumerate(style_ids)) + " END"
        sql = f"""
        SELECT s.*, 
               o.po_number, o.customer_name, o.brand, o.season, o.delivery_date, o.order_date
        FROM styles s
        JOIN orders o ON s.order_id = o.id
        WHERE s.company_id = ? AND s.id IN ({placeholders})
        ORDER BY {order_by_clause}
        """
        params = [company_id] + style_ids
    else:
        sql = """
        SELECT s.*, 
               o.po_number, o.customer_name, o.brand, o.season, o.delivery_date, o.order_date
        FROM styles s
        JOIN orders o ON s.order_id = o.id
        WHERE s.company_id = ?
        ORDER BY o.id DESC, s.id ASC
        """
        params = [company_id]
        
    c.execute(sql, tuple(params))
    styles = [dict(r) for r in c.fetchall()]
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "İmalat Çarşaf Özeti"
    ws.views.sheetView[0].showGridLines = True
    
    fill_yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_title = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_total = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    
    border_thin = Border(
        left=Side(style="thin", color="64748B"),
        right=Side(style="thin", color="64748B"),
        top=Side(style="thin", color="64748B"),
        bottom=Side(style="thin", color="64748B")
    )
    border_total = Border(
        left=Side(style="thin", color="64748B"),
        right=Side(style="thin", color="64748B"),
        top=Side(style="thin", color="64748B"),
        bottom=Side(style="double", color="0F172A")
    )
    
    font_title = Font(name="Segoe UI", size=13, bold=True, color="0F172A")
    font_header = Font(name="Segoe UI", size=10.5, bold=True, color="000000")
    font_bold = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    font_normal = Font(name="Segoe UI", size=9.5, color="1E293B")
    font_total = Font(name="Segoe UI", size=10.5, bold=True, color="0F172A")
    font_muted = Font(name="Segoe UI", size=8.5, color="94A3B8")
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # 1. Başlık Banner'ı
    header_title_text = (title or "MÜŞTERİ SEZON İMALAT ÇARŞAF LİSTESİ").strip().upper()
    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = header_title_text
    title_cell.font = font_title
    title_cell.alignment = align_center
    title_cell.fill = fill_title
    ws.row_dimensions[1].height = 32
    for c_idx in range(1, 8):
        ws.cell(row=1, column=c_idx).border = border_thin
        
    # 2. Sütun Başlıkları
    headers = ["GÖRSEL", "STYLE", "RENK", "KUMAŞ", "ADET", "PRICE", "NOTLAR"]
    ws.row_dimensions[2].height = 25
    for c_idx, h_name in enumerate(headers, start=1):
        c_cell = ws.cell(row=2, column=c_idx, value=h_name)
        c_cell.fill = fill_yellow
        c_cell.font = font_header
        c_cell.alignment = align_center if c_idx in (1, 2, 5, 6) else align_left
        c_cell.border = border_thin
        
    col_widths = {
        "A": 16,
        "B": 18,
        "C": 24,
        "D": 32,
        "E": 12,
        "F": 14,
        "G": 28
    }
    for col_let, w in col_widths.items():
        ws.column_dimensions[col_let].width = w
        
    # 3. Satırlar
    curr_row = 2
    for s in styles:
        curr_row += 1
        row_fill = fill_white if (curr_row % 2 == 1) else fill_zebra
        ws.row_dimensions[curr_row].height = 75
        
        for c_idx in range(1, 8):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.fill = row_fill
            cell.border = border_thin
            
        # Görsel
        img_added = False
        img_url = s.get("image_url") or s.get("image_url_2")
        if img_url and os.path.exists(UPLOADS_DIR):
            try:
                fn = os.path.basename(img_url.split("?")[0])
                possible_paths = [
                    os.path.join(UPLOADS_DIR, fn),
                    os.path.join(UPLOADS_DIR, "styles", fn)
                ]
                img_path = next((p for p in possible_paths if os.path.isfile(p)), None)
                if img_path:
                    with PILImage.open(img_path) as pil_img:
                        orig_w, orig_h = pil_img.size
                        max_w, max_h = 95, 90
                        ratio = min(max_w / max(1, orig_w), max_h / max(1, orig_h))
                        new_w = max(1, int(orig_w * ratio))
                        new_h = max(1, int(orig_h * ratio))

                    xl_img = OpenpyxlImage(img_path)
                    xl_img.width = new_w
                    xl_img.height = new_h

                    col_w_px = int(16 * 7.5)
                    row_h_px = int(75 * 1.33)
                    x_off_px = max(0, (col_w_px - new_w) // 2)
                    y_off_px = max(0, (row_h_px - new_h) // 2)

                    marker = AnchorMarker(col=0, colOff=pixels_to_EMU(x_off_px), row=curr_row - 1, rowOff=pixels_to_EMU(y_off_px))
                    size = XDRPositiveSize2D(pixels_to_EMU(new_w), pixels_to_EMU(new_h))
                    xl_img.anchor = OneCellAnchor(_from=marker, ext=size)
                    ws.add_image(xl_img)
                    img_added = True
            except Exception:
                pass

        if not img_added:
            cell_a = ws.cell(row=curr_row, column=1, value="[Resim Yok]")
            cell_a.font = font_muted
            cell_a.alignment = align_center

        # Style No
        cell_b = ws.cell(row=curr_row, column=2, value=s.get("style_no") or "")
        cell_b.font = font_bold
        cell_b.alignment = align_center

        # Renk
        c_code = (s.get("color_code") or "").strip()
        c_name = (s.get("color_name") or "").strip()
        color_disp = f"{c_name} ({c_code})" if (c_code and c_name) else (c_name or c_code or "-")
        cell_c = ws.cell(row=curr_row, column=3, value=color_disp)
        cell_c.font = font_bold
        cell_c.alignment = align_left_wrap

        # Kumaş
        fabric_text = s.get("fabric_article") or s.get("fabric_type") or "-"
        cell_d = ws.cell(row=curr_row, column=4, value=fabric_text)
        cell_d.font = font_normal
        cell_d.alignment = align_left_wrap

        # Adet
        qty = int(s.get("total_quantity") or 0)
        cell_e = ws.cell(row=curr_row, column=5, value=qty)
        cell_e.font = font_bold
        cell_e.alignment = align_center
        cell_e.number_format = "#,##0"

        # Price (Sütun 6 / F Sütunu - SEMBOL ÖNDE, VİRGÜLLÜ, GERÇEK SAYI)
        p_val = s.get("unit_price")
        cur_sym, cur_fmt = get_excel_currency_format(s.get("currency"))
        
        cell_f = ws.cell(row=curr_row, column=6)
        cell_f.font = font_bold
        cell_f.alignment = align_center
        
        pf = safe_price_float(p_val)
        if pf is not None:
            cell_f.value = pf
            cell_f.number_format = cur_fmt
        elif p_val is not None and str(p_val).strip() != "":
            s_clean = str(p_val).strip().replace(".", ",")
            cell_f.value = f"{cur_sym} {s_clean}"
        else:
            cell_f.value = None

        # Notlar
        notes_val = s.get("notes") or ""
        if not notes_val and s.get("custom_fields_json"):
            try:
                cf = json.loads(s["custom_fields_json"])
                notes_val = cf.get("notes") or cf.get("notlar") or ""
            except Exception:
                pass
        cell_g = ws.cell(row=curr_row, column=7, value=notes_val)
        cell_g.font = font_normal
        cell_g.alignment = align_left_wrap

    # 4. Toplam Satırı
    if curr_row >= 3:
        sum_row = curr_row + 1
        ws.row_dimensions[sum_row].height = 28
        for c_idx in range(1, 8):
            c = ws.cell(row=sum_row, column=c_idx)
            c.fill = fill_total
            c.border = border_total

        ws.cell(row=sum_row, column=2, value=f"TOPLAM ({len(styles)} Model)").font = font_total
        ws.cell(row=sum_row, column=2).alignment = align_center
        
        cell_sum_qty = ws.cell(row=sum_row, column=5, value=f"=SUM(E3:E{curr_row})")
        cell_sum_qty.font = font_total
        cell_sum_qty.alignment = align_center
        cell_sum_qty.number_format = "#,##0"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    val = output.getvalue()
    raw_fname = (title or "Imalat_Ozet_Listesi").strip()
    safe_fname = re.sub(r'[^a-zA-Z0-9_\-]', '_', raw_fname)[:40]
    filename = f"{safe_fname}.xlsx"
    
    return Response(
        content=val,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Content-Length": str(len(val)),
            "Cache-Control": "no-cache, no-store, must-revalidate"
        }
    )

@app.get("/api/styles/export-carsaf-excel")
def api_export_carsaf_excel_get(
    ids: Optional[str] = Query(None),
    title: Optional[str] = Query(None),
    authorization: Optional[str] = Header(None)
):
    user = None
    if authorization:
        t = authorization.replace("Bearer ", "").strip()
        user = get_current_user(t)
    company_id = (user.get("company_id") if user and user.get("company_id") is not None else 1)
    
    style_ids = None
    if ids:
        try:
            style_ids = [int(x.strip()) for x in ids.split(",") if x.strip().isdigit()]
        except Exception:
            pass
            
    return generate_carsaf_print_excel_workbook(style_ids, title, company_id)


# ----------------- BUDGET EXCEL EXPORT (ÖRNEK BÜTÇE FORMATI) -----------------

def generate_budget_excel_workbook(style_ids: Optional[List[int]] = None, company_id: int = 1) -> Response:
    conn = get_db()
    c = conn.cursor()
    
    # Safe company_id fallback
    cid = company_id if company_id is not None else 1

    if style_ids and len(style_ids) > 0:
        placeholders = ",".join("?" for _ in style_ids)
        order_by_clause = "CASE s.id " + " ".join(f"WHEN {sid} THEN {i}" for i, sid in enumerate(style_ids)) + " END"
        sql = f"""
        SELECT s.*, 
               o.po_number, o.customer_name, o.brand, o.season, o.delivery_date, o.order_date
        FROM styles s
        JOIN orders o ON s.order_id = o.id
        WHERE (s.company_id = ? OR s.company_id IS NULL) AND s.id IN ({placeholders})
        ORDER BY {order_by_clause}
        """
        params = [cid] + style_ids
    else:
        sql = """
        SELECT s.*, 
               o.po_number, o.customer_name, o.brand, o.season, o.delivery_date, o.order_date
        FROM styles s
        JOIN orders o ON s.order_id = o.id
        WHERE (s.company_id = ? OR s.company_id IS NULL)
        ORDER BY o.id DESC, s.id ASC
        """
        params = [cid]
        
    c.execute(sql, tuple(params))
    styles = [dict(r) for r in c.fetchall()]

    all_style_ids = [s["id"] for s in styles]
    fl_by_style = {}
    if all_style_ids:
        pl_ids = ",".join("?" for _ in all_style_ids)
        c.execute(f"SELECT * FROM fabric_links WHERE style_id IN ({pl_ids}) ORDER BY id ASC", tuple(all_style_ids))
        for fl in c.fetchall():
            sid = fl["style_id"]
            if sid not in fl_by_style:
                fl_by_style[sid] = dict(fl)
    conn.close()

    # FabricTag global cache
    ft_map = {}
    try:
        ft_conn = FabricTagConnector.get_connection()
        if ft_conn:
            ft_c = ft_conn.cursor()
            ft_c.execute("SELECT id, internal_code, company_name, quality_name, quality_code, design_code, color FROM fabrics")
            for f_row in ft_c.fetchall():
                f_d = dict(f_row)
                if f_d.get("id"):
                    ft_map[str(f_d["id"])] = f_d
                if f_d.get("internal_code"):
                    ft_map[f_d["internal_code"].upper()] = f_d
                if f_d.get("quality_code"):
                    ft_map[f_d["quality_code"].upper()] = f_d
            ft_conn.close()
    except Exception:
        pass

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bütçe Tablosu"
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "K2"

    first_cust = (styles[0].get("customer_name") or styles[0].get("brand") or "ÜRETİM") if styles else "ÜRETİM"
    first_season = (styles[0].get("season") or "") if styles else ""
    title_col_1 = f"{first_cust}  {first_season} SMS".strip() or "BÜTÇE LİSTESİ"

    # Headers: Style removed, Kumaş expanded into Kumaşçı, Kalite Adı, Kalite Kodu, Varyant, Renk (49 columns total)
    headers = [
        "Görsel",
        title_col_1,
        "Model - Renk",
        "Renk",
        "Kumaşçı",
        "Kalite Adı",
        "Kalite Kodu",
        "Varyant",
        "Kumaş Rengi",
        "QUANTITY",
        "PRICE ",
        "KUMAŞ 1  FİYAT",
        "1 iç  CM",
        "KUMAŞ 1 TOTAL",
        "KUMAŞ 2 FİYAT",
        "1 iç gramaj ",
        "KUMAŞ 2 TOTAL ",
        "kumaş toplam ",
        "KUMAŞ 3 FİYAT",
        "1 iç gramaj ",
        "KUMAŞ 3 TOTAL ",
        "FASON",
        "KES.İŞT",
        "AKS",
        "ETKOL",
        "LOGO",
        "BRİT",
        "LASTİK",
        "Düğme",
        "Çıtçıt",
        "DIGER",
        "KARGO",
        "NAKLIYE",
        "KUMAŞ ",
        "ASTAR",
        "DİKME",
        "İlik-Düğme",
        "Çıtçıt Çakım",
        "",
        "",
        "",
        "",
        "maliyet TL",
        "maliyet  EURO ",
        "Euro",
        "PRICE",
        "TOPLAM maliyet TL +   6%",
        "TOPLAM maliyet EUR +   6%",
        "Order Toplamı "
    ]

    header_font = Font(name="Segoe UI", size=9.5, bold=True, color="000000")
    header_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="medium", color="475569"),
        bottom=Side(style="medium", color="475569")
    )

    ws.row_dimensions[1].height = 34
    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = header_border

    cell_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    font_normal = Font(name="Segoe UI", size=9, color="1E293B")
    font_bold = Font(name="Segoe UI", size=9, bold=True, color="0F172A")
    font_muted = Font(name="Segoe UI", size=8.5, color="94A3B8")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    last_row = 1
    total_qty = 0
    total_maliyet_tl_sum = 0.0
    total_maliyet_eur_sum = 0.0
    total_order_sum = 0.0
    img_buffers = []

    for idx, s in enumerate(styles, start=2):
        last_row = idx
        ws.row_dimensions[idx].height = 75
        current_fill = fill_white if (idx % 2 == 0) else fill_zebra

        # 0. Görsel (A Sütunu / Column 1)
        img_url = s.get("image_url") or s.get("image_url_2")
        cell_a = ws.cell(row=idx, column=1)
        cell_a.fill = current_fill
        cell_a.border = cell_border
        img_added = False
        if img_url:
            try:
                pil_img = None
                if img_url.startswith("data:image"):
                    hdr, enc = img_url.split(",", 1)
                    pil_img = PILImage.open(io.BytesIO(base64.b64decode(enc)))
                elif img_url.startswith("/uploads/"):
                    p = UPLOADS_DIR / img_url.replace("/uploads/", "").lstrip("/")
                    if p.exists():
                        pil_img = PILImage.open(p)
                elif img_url.startswith("/data/uploads/") or img_url.startswith("data/uploads/"):
                    rel_path = img_url.lstrip("/")
                    p = Path(rel_path)
                    if p.exists():
                        pil_img = PILImage.open(p)
                elif Path(img_url).exists():
                    pil_img = PILImage.open(img_url)
                    
                if pil_img:
                    pil_img = pil_img.convert("RGB")
                    pil_img.thumbnail((105, 88), PILImage.Resampling.LANCZOS)
                    buf = io.BytesIO()
                    pil_img.save(buf, format="JPEG", quality=88)
                    buf.seek(0)
                    img_buffers.append(buf)
                    xl_img = OpenpyxlImage(buf)
                    img_w = pil_img.width
                    img_h = pil_img.height

                    # Hücrede En ve Boydan Ortalama (Horizontal & Vertical Center)
                    col_w_px = 140
                    row_h_px = 100
                    x_off_px = max(0, (col_w_px - img_w) // 2)
                    y_off_px = max(0, (row_h_px - img_h) // 2)

                    marker = AnchorMarker(col=0, colOff=pixels_to_EMU(x_off_px), row=idx - 1, rowOff=pixels_to_EMU(y_off_px))
                    size = XDRPositiveSize2D(pixels_to_EMU(img_w), pixels_to_EMU(img_h))
                    xl_img.anchor = OneCellAnchor(_from=marker, ext=size)
                    ws.add_image(xl_img)
                    img_added = True
            except Exception:
                pass
                
        if not img_added:
            cell_a.value = "[Görsel Yok]"
            cell_a.alignment = align_center
            cell_a.font = font_muted

        cost_data = {}
        if s.get("cost_data_json"):
            try:
                cost_data = json.loads(s["cost_data_json"])
            except Exception:
                cost_data = {}

        def safe_num(val, default=0.0):
            if val is None or val == "":
                return default
            try:
                if isinstance(val, (int, float)):
                    return float(val)
                s_clean = str(val).strip().replace(",", ".")
                return float(s_clean)
            except Exception:
                return default

        rates = cost_data.get("exchangeRates") or {"EUR": 55.75, "USD": 48.9, "TL": 1.0}
        eur_rate = safe_num(rates.get("EUR"), 55.75)

        fabrics = cost_data.get("fabrics") or []
        items = cost_data.get("items") or []

        def get_item_val(keywords):
            for it in items:
                n = (it.get("name") or "").lower()
                if any(k.lower() in n for k in keywords):
                    v = it.get("total_tl")
                    if v is None or v == "":
                        v = it.get("evaluated_price")
                    if v is None or v == "":
                        v = it.get("price")
                    return safe_num(v, 0.0)
            return 0.0

        fb1 = fabrics[0] if len(fabrics) > 0 else {}
        fb1_price = safe_num(fb1.get("price") or s.get("fabric_price_1"))
        fb1_meters = safe_num(fb1.get("meters") or s.get("unit_meters"))
        fb1_grams = safe_num(fb1.get("grams") or s.get("unit_grams"))
        fb1_cm = round(fb1_meters * 100) if fb1.get("unit_type") != "KG" else fb1_grams
        fb1_total = safe_num(fb1.get("total_tl"))

        fb2 = fabrics[1] if len(fabrics) > 1 else {}
        fb2_price = safe_num(fb2.get("price") or s.get("fabric_price_2"))
        fb2_meters = safe_num(fb2.get("meters") or s.get("unit_meters_2"))
        fb2_grams = safe_num(fb2.get("grams") or s.get("unit_grams_2"))
        fb2_gramaj = fb2_grams if fb2_grams > 0 else (fb2_meters * 100)
        fb2_total = safe_num(fb2.get("total_tl"))

        fb3 = fabrics[2] if len(fabrics) > 2 else {}
        fb3_price = safe_num(fb3.get("price"))
        fb3_meters = safe_num(fb3.get("meters"))
        fb3_grams = safe_num(fb3.get("grams"))
        fb3_gramaj = fb3_grams if fb3_grams > 0 else (fb3_meters * 100)
        fb3_total = safe_num(fb3.get("total_tl"))

        kumas_toplam = safe_num(cost_data.get("totals", {}).get("toplam_kumas_tl") or (fb1_total + fb2_total + fb3_total))

        v_fason = get_item_val(["fason", "dikim"])
        v_kesim = get_item_val(["kesim"])
        v_aks = get_item_val(["aksesuar", "aks"])
        v_etkol = get_item_val(["etiket", "kol", "etkol"])
        v_logo = get_item_val(["logo", "baskı", "nakış"])
        v_brit = get_item_val(["brit"])
        v_lastik = get_item_val(["lastik"])
        v_dugme = get_item_val(["düğme", "fermuar"])
        v_citcit = get_item_val(["çıtçıt"])
        v_diger = get_item_val(["diğer", "tasarım", "kalıp"])
        v_kargo = get_item_val(["kargo"])
        v_nakliye = get_item_val(["nakliye"])
        v_kumas_extra = get_item_val(["yıkama", "taş"]) or (fb1_total if not get_item_val(["yıkama", "taş"]) else 0.0)
        v_astar = get_item_val(["astar", "tela"])
        v_dikme = get_item_val(["dikme", "el dikişi"])
        v_ilik_dugme = get_item_val(["ilik", "ilik-düğme"])
        v_citcit_cakim = get_item_val(["çakım", "çıtçıt çakım"])

        totals = cost_data.get("totals") or {}
        maliyet_tl = safe_num(totals.get("maliyet_1_tl") or totals.get("maliyet_2_tl"))
        if maliyet_tl == 0.0 and (kumas_toplam > 0 or v_fason > 0):
            maliyet_tl = round((kumas_toplam + v_fason + v_kesim + v_aks + v_etkol + v_kargo + v_nakliye + v_diger) * 1.06, 2)
            
        maliyet_eur = safe_num(totals.get("maliyet_1_doviz") or (round(maliyet_tl / eur_rate, 2) if eur_rate > 0 else 0.0))
        price = safe_num(totals.get("toplam_satis_doviz") or s.get("unit_price"))
        qty = int(safe_num(s.get("total_quantity")))
        total_qty += qty

        tot_maliyet_tl = round(maliyet_tl * qty, 2)
        tot_maliyet_eur = round(maliyet_eur * qty, 2)
        order_toplam = round(price * qty, 2)

        total_maliyet_tl_sum += tot_maliyet_tl
        total_maliyet_eur_sum += tot_maliyet_eur
        total_order_sum += order_toplam

        style_no = str(s.get("style_no") or "")
        color_code = str(s.get("color_code") or "").strip()
        color_name = str(s.get("color_name") or "").strip()
        model_color = f"{style_no}-{color_code or color_name}"
        cust_season = f"{s.get('customer_name') or ''} {s.get('season') or ''}".strip()
        display_color = (f"{color_code} {color_name}".strip()) if (color_code and color_name and not color_name.startswith(color_code)) else (color_name or color_code)

        # Resolve Fabric Details (Kumaşçı, Kalite Adı, Kalite Kodu, Varyant, Kumaş Rengi)
        fl = fl_by_style.get(s["id"], {})
        f_meta = None
        if fl.get("fabrictag_id") and str(fl["fabrictag_id"]) in ft_map:
            f_meta = ft_map[str(fl["fabrictag_id"])]
        if not f_meta:
            raw_art = str(s.get("fabric_article") or s.get("fabric_type") or "")
            m_elt = re.search(r'ELT\d+', raw_art, re.IGNORECASE)
            if m_elt and m_elt.group(0).upper() in ft_map:
                f_meta = ft_map[m_elt.group(0).upper()]

        def clean_kumas_val(v):
            if not v: return ""
            sv = str(v).strip()
            if re.match(r'^ELT\d+$', sv, re.IGNORECASE):
                return ""
            if sv.upper() in ["KODSUZ", "YOK", "NONE", "NULL", "-", "—", "İSİMSİZ", "ISIMSIZ", "TANIMSIZ", "BİLGİSİ YOK", "BILGISI YOK", "BİLGİ YOK", "BILGI YOK", "BELİRTİLMEDİ", "BELIRTILMEDI"]:
                return ""
            return sv

        kumasci = ""
        kalite_adi = ""
        kalite_kodu = ""
        varyant = ""
        kumas_renk = ""

        if f_meta:
            kumasci = clean_kumas_val(f_meta.get("company_name"))
            kalite_adi = clean_kumas_val(f_meta.get("quality_name"))
            kalite_kodu = clean_kumas_val(f_meta.get("quality_code"))
            varyant = clean_kumas_val(f_meta.get("design_code"))
            kumas_renk = clean_kumas_val(f_meta.get("color") or display_color)
        else:
            raw = str(s.get("fabric_article") or s.get("fabric_type") or "")
            raw = re.sub(r'\bELT\d+\b\s*[-–—:]*\s*', '', raw, flags=re.IGNORECASE).strip()
            parts = [p.strip() for p in raw.split(" - ") if p.strip()]
            if len(parts) >= 2:
                kumasci = clean_kumas_val(parts[0])
                rest = parts[1]
                if "(" in rest and ")" in rest:
                    kalite_adi = clean_kumas_val(rest.split("(")[0].strip())
                    kalite_kodu = clean_kumas_val(rest.split("(")[1].split(")")[0].strip())
                else:
                    kalite_adi = clean_kumas_val(rest)
            elif len(parts) == 1:
                kumasci = clean_kumas_val(parts[0])
            varyant = clean_kumas_val(s.get("fabric_variant_1"))
            kumas_renk = clean_kumas_val(display_color)

        row_data = [
            cust_season,
            model_color,
            display_color,
            kumasci,
            kalite_adi,
            kalite_kodu,
            varyant,
            kumas_renk,
            qty if qty else None,
            price if price else None,
            fb1_price if fb1_price else 0,
            fb1_cm if fb1_cm else 0,
            fb1_total if fb1_total else 0,
            fb2_price if fb2_price else 0,
            fb2_gramaj if fb2_gramaj else 0,
            fb2_total if fb2_total else 0,
            kumas_toplam if kumas_toplam else 0,
            fb3_price if fb3_price else None,
            fb3_gramaj if fb3_gramaj else None,
            fb3_total if fb3_total else None,
            v_fason if v_fason else 0,
            v_kesim if v_kesim else 0,
            v_aks if v_aks else 0,
            v_etkol if v_etkol else 0,
            v_logo if v_logo else 0,
            v_brit if v_brit else 0,
            v_lastik if v_lastik else 0,
            v_dugme if v_dugme else 0,
            v_citcit if v_citcit else 0,
            v_diger if v_diger else 0,
            v_kargo if v_kargo else 0,
            v_nakliye if v_nakliye else 0,
            v_kumas_extra if v_kumas_extra else 0,
            v_astar if v_astar else 0,
            v_dikme if v_dikme else 0,
            v_ilik_dugme if v_ilik_dugme else 0,
            v_citcit_cakim if v_citcit_cakim else 0,
            None,
            None,
            None,
            None,
            maliyet_tl if maliyet_tl else 0,
            maliyet_eur if maliyet_eur else 0,
            eur_rate,
            price if price else 0,
            tot_maliyet_tl if tot_maliyet_tl else 0,
            tot_maliyet_eur if tot_maliyet_eur else 0,
            order_toplam if order_toplam else 0
        ]

        # Populate columns 2 to 49 (column 1 is Görsel)
        for c_idx_offset, val in enumerate(row_data, start=2):
            c_cell = ws.cell(row=idx, column=c_idx_offset, value=val)
            c_cell.fill = current_fill
            c_cell.border = cell_border
            c_cell.font = font_normal

            if c_idx_offset in [2, 3, 5, 6]:
                c_cell.alignment = align_left
            elif c_idx_offset in [4, 7, 8, 9]:
                c_cell.alignment = align_center
                c_cell.font = font_bold
            elif c_idx_offset == 10:  # QUANTITY
                c_cell.alignment = align_center
                c_cell.font = font_bold
                c_cell.number_format = "#,##0"
            elif c_idx_offset in [11, 46]:  # PRICE
                c_cell.alignment = align_right
                c_cell.font = font_bold
                _, b_prc_fmt = get_excel_currency_format(s.get("currency") or "EUR")
                c_cell.number_format = b_prc_fmt
            elif c_idx_offset in [45]:  # EUR RATE
                c_cell.alignment = align_center
                c_cell.number_format = "#,##0.00"
            elif c_idx_offset in [43, 47]:  # Maliyet TL ve TOPLAM maliyet TL + 6%
                c_cell.alignment = align_right
                c_cell.font = font_bold
                c_cell.number_format = '"₺" #,##0.00'
            elif c_idx_offset in [44, 48, 49]:  # Maliyet EUR, TOPLAM EUR, Order Toplamı
                c_cell.alignment = align_right
                c_cell.font = font_bold
                c_cell.number_format = '"€" #,##0.00'
            elif isinstance(val, (int, float)):
                c_cell.alignment = align_right
                c_cell.number_format = "#,##0.00"

    # Summary Total Row
    tot_row = last_row + 1
    ws.row_dimensions[tot_row].height = 28
    tot_fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    tot_font = Font(name="Segoe UI", size=10, bold=True, color="000000")
    
    ws.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=9)
    cell_tot_lbl = ws.cell(row=tot_row, column=1, value="GENEL TOPLAM:")
    cell_tot_lbl.font = tot_font
    cell_tot_lbl.alignment = Alignment(horizontal="right", vertical="center")
    
    c_tot_qty = ws.cell(row=tot_row, column=10, value=total_qty)
    c_tot_qty.font = tot_font
    c_tot_qty.alignment = align_center
    c_tot_qty.number_format = "#,##0"
    
    c_tot_tl = ws.cell(row=tot_row, column=47, value=total_maliyet_tl_sum)
    c_tot_tl.font = tot_font
    c_tot_tl.alignment = align_right
    c_tot_tl.number_format = '"₺" #,##0.00'
    
    c_tot_eur = ws.cell(row=tot_row, column=48, value=total_maliyet_eur_sum)
    c_tot_eur.font = tot_font
    c_tot_eur.alignment = align_right
    c_tot_eur.number_format = '"€" #,##0.00'
    
    c_tot_order = ws.cell(row=tot_row, column=49, value=total_order_sum)
    c_tot_order.font = tot_font
    c_tot_order.alignment = align_right
    c_tot_order.number_format = '"€" #,##0.00'
    
    for c_i in range(1, 50):
        c_tot = ws.cell(row=tot_row, column=c_i)
        c_tot.fill = tot_fill
        c_tot.border = Border(top=Side(style="medium", color="475569"), bottom=Side(style="double", color="475569"))

    ws.column_dimensions["A"].width = 18  # Görsel
    ws.column_dimensions["B"].width = 24  # Müşteri / Sezon
    ws.column_dimensions["C"].width = 20  # Model - Renk
    ws.column_dimensions["D"].width = 16  # Renk
    ws.column_dimensions["E"].width = 22  # Kumaşçı
    ws.column_dimensions["F"].width = 22  # Kalite Adı
    ws.column_dimensions["G"].width = 16  # Kalite Kodu
    ws.column_dimensions["H"].width = 16  # Varyant
    ws.column_dimensions["I"].width = 18  # Kumaş Rengi
    ws.column_dimensions["J"].width = 14  # Quantity
    ws.column_dimensions["K"].width = 14  # Price
    ws.column_dimensions["R"].width = 16  # Kumaş Toplam
    ws.column_dimensions["AQ"].width = 16 # Maliyet TL
    ws.column_dimensions["AR"].width = 16 # Maliyet EUR
    ws.column_dimensions["AS"].width = 12 # Euro Kuru
    ws.column_dimensions["AT"].width = 14 # Price
    ws.column_dimensions["AU"].width = 20 # Toplam Maliyet TL +6%
    ws.column_dimensions["AV"].width = 20 # Toplam Maliyet EUR +6%
    ws.column_dimensions["AW"].width = 18 # Order Toplamı

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    
    filename = f"Butce_Tablosu_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"}
    )

@app.get("/api/styles/export-budget-excel")
def api_export_budget_excel_get(ids: Optional[str] = Query(None), authorization: Optional[str] = Header(None)):
    user = None
    if authorization:
        t = authorization.replace("Bearer ", "").strip()
        user = get_current_user(t)
    company_id = (user.get("company_id") if user and user.get("company_id") is not None else 1)
    
    style_ids = None
    if ids:
        try:
            style_ids = [int(x.strip()) for x in ids.split(",") if x.strip().isdigit()]
        except Exception:
            pass
            
    return generate_budget_excel_workbook(style_ids, company_id)

@app.post("/api/styles/export-budget-excel")
def api_export_budget_excel_post(req: ExportExcelRequest, authorization: Optional[str] = Header(None)):
    user = None
    if authorization:
        t = authorization.replace("Bearer ", "").strip()
        user = get_current_user(t)
    company_id = (user.get("company_id") if user and user.get("company_id") is not None else 1)
    
    return generate_budget_excel_workbook(req.style_ids, company_id)



# ----------------- SINGLE SIZE INLINE UPDATE -----------------

class UpdateSingleSizeRequest(BaseModel):
    style_id: int
    size_name: str
    quantity: int
    reason: Optional[str] = "Hızlı Beden Revizyonu"

@app.post("/api/styles/update-single-size")
def api_update_single_size(req: UpdateSingleSizeRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()

    c.execute("SELECT total_quantity, order_id, style_no, color_name FROM styles WHERE id = ?", (req.style_id,))
    st = c.fetchone()
    if not st:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    
    old_total = st[0] or 0
    order_id = st[1]
    style_no = st[2]
    color_name = st[3]

    c.execute("SELECT id, quantity FROM size_distributions WHERE style_id = ? AND size_name = ?", (req.style_id, req.size_name))
    row = c.fetchone()
    old_size_qty = row[1] if row else 0

    if row:
        c.execute("UPDATE size_distributions SET quantity = ? WHERE id = ?", (req.quantity, row[0]))
    else:
        c.execute("INSERT INTO size_distributions (style_id, size_name, quantity) VALUES (?, ?, ?)", (req.style_id, req.size_name, req.quantity))

    c.execute("SELECT SUM(quantity) FROM size_distributions WHERE style_id = ?", (req.style_id,))
    new_total = c.fetchone()[0] or 0

    c.execute("UPDATE styles SET total_quantity = ?, total_amount = ? * unit_price, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (new_total, new_total, req.style_id))

    c.execute("""
    UPDATE orders 
    SET total_quantity = (SELECT SUM(total_quantity) FROM styles WHERE order_id = ?),
        total_amount = (SELECT SUM(total_quantity * unit_price) FROM styles WHERE order_id = ?)
    WHERE id = ?
    """, (order_id, order_id, order_id))

    user_name = user.get("full_name") or user.get("username")
    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (req.style_id, order_id, user.get("user_id"), user_name, f"size_{req.size_name}", f"{req.size_name}: {old_size_qty}", f"{req.size_name}: {req.quantity}", f"{req.reason} (Toplam: {old_total} -> {new_total})"))

    c.execute("""
    INSERT INTO production_status_logs (style_id, user_name, category, title, message)
    VALUES (?, ?, 'Beden Revizyonu', ?, ?)
    """, (req.style_id, user_name, f"Beden {req.size_name} Değiştirildi", f"{req.size_name}: {old_size_qty} -> {req.quantity} Adet (Yeni Toplam: {new_total})"))

    conn.commit()
    conn.close()
    return {"status": "success", "new_total": new_total, "message": f"{req.size_name} bedeni güncellendi. Yeni toplam: {new_total}"}

# ----------------- DUAL IMAGE UPLOAD (SLOT 1 & SLOT 2) -----------------

@app.post("/api/styles/{style_id}/image")
async def api_upload_style_image(
    style_id: int, 
    slot: int = Query(1),
    file: UploadFile = File(...), 
    user: Dict[str, Any] = Depends(require_user)
):
    ext = Path(file.filename).suffix.lower()
    if ext not in [".jpg", ".jpeg", ".png", ".webp", ".gif"]:
        raise HTTPException(status_code=400, detail="Geçersiz görsel formatı. Lütfen JPG veya PNG yükleyin.")
    
    filename = f"style_{style_id}_s{slot}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
    dest_path = STYLE_IMAGES_DIR / filename
    
    with open(dest_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
        
    rel_url = f"/uploads/styles/{filename}"
    
    col_name = "image_url_2" if slot == 2 else "image_url"
    
    conn = get_db()
    c = conn.cursor()
    c.execute(f"UPDATE styles SET {col_name} = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (rel_url, style_id))
    
    user_name = user.get("full_name") or user.get("username")
    c.execute("""
    INSERT INTO audit_logs (style_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, 'Eski Görsel', ?, ?)
    """, (style_id, user.get("user_id"), user_name, col_name, rel_url, f"Model Fotoğrafı Yüklendi (Görsel {slot})"))

    conn.commit()
    conn.close()
    
    return {"status": "success", "image_url": rel_url, "slot": slot}

@app.delete("/api/styles/{style_id}/image/{slot}")
def api_delete_style_image(style_id: int, slot: int, user: Dict[str, Any] = Depends(require_user)):
    col_name = "image_url_2" if slot == 2 else "image_url"
    conn = get_db()
    c = conn.cursor()
    c.execute(f"UPDATE styles SET {col_name} = NULL, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (style_id,))
    conn.commit()
    conn.close()
    return {"status": "success", "message": f"Görsel {slot} silindi."}

@app.post("/api/styles/{style_id}/swap-images")
def api_swap_style_images(style_id: int, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT image_url, image_url_2 FROM styles WHERE id = ?", (style_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    img1, img2 = row[0], row[1]
    c.execute("UPDATE styles SET image_url = ?, image_url_2 = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (img2, img1, style_id))
    user_name = user.get("full_name") or user.get("username")
    c.execute("""
    INSERT INTO audit_logs (style_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, 'image_url', ?, ?, 'Görsel 1 ve 2 Yer Değiştirildi')
    """, (style_id, user.get("user_id"), user_name, str(img1), str(img2)))
    conn.commit()
    conn.close()
    return {"status": "success", "image_url": img2, "image_url_2": img1, "message": "Görseller yer değiştirildi."}


# ----------------- UPDATE CELL WITH AUDIT LOG & TIMESTAMPS -----------------

class StyleUpdateFieldRequest(BaseModel):
    style_id: int
    field: str
    value: Any
    reason: Optional[str] = "Kullanıcı Düzenlemesi"

@app.post("/api/styles/update-cell")
def api_update_style_cell(req: StyleUpdateFieldRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()

    if req.field in ["order_date", "delivery_date", "season"]:
        c.execute("SELECT order_id FROM styles WHERE id = ?", (req.style_id,))
        order_id = c.fetchone()[0]
        c.execute(f"UPDATE orders SET {req.field} = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (req.value, order_id))
        user_name = user.get("full_name") or user.get("username")
        c.execute("""
        INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
        VALUES (?, ?, ?, ?, ?, '', ?, ?)
        """, (req.style_id, order_id, user.get("user_id"), user_name, req.field, str(req.value), req.reason))
        conn.commit()
        conn.close()
        return {"status": "success", "message": f"{req.field} güncellendi"}

    c.execute(f"SELECT {req.field}, order_id FROM styles WHERE id = ?", (req.style_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    
    old_val = str(row[0] or "")
    order_id = row[1]
    new_val = str(req.value or "")

    user_name = user.get("full_name") or user.get("username")
    now_date_str = datetime.now().strftime('%d.%m.%Y')
    now_time_str = datetime.now().strftime('%d.%m.%Y %H:%M')

    extra_sql = ""
    extra_params = []

    if req.field == "fabric_ordered":
        val_int = int(req.value or 0)
        date_val = now_date_str if val_int == 1 else None
        extra_sql = ", fabric_ordered_date = ?"
        extra_params.append(date_val)
    elif req.field == "pps_sent":
        str_val = str(req.value or "").strip()
        date_val = now_date_str if str_val and str_val != "-" and str_val != "0" else None
        extra_sql = ", pps_sent_date = ?"
        extra_params.append(date_val)
    elif req.field == "shipping_sample_sent":
        val_int = int(req.value or 0)
        date_val = now_date_str if val_int == 1 else None
        extra_sql = ", shipping_sample_sent_date = ?"
        extra_params.append(date_val)
    elif req.field == "status":
        extra_sql = ", last_status_updated_by = ?, last_status_updated_at = ?"
        extra_params.extend([user_name, now_time_str])

    sql = f"UPDATE styles SET {req.field} = ? {extra_sql}, updated_at = CURRENT_TIMESTAMP WHERE id = ?"
    c.execute(sql, [req.value] + extra_params + [req.style_id])
    
    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (req.style_id, order_id, user.get("user_id"), user_name, req.field, old_val, new_val, req.reason))

    if req.field == "unit_price":
        c.execute("""
        UPDATE orders 
        SET total_amount = (SELECT SUM(total_quantity * unit_price) FROM styles WHERE order_id = ?)
        WHERE id = ?
        """, (order_id, order_id))
    elif req.field == "fabric_received_meters":
        rec_meters = float(req.value or 0)
        if rec_meters > 0:
            log_msg = f"{now_date_str} tarihinde {rec_meters} metre kumaş geldi."
            c.execute("""
            INSERT INTO production_status_logs (style_id, user_name, category, title, message)
            VALUES (?, ?, 'Kumaş Girişi', 'Kumaş Depoya Geldi', ?)
            """, (req.style_id, user_name, log_msg))

    c.execute("""
    INSERT INTO production_status_logs (style_id, user_name, category, title, message)
    VALUES (?, ?, 'İşlem Aşaması', ?, ?)
    """, (req.style_id, user_name, f"'{req.field}' Değiştirildi", f"{old_val} ➜ {new_val} ({req.reason})"))


    conn.commit()
    conn.close()
    return {
        "status": "success", 
        "message": f"{req.field} güncellendi",
        "user_name": user_name,
        "date_str": date_val if 'date_val' in locals() else now_date_str
    }


# ----------------- KESİMHANE & PASTAL GÜNCELLEME -----------------

class CuttingSizeItem(BaseModel):
    size_name: str
    cut_quantity: int = 0

class CuttingUpdateRequest(BaseModel):
    style_id: int
    cut_meters: Optional[float] = 0.0
    sizes: List[CuttingSizeItem] = []

@app.post("/api/cutting/update")
def api_update_cutting(req: CuttingUpdateRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()

    c.execute("SELECT id, order_id, style_no, total_quantity, cut_meters, actual_unit_meters, cutting_date FROM styles WHERE id = ?", (req.style_id,))
    style_row = c.fetchone()
    if not style_row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    
    order_id = style_row["order_id"]
    user_name = user.get("full_name") or user.get("username") or "Kesim Şefi"
    now_str = datetime.now().strftime("%d.%m.%Y %H:%M")

    total_cut_qty = 0
    for sz in req.sizes:
        c.execute("""
        UPDATE size_distributions 
        SET cut_quantity = ? 
        WHERE style_id = ? AND size_name = ?
        """, (int(sz.cut_quantity or 0), req.style_id, sz.size_name))
        total_cut_qty += int(sz.cut_quantity or 0)

    cut_m = float(req.cut_meters or 0.0)
    actual_unit_m = round(cut_m / total_cut_qty, 3) if total_cut_qty > 0 and cut_m > 0 else 0.0

    c.execute("""
    UPDATE styles 
    SET cut_meters = ?, actual_unit_meters = ?, cutting_date = ?, updated_at = CURRENT_TIMESTAMP 
    WHERE id = ?
    """, (cut_m, actual_unit_m, now_str, req.style_id))

    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, 'cutting_update', ?, ?, 'Kesimhane Beden & Metraj Güncellemesi')
    """, (req.style_id, order_id, user.get("user_id"), user_name, 
          f"M:{style_row['cut_meters']}", f"M:{cut_m}, Adet:{total_cut_qty}, Birim:{actual_unit_m}, Tarih:{now_str}"))

    conn.commit()
    conn.close()

    return {
        "status": "success",
        "style_id": req.style_id,
        "total_cut_quantity": total_cut_qty,
        "cut_meters": cut_m,
        "actual_unit_meters": actual_unit_m,
        "cutting_date": now_str
    }


# ----------------- KESİM FİŞİ EXCEL DIŞA AKTARIM -----------------

class CuttingSlipColorItem(BaseModel):
    name: str
    channel: Optional[str] = ""
    quantities: List[Any] = []
    total: Optional[int] = 0

class CuttingSlipExportRequest(BaseModel):
    customer_name: Optional[str] = ""
    brand: Optional[str] = ""
    style_no: Optional[str] = ""
    channel: Optional[str] = ""
    date: Optional[str] = ""
    image_url: Optional[str] = ""
    sizes: List[str] = []
    has_channel_info: Optional[bool] = False
    colors: List[CuttingSlipColorItem] = []

@app.post("/api/cutting/export-slip-excel")
def api_export_cutting_slip_excel(req: CuttingSlipExportRequest):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sayfa1"

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    thin = Side(border_style="thin", color="000000")
    medium = Side(border_style="medium", color="000000")
    box_border = Border(top=thin, left=thin, right=thin, bottom=thin)

    has_channel = bool(req.has_channel_info or any(bool(c.channel and str(c.channel).strip()) for c in req.colors))
    num_sizes = len(req.sizes)

    # Kolon yerleşimi:
    # Boş / adetsiz sütunlar kaldırıldı.
    # Kanal varsa: Col 1: RENK/VARYANT, Col 2: KANAL, Col 3..(2+num_sizes): Bedenler, Son Kolon: TOPLAM
    # Kanal yoksa: Col 1: RENK/VARYANT, Col 2..(1+num_sizes): Bedenler, Son Kolon: TOPLAM
    if has_channel:
        total_cols = max(4, 1 + 1 + num_sizes + 1)
        size_col_letters = [get_column_letter(3 + idx) for idx in range(num_sizes)]
        ws.column_dimensions['A'].width = 22
        ws.column_dimensions['B'].width = 16
        for sc in size_col_letters:
            ws.column_dimensions[sc].width = 12
    else:
        total_cols = max(4, 1 + num_sizes + 1)
        size_col_letters = [get_column_letter(2 + idx) for idx in range(num_sizes)]
        ws.column_dimensions['A'].width = 24
        for sc in size_col_letters:
            ws.column_dimensions[sc].width = 12

    last_col_letter = get_column_letter(total_cols)
    ws.column_dimensions[last_col_letter].width = 15

    # Row 1: KESİM FİŞİ
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A1:{last_col_letter}1")
    cell_a1 = ws["A1"]
    cell_a1.value = "KESİM FİŞİ"
    cell_a1.font = Font(name="Arial", size=20, bold=True)
    cell_a1.alignment = Alignment(horizontal="center", vertical="center")

    # Row 2: MÜŞTERİ & TARİH
    ws.row_dimensions[2].height = 25
    ws["A2"].value = "MÜŞTERİ:"
    ws["A2"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"].alignment = Alignment(horizontal="right", vertical="center")

    second_last_col_letter = get_column_letter(max(2, total_cols - 1))
    if total_cols > 4:
        ws.merge_cells(f"B2:{get_column_letter(total_cols - 2)}2")
    ws["B2"].value = req.customer_name or ""
    ws["B2"].font = Font(name="Arial", size=14, bold=False)
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center")

    for ci in range(1, total_cols + 1):
        ws[f"{get_column_letter(ci)}2"].border = box_border

    ws[f"{second_last_col_letter}2"].value = "TARİH"
    ws[f"{second_last_col_letter}2"].font = Font(name="Arial", size=14, bold=True)
    ws[f"{second_last_col_letter}2"].alignment = Alignment(horizontal="center", vertical="center")

    ws[f"{last_col_letter}2"].value = req.date or datetime.now().strftime("%d.%m.%Y")
    ws[f"{last_col_letter}2"].font = Font(name="Arial", size=14, bold=False)
    ws[f"{last_col_letter}2"].alignment = Alignment(horizontal="center", vertical="center")

    # Row 3: MARKA
    ws.row_dimensions[3].height = 25
    ws["A3"].value = "MARKA:"
    ws["A3"].font = Font(name="Arial", size=14, bold=True)
    ws["A3"].alignment = Alignment(horizontal="right", vertical="center")

    brand_val = req.brand or ""
    if req.channel:
        brand_val = f"{brand_val} (Kanal: {req.channel})" if brand_val else f"Kanal: {req.channel}"
    if total_cols > 2:
        ws.merge_cells(f"B3:{last_col_letter}3")
    ws["B3"].value = brand_val
    ws["B3"].font = Font(name="Arial", size=14, bold=False)
    ws["B3"].alignment = Alignment(horizontal="left", vertical="center")

    for ci in range(1, total_cols + 1):
        ws[f"{get_column_letter(ci)}3"].border = box_border

    # Rows 4 to 8: Image box
    ws.merge_cells("A4:A8")
    for r in range(4, 9):
        ws.row_dimensions[r].height = 26
        for ci in range(1, total_cols + 1):
            ws[f"{get_column_letter(ci)}{r}"].border = box_border

    img_added = False
    if req.image_url:
        clean_img_url = req.image_url.split("?")[0].lstrip("/")
        possible_paths = [
            BASE_DIR / clean_img_url,
            BASE_DIR / "data" / clean_img_url,
            BASE_DIR / "data" / "uploads" / clean_img_url.replace("uploads/", ""),
            BASE_DIR / clean_img_url.replace("uploads/", "data/uploads/")
        ]
        for p in possible_paths:
            if p.exists() and p.is_file():
                try:
                    openpyxl_img = OpenpyxlImage(str(p))
                    with PILImage.open(str(p)) as pil_img:
                        orig_w, orig_h = pil_img.size
                        ratio = min(150.0 / max(1, orig_w), 125.0 / max(1, orig_h))
                        openpyxl_img.width = int(orig_w * ratio)
                        openpyxl_img.height = int(orig_h * ratio)
                    ws.add_image(openpyxl_img, "A4")
                    img_added = True
                    break
                except Exception as ex:
                    print("Error adding image to excel:", ex)
    if not img_added:
        ws["A4"].value = "[ RESİM ]"
        ws["A4"].font = Font(name="Arial", size=12, color="888888")
        ws["A4"].alignment = Alignment(horizontal="center", vertical="center")

    # Row 4: ARTICLE
    ws["B4"].value = "ARTICLE     :"
    ws["B4"].font = Font(name="Arial", size=14, bold=True)
    ws["B4"].alignment = Alignment(horizontal="left", vertical="center")
    
    if total_cols > 2:
        ws.merge_cells(f"C4:{last_col_letter}4")
        ws["C4"].value = req.style_no or ""
        ws["C4"].font = Font(name="Calibri", size=18, bold=True)
        ws["C4"].alignment = Alignment(horizontal="left", vertical="center")
    else:
        ws["B4"].value = f"ARTICLE: {req.style_no or ''}"

    # Row 5: KUMAŞ KG
    ws["B5"].value = "KUMAŞ KG :"
    ws["B5"].font = Font(name="Arial", size=14, bold=False)
    ws["B5"].alignment = Alignment(horizontal="left", vertical="center")
    if total_cols > 2:
        ws.merge_cells(f"C5:{last_col_letter}5")

    # Row 6: METRAJ
    ws["B6"].value = "METRAJ      :"
    ws["B6"].font = Font(name="Arial", size=14, bold=False)
    ws["B6"].alignment = Alignment(horizontal="left", vertical="center")
    if total_cols > 2:
        ws.merge_cells(f"C6:{last_col_letter}6")

    # Row 7: EN
    ws["B7"].value = "EN :"
    ws["B7"].font = Font(name="Arial", size=14, bold=False)
    ws["B7"].alignment = Alignment(horizontal="left", vertical="center")
    if total_cols > 2:
        ws.merge_cells(f"C7:{last_col_letter}7")

    # Row 8: GRAMAJ
    ws["B8"].value = "GRAMAJ:"
    ws["B8"].font = Font(name="Arial", size=14, bold=False)
    ws["B8"].alignment = Alignment(horizontal="left", vertical="center")
    if total_cols > 2:
        ws.merge_cells(f"C8:{last_col_letter}8")

    # Row 9: spacer
    ws.row_dimensions[9].height = 12

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    # Row 10: Size Headers (KOMPLE SARI ZEMİN #FFFF00)
    ws.row_dimensions[10].height = 28
    for ci in range(1, total_cols + 1):
        c_letter = get_column_letter(ci)
        ws[f"{c_letter}10"].fill = yellow_fill
        ws[f"{c_letter}10"].border = Border(
            top=medium,
            bottom=thin,
            left=medium if ci == 1 else thin,
            right=medium if ci == total_cols else thin
        )

    ws["A10"].value = "RENK / VARYANT"
    ws["A10"].font = Font(name="Arial", size=11, bold=True)
    ws["A10"].alignment = Alignment(horizontal="center", vertical="center")

    if has_channel:
        ws["B10"].value = "KANAL"
        ws["B10"].font = Font(name="Arial", size=12, bold=True)
        ws["B10"].alignment = Alignment(horizontal="center", vertical="center")

    for sc, sz in zip(size_col_letters, req.sizes):
        ws[f"{sc}10"].value = sz
        ws[f"{sc}10"].font = Font(name="Arial", size=16, bold=True)
        ws[f"{sc}10"].alignment = Alignment(horizontal="center", vertical="center")

    ws[f"{last_col_letter}10"].value = "TOPLAM"
    ws[f"{last_col_letter}10"].font = Font(name="Arial", size=16, bold=True)
    ws[f"{last_col_letter}10"].alignment = Alignment(horizontal="center", vertical="center")

    # Dynamic row heights based on color count so 9 colors fit on 1 A4 page
    num_colors = max(1, len(req.colors))
    if num_colors >= 7:
        row_h_qty = 22
        row_h_empty = 20
    elif num_colors >= 4:
        row_h_qty = 26
        row_h_empty = 24
    else:
        row_h_qty = 30
        row_h_empty = 28

    cur_row = 11
    for c_idx, col_info in enumerate(req.colors):
        # 1. SATIR: Adet satırı
        ws.row_dimensions[cur_row].height = row_h_qty
        top_border_style = thin if c_idx == 0 else medium
        for ci in range(1, total_cols + 1):
            c_letter = get_column_letter(ci)
            ws[f"{c_letter}{cur_row}"].border = Border(
                top=top_border_style,
                bottom=thin,
                left=medium if ci == 1 else thin,
                right=medium if ci == total_cols else thin
            )

        ws[f"A{cur_row}"].value = col_info.name
        ws[f"A{cur_row}"].font = Font(name="Arial", size=10 if num_colors >= 7 else 11, bold=True)
        ws[f"A{cur_row}"].alignment = Alignment(horizontal="center", vertical="center")
        
        if has_channel:
            ws[f"B{cur_row}"].value = col_info.channel or ""
            ws[f"B{cur_row}"].font = Font(name="Arial", size=11, bold=True)
            ws[f"B{cur_row}"].alignment = Alignment(horizontal="center", vertical="center")

        for sc, q in zip(size_col_letters, col_info.quantities):
            val = int(q) if (q is not None and str(q).isdigit()) else (q if q else "")
            ws[f"{sc}{cur_row}"].value = val
            ws[f"{sc}{cur_row}"].font = Font(name="Arial", size=13 if num_colors >= 7 else 16, bold=False)
            ws[f"{sc}{cur_row}"].alignment = Alignment(horizontal="center", vertical="center")
        
        ws[f"{last_col_letter}{cur_row}"].value = col_info.total
        ws[f"{last_col_letter}{cur_row}"].font = Font(name="Arial", size=13 if num_colors >= 7 else 16, bold=True)
        ws[f"{last_col_letter}{cur_row}"].alignment = Alignment(horizontal="center", vertical="center")
        
        # 2. SATIR: 1. Extra Boş Kutucuk Satırı
        cur_row += 1
        ws.row_dimensions[cur_row].height = row_h_empty
        for ci in range(1, total_cols + 1):
            c_letter = get_column_letter(ci)
            ws[f"{c_letter}{cur_row}"].border = Border(
                top=thin,
                bottom=thin,
                left=medium if ci == 1 else thin,
                right=medium if ci == total_cols else thin
            )

        # 3. SATIR: 2. Extra Boş Kutucuk Satırı (HER RENK BİTİMİNDE ALT ÇİZGİSİ BOLD / MEDIUM)
        cur_row += 1
        ws.row_dimensions[cur_row].height = row_h_empty
        for ci in range(1, total_cols + 1):
            c_letter = get_column_letter(ci)
            ws[f"{c_letter}{cur_row}"].border = Border(
                top=thin,
                bottom=medium,
                left=medium if ci == 1 else thin,
                right=medium if ci == total_cols else thin
            )
        
        # Renk bitiminden sonra doğrudan sonraki satıra geç (aradaki boşluk kaldırıldı)
        cur_row += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    clean_style = re.sub(r'[^a-zA-Z0-9_-]', '_', req.style_no or 'Model')
    filename = f"KESIM_FISI_{clean_style}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


class ShippingSizeItem(BaseModel):
    size_name: str
    shipped_quantity: int = 0

class ShippingUpdateRequest(BaseModel):
    style_id: int
    sizes: List[ShippingSizeItem] = []

@app.post("/api/shipping/update")
def api_update_shipping(req: ShippingUpdateRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()

    c.execute("SELECT id, order_id, style_no, total_quantity, shipping_date FROM styles WHERE id = ?", (req.style_id,))
    style_row = c.fetchone()
    if not style_row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    
    order_id = style_row["order_id"]
    user_name = user.get("full_name") or user.get("username") or "Sevkiyat Sorumlusu"
    now_str = datetime.now().strftime("%d.%m.%Y %H:%M")

    total_shipped_qty = 0
    for sz in req.sizes:
        c.execute("""
        UPDATE size_distributions 
        SET shipped_quantity = ? 
        WHERE style_id = ? AND size_name = ?
        """, (int(sz.shipped_quantity or 0), req.style_id, sz.size_name))
        total_shipped_qty += int(sz.shipped_quantity or 0)

    c.execute("""
    UPDATE styles 
    SET shipping_date = ?, updated_at = CURRENT_TIMESTAMP 
    WHERE id = ?
    """, (now_str, req.style_id))

    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, 'shipping_update', ?, ?, 'Yükleme Adetleri Güncellendi')
    """, (req.style_id, order_id, user.get("user_id"), user_name, 
          f"Tarih:{style_row['shipping_date']}", f"Adet:{total_shipped_qty}, Tarih:{now_str}"))

    conn.commit()
    conn.close()

    return {
        "status": "success",
        "style_id": req.style_id,
        "total_shipped_quantity": total_shipped_qty,
        "shipping_date": now_str
    }


# ----------------- FABRIC ASSIGNMENT -----------------

class AssignFabricRequest(BaseModel):
    style_id: int
    fabrictag_id: Optional[int] = None
    fabric_name: str
    composition: str
    supplier: Optional[str] = ""
    fabric_slot: Optional[int] = 1

@app.post("/api/styles/assign-fabric")
def api_assign_fabric(req: AssignFabricRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    slot = req.fabric_slot or 1
    
    c.execute("SELECT fabric_article, fabric_composition, fabric_article_2, fabric_composition_2, order_id FROM styles WHERE id = ?", (req.style_id,))
    row = c.fetchone()
    order_id = row[4] if row else None
    
    if slot == 2:
        old_article = (row[2] or "") if row else ""
        old_comp = (row[3] or "") if row else ""
        c.execute("""
        UPDATE styles 
        SET fabric_article_2 = ?, fabric_composition_2 = ?, updated_at = CURRENT_TIMESTAMP 
        WHERE id = ?
        """, (req.fabric_name, req.composition, req.style_id))
    else:
        old_article = (row[0] or "") if row else ""
        old_comp = (row[1] or "") if row else ""
        c.execute("""
        UPDATE styles 
        SET fabric_article = ?, fabric_composition = ?, updated_at = CURRENT_TIMESTAMP 
        WHERE id = ?
        """, (req.fabric_name, req.composition, req.style_id))

    c.execute("DELETE FROM fabric_links WHERE style_id = ? AND (fabric_slot = ? OR (? = 1 AND fabric_slot IS NULL))", (req.style_id, slot, slot))
    c.execute("""
    INSERT INTO fabric_links (style_id, fabrictag_id, fabric_name, composition, supplier, fabric_slot, status)
    VALUES (?, ?, ?, ?, ?, ?, 'Atandı')
    """, (req.style_id, req.fabrictag_id, req.fabric_name, req.composition, req.supplier, slot))

    user_name = user.get("full_name") or user.get("username")
    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, 'fabric_assignment', ?, ?, ?)
    """, (req.style_id, order_id, user.get("user_id"), user_name, f"{old_article} ({old_comp})", f"{req.fabric_name} ({req.composition})", f"FabricTag Kumaş {slot} Ataması Yapıldı"))


    c.execute("""
    INSERT INTO production_status_logs (style_id, user_name, category, title, message)
    VALUES (?, ?, 'Kumaş', 'Kumaş Değiştirildi / Atandı', ?)
    """, (req.style_id, user_name, f"Yeni Kumaş: {req.fabric_name} | {req.composition} ({req.supplier})"))

    conn.commit()
    conn.close()
    return {"status": "success", "message": "Kumaş başarıyla atandı"}

class DetachFabricRequest(BaseModel):
    style_id: int
    fabric_slot: Optional[int] = 1

@app.post("/api/styles/detach-fabric")
def api_detach_fabric(req: DetachFabricRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    slot = req.fabric_slot or 1
    
    c.execute("SELECT fabric_article, fabric_composition, fabric_article_2, fabric_composition_2, order_id FROM styles WHERE id = ?", (req.style_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    
    order_id = row[4]
    if slot == 2:
        old_article = row[2] or ""
        old_comp = row[3] or ""
        c.execute("""
        UPDATE styles 
        SET fabric_article_2 = '', fabric_composition_2 = '', fabric_type_2 = '', updated_at = CURRENT_TIMESTAMP 
        WHERE id = ?
        """, (req.style_id,))
    else:
        old_article = row[0] or ""
        old_comp = row[1] or ""
        c.execute("""
        UPDATE styles 
        SET fabric_article = '', fabric_composition = '', fabric_type = '', updated_at = CURRENT_TIMESTAMP 
        WHERE id = ?
        """, (req.style_id,))

    c.execute("DELETE FROM fabric_links WHERE style_id = ? AND (fabric_slot = ? OR (? = 1 AND fabric_slot IS NULL))", (req.style_id, slot, slot))

    user_name = user.get("full_name") or user.get("username")
    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, 'fabric_detachment', ?, '', ?)
    """, (req.style_id, order_id, user.get("user_id"), user_name, f"{old_article} ({old_comp})", f"Kumaş {slot} Ataması İptal Edildi"))

    c.execute("""
    INSERT INTO production_status_logs (style_id, user_name, category, title, message)
    VALUES (?, ?, 'Kumaş', 'Kumaş İptal Edildi / Kaldırıldı', ?)
    """, (req.style_id, user_name, f"Kumaş {slot} ataması kaldırıldı. Önceki: {old_article} ({old_comp})"))

    conn.commit()
    conn.close()
    return {"status": "success", "message": f"Kumaş {slot} ataması kaldırıldı"}

# ----------------- COST CALCULATION ENDPOINTS -----------------

def check_can_view_costs(user: Dict[str, Any]):
    role = user.get("role")
    perms = user.get("permissions") or {}
    if role in ["superadmin", "masterdeveloper", "admin"]:
        return True
    if perms.get("can_view_costs"):
        return True
    raise HTTPException(status_code=403, detail="Maliyet verilerini görüntüleme yetkiniz bulunmamaktadır.")

def check_can_edit_costs(user: Dict[str, Any]):
    role = user.get("role")
    perms = user.get("permissions") or {}
    if role in ["superadmin", "masterdeveloper", "admin"]:
        return True
    if perms.get("can_edit_costs"):
        return True
    raise HTTPException(status_code=403, detail="Maliyet verilerini düzenleme yetkiniz bulunmamaktadır.")

class CostSaveRequest(BaseModel):
    cost_data: Dict[str, Any]
    apply_to_unit_price: bool = False
    calculated_price: Optional[float] = None
    currency: Optional[str] = None

@app.get("/api/styles/{style_id}/cost")
def api_get_style_cost(style_id: int, user: Dict[str, Any] = Depends(require_user)):
    check_can_view_costs(user)
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT cost_data_json, unit_price, currency FROM styles WHERE id = ?", (style_id,))
    row = c.fetchone()
    conn.close()
    if not row:
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    cost_json = row[0]
    data = json.loads(cost_json) if cost_json else None
    return {
        "status": "success",
        "cost_data": data,
        "unit_price": row[1],
        "currency": row[2] or "€"
    }

@app.post("/api/styles/{style_id}/cost")
def api_save_style_cost(style_id: int, req: CostSaveRequest, user: Dict[str, Any] = Depends(require_user)):
    check_can_edit_costs(user)
    conn = get_db()
    c = conn.cursor()
    user_name = user.get("full_name") or user.get("username")

    c.execute("SELECT order_id, unit_price, currency FROM styles WHERE id = ?", (style_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    order_id, old_price, cur_currency = row[0], row[1], row[2]

    cost_str = json.dumps(req.cost_data)
    
    if req.apply_to_unit_price and req.calculated_price is not None:
        c.execute("UPDATE styles SET cost_data_json = ?, unit_price = ?, currency = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                  (cost_str, req.calculated_price, req.currency or cur_currency, style_id))
    else:
        c.execute("UPDATE styles SET cost_data_json = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (cost_str, style_id))

    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, 'cost_calculation', '', ?, 'Maliyet Hesaplama Verileri Güncellendi')
    """, (style_id, order_id, user.get("user_id"), user_name, str(req.calculated_price or '')))

    conn.commit()
    conn.close()
    return {"status": "success", "message": "Maliyet başarıyla kaydedildi"}

class ActualCostSaveRequest(BaseModel):
    actual_cost_data: Dict[str, Any]
    calculated_total: Optional[float] = None
    currency: Optional[str] = "TL"

@app.post("/api/styles/{style_id}/actual-cost")
def api_save_style_actual_cost(style_id: int, req: ActualCostSaveRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    user_name = user.get("full_name") or user.get("username")

    c.execute("SELECT order_id FROM styles WHERE id = ?", (style_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Model bulunamadı.")
    order_id = row[0]

    cost_str = json.dumps(req.actual_cost_data)
    c.execute("UPDATE styles SET actual_cost_data_json = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (cost_str, style_id))

    c.execute("""
    INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
    VALUES (?, ?, ?, ?, 'actual_cost', '', ?, 'Gerçekleşen Maliyet & Fatura Verileri Kaydedildi')
    """, (style_id, order_id, user.get("user_id"), user_name, str(req.calculated_total or '')))

    conn.commit()
    conn.close()
    return {"status": "success", "message": "Gerçekleşen maliyet ve fatura verileri başarıyla kaydedildi."}

# ----------------- FREE COST STUDIES (SERBEST FİYAT ÇALIŞMALARI) -----------------

class FreeCostStudySaveRequest(BaseModel):
    season: Optional[str] = ""
    customer_name: Optional[str] = ""
    brand: Optional[str] = ""
    model_no: Optional[str] = ""
    color: Optional[str] = ""
    fabric_supplier: Optional[str] = ""
    fabric_name: Optional[str] = ""
    fabric_code: Optional[str] = ""
    fabric_color: Optional[str] = ""
    fabric_width: Optional[str] = ""
    cost_data: Dict[str, Any] = {}
    kumas_total_tl: Optional[float] = 0.0
    imalat_total_tl: Optional[float] = 0.0
    unit_cost_tl: Optional[float] = 0.0
    unit_price_target: Optional[float] = 0.0
    target_currency: Optional[str] = "EUR"

@app.get("/api/free-cost-studies")
def api_get_free_cost_studies(q: Optional[str] = None, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1
    
    if q and q.strip():
        term = f"%{q.strip().lower()}%"
        c.execute("""
        SELECT * FROM free_cost_studies
        WHERE (company_id = ? OR company_id IS NULL)
          AND (LOWER(season) LIKE ? OR LOWER(customer_name) LIKE ? OR LOWER(brand) LIKE ? OR LOWER(model_no) LIKE ? OR LOWER(color) LIKE ? OR LOWER(fabric_name) LIKE ? OR LOWER(fabric_supplier) LIKE ?)
        ORDER BY id DESC
        """, (company_id, term, term, term, term, term, term, term))
    else:
        c.execute("""
        SELECT * FROM free_cost_studies
        WHERE (company_id = ? OR company_id IS NULL)
        ORDER BY id DESC
        """, (company_id,))
    
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    
    for r in rows:
        try:
            r["cost_data"] = json.loads(r.get("cost_data_json") or "{}")
        except:
            r["cost_data"] = {}
            
    return {"status": "success", "studies": rows}

@app.post("/api/free-cost-studies")
def api_create_free_cost_study(req: FreeCostStudySaveRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1
    user_name = user.get("full_name") or user.get("username") or "Kullanıcı"
    cost_str = json.dumps(req.cost_data)

    c.execute("""
    INSERT INTO free_cost_studies (
        company_id, season, customer_name, brand, model_no, color,
        fabric_supplier, fabric_name, fabric_code, fabric_color, fabric_width,
        cost_data_json, kumas_total_tl, imalat_total_tl, unit_cost_tl, unit_price_target,
        target_currency, created_by, created_at, updated_at
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
    """, (
        company_id, req.season or "", req.customer_name or "", req.brand or "", req.model_no or "", req.color or "",
        req.fabric_supplier or "", req.fabric_name or "", req.fabric_code or "", req.fabric_color or "", req.fabric_width or "",
        cost_str, req.kumas_total_tl or 0.0, req.imalat_total_tl or 0.0, req.unit_cost_tl or 0.0, req.unit_price_target or 0.0,
        req.target_currency or "EUR", user_name
    ))
    new_id = c.lastrowid
    conn.commit()
    conn.close()
    return {"status": "success", "message": "Serbest maliyet çalışması başarıyla kaydedildi.", "id": new_id}

@app.get("/api/free-cost-studies/export-list-excel")
def api_export_free_cost_list_excel(q: Optional[str] = None, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1

    if q and q.strip():
        term = f"%{q.strip().lower()}%"
        c.execute("""
        SELECT * FROM free_cost_studies
        WHERE (company_id = ? OR company_id IS NULL)
          AND (LOWER(season) LIKE ? OR LOWER(customer_name) LIKE ? OR LOWER(brand) LIKE ? OR LOWER(model_no) LIKE ? OR LOWER(color) LIKE ? OR LOWER(fabric_name) LIKE ? OR LOWER(fabric_supplier) LIKE ?)
        ORDER BY id DESC
        """, (company_id, term, term, term, term, term, term, term))
    else:
        c.execute("""
        SELECT * FROM free_cost_studies
        WHERE (company_id = ? OR company_id IS NULL)
        ORDER BY id DESC
        """, (company_id,))

    rows = [dict(r) for r in c.fetchall()]
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Serbest Maliyet Listesi"
    ws.views.sheetView[0].showGridLines = True

    # Stiller
    f_title = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    f_th = Font(name="Segoe UI", size=10, bold=True, color="000000")
    f_td = Font(name="Segoe UI", size=9, color="1E293B")
    f_td_bold = Font(name="Segoe UI", size=9, bold=True, color="0F172A")
    f_theme = Font(name="Segoe UI", size=9, bold=True, color="0284C7")
    f_green = Font(name="Segoe UI", size=10, bold=True, color="15803D")

    fill_banner = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_th = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    bd_thin = Side(style="thin", color="CBD5E1")
    cell_border = Border(left=bd_thin, right=bd_thin, top=bd_thin, bottom=bd_thin)

    # Başlık Banner
    ws.merge_cells("A1:Q1")
    banner_cell = ws["A1"]
    banner_cell.value = "TEXFLOW  |  SERBEST FİYAT ÇALIŞMALARI LİSTESİ"
    banner_cell.font = f_title
    banner_cell.fill = fill_banner
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    headers = [
        ("No", 8),
        ("Sezon", 12),
        ("Müşteri", 20),
        ("Marka", 16),
        ("Model No", 16),
        ("Renk", 14),
        ("Kumaşçı", 18),
        ("Kalite Adı", 20),
        ("Kalite Kodu", 15),
        ("Kumaş Rengi", 14),
        ("Kumaş Eni", 12),
        ("Kumaş Maliyeti (₺)", 18),
        ("İmalat & Aks. (₺)", 18),
        ("Birim Maliyet (₺)", 18),
        ("Birim Satış Fiyatı", 18),
        ("Para Birimi", 12),
        ("Tarih", 14)
    ]

    ws.row_dimensions[2].height = 24
    fill_yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    f_th_yellow = Font(name="Segoe UI", size=10, bold=True, color="000000")
    for col_idx, (h_text, width) in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=h_text)
        cell.font = f_th_yellow
        cell.fill = fill_yellow
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    curr_row = 3
    num_fmt_curr = "₺ #,##0.00"
    num_fmt_2dec = "#,##0.00"

    for idx, r in enumerate(rows):
        r_fill = fill_zebra if idx % 2 == 1 else None
        ws.row_dimensions[curr_row].height = 20

        vals = [
            (r.get("id"), "center", f_td),
            (r.get("season") or "-", "center", f_td),
            (r.get("customer_name") or "-", "left", f_theme),
            (r.get("brand") or "-", "left", f_td),
            (r.get("model_no") or "-", "left", f_td_bold),
            (r.get("color") or "-", "left", f_td),
            (r.get("fabric_supplier") or "-", "left", f_td),
            (r.get("fabric_name") or "-", "left", f_td),
            (r.get("fabric_code") or "-", "left", f_td),
            (r.get("fabric_color") or "-", "left", f_td),
            (r.get("fabric_width") or "-", "center", f_td),
            (float(r.get("kumas_total_tl") or 0.0), "right", f_td, num_fmt_curr),
            (float(r.get("imalat_total_tl") or 0.0), "right", f_td, num_fmt_curr),
            (float(r.get("unit_cost_tl") or 0.0), "right", f_td_bold, num_fmt_curr),
            (float(r.get("unit_price_target") or 0.0), "right", f_green, num_fmt_2dec),
            (r.get("target_currency") or "EUR", "center", f_td_bold),
            (str(r.get("created_at") or "")[:10], "center", f_td)
        ]

        for col_idx, item in enumerate(vals, start=1):
            val = item[0]
            align = item[1]
            font = item[2]
            c = ws.cell(row=curr_row, column=col_idx, value=val)
            c.font = font
            c.alignment = Alignment(horizontal=align, vertical="center")
            c.border = cell_border
            if len(item) > 3 and item[3]:
                c.number_format = item[3]
            if r_fill:
                c.fill = r_fill

        curr_row += 1

    buf = io.BytesIO()
    wb.save(buf)
    val = buf.getvalue()
    dStr = datetime.now().strftime('%Y%m%d')
    filename = f"Serbest_Maliyet_Listesi_{dStr}.xlsx"
    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Content-Length": str(len(val))
    }
    return Response(
        content=val,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )

@app.post("/api/free-cost-studies/export-study-excel")
def api_export_free_cost_study_excel(payload: Dict[str, Any], user: Dict[str, Any] = Depends(require_user)):
    header = payload.get("header") or {}
    form = payload.get("form") or {}
    rates = form.get("rates") or {}
    cur_rates = form.get("exchangeRates") or {"EUR": 56.0, "USD": 48.5, "GBP": 65.5, "TL": 1.0}
    target_curr = form.get("targetCurrency") or "EUR"
    totals = form.get("totals") or {}
    fabrics = form.get("fabrics") or []
    items = form.get("items") or []

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Maliyet Tablosu"
    ws.views.sheetView[0].showGridLines = True

    # Font ve Dolgu Stilleri
    f_banner = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    f_sec = Font(name="Segoe UI", size=10, bold=True, color="000000")
    f_lbl = Font(name="Segoe UI", size=9, bold=True, color="475569")
    f_val = Font(name="Segoe UI", size=9, bold=False, color="0F172A")
    f_th = Font(name="Segoe UI", size=9, bold=True, color="000000")
    f_td = Font(name="Segoe UI", size=9, color="1E293B")
    f_td_bold = Font(name="Segoe UI", size=9, bold=True, color="0F172A")
    f_total_lbl = Font(name="Segoe UI", size=9, bold=True, color="0F172A")
    f_grand = Font(name="Segoe UI", size=11, bold=True, color="15803D")

    fill_banner = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_sec = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid") # #FFFF00 Sarı Başlık
    fill_subth = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid") # #FEF9C3 Açık Sarı Tablo Başlığı
    fill_meta = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_tot = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid") # Sarımsı Ara Toplam
    fill_grand = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

    bd_thin = Side(style="thin", color="CBD5E1")
    bd_thick = Side(style="medium", color="475569")
    bd_double = Side(style="double", color="1E293B")
    cell_border = Border(left=bd_thin, right=bd_thin, top=bd_thin, bottom=bd_thin)
    tot_border = Border(left=bd_thin, right=bd_thin, top=bd_thick, bottom=bd_double)

    # 1. Başlık
    ws.merge_cells("A1:G1")
    b_cell = ws["A1"]
    b_cell.value = "TEXFLOW  |  SERBEST FİYAT & BİRİM MALİYET HESAPLAMA FORMU"
    b_cell.font = f_banner
    b_cell.fill = fill_banner
    b_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # 2. Model & Kumaş Künyesi
    ws.merge_cells("A2:G2")
    s1 = ws["A2"]
    s1.value = "MODEL VE KUMAŞ KÜNYESİ"
    s1.font = f_sec
    s1.fill = fill_sec
    s1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 20

    meta_pairs = [
        [("Müşteri", header.get("customer_name") or "-"), ("Marka", header.get("brand") or "-"), ("Sezon", header.get("season") or "-")],
        [("Model No", header.get("model_no") or "-"), ("Model Rengi", header.get("color") or "-"), ("Döviz Kurları", f"EUR: {cur_rates.get('EUR', 56.0)} | USD: {cur_rates.get('USD', 48.5)} | GBP: {cur_rates.get('GBP', 65.5)}")],
        [("Kumaşçı", header.get("fabric_supplier") or "-"), ("Kalite Adı", header.get("fabric_name") or "-"), ("Kalite Kodu", header.get("fabric_code") or "-")],
        [("Kumaş Rengi", header.get("fabric_color") or "-"), ("Kumaş Eni", header.get("fabric_width") or "-"), ("Tarih", form.get("rates_snapshot_date") or datetime.now().strftime('%d.%m.%Y'))]
    ]

    r_idx = 3
    for pair_row in meta_pairs:
        ws.row_dimensions[r_idx].height = 19
        c_lbl1 = ws.cell(row=r_idx, column=1, value=pair_row[0][0])
        c_lbl1.font = f_lbl
        c_lbl1.fill = fill_meta
        c_lbl1.border = cell_border
        
        c_val1 = ws.cell(row=r_idx, column=2, value=pair_row[0][1])
        c_val1.font = f_val
        c_val1.border = cell_border

        c_lbl2 = ws.cell(row=r_idx, column=3, value=pair_row[1][0])
        c_lbl2.font = f_lbl
        c_lbl2.fill = fill_meta
        c_lbl2.border = cell_border

        c_val2 = ws.cell(row=r_idx, column=4, value=pair_row[1][1])
        c_val2.font = f_val
        c_val2.border = cell_border

        c_lbl3 = ws.cell(row=r_idx, column=5, value=pair_row[2][0])
        c_lbl3.font = f_lbl
        c_lbl3.fill = fill_meta
        c_lbl3.border = cell_border

        ws.merge_cells(start_row=r_idx, start_column=6, end_row=r_idx, end_column=7)
        c_val3 = ws.cell(row=r_idx, column=6, value=pair_row[2][1])
        c_val3.font = f_val
        c_val3.border = cell_border
        ws.cell(row=r_idx, column=7).border = cell_border
        r_idx += 1

    r_idx += 1

    # 3. Kumaş Maliyetleri Tablosu
    ws.merge_cells(f"A{r_idx}:G{r_idx}")
    s2 = ws[f"A{r_idx}"]
    s2.value = f"1. KUMAŞ MALİYETLERİ ({len(fabrics)} Kumaş)"
    s2.font = f_sec
    s2.fill = fill_sec
    s2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[r_idx].height = 20
    r_idx += 1

    fab_cols = [
        ("Kumaş Tanımı", 26),
        ("Kalite / Kod", 20),
        ("Birim", 10),
        ("Sarfiyat (M/Gr)", 15),
        ("Birim Fiyat", 14),
        ("Fire %", 10),
        ("Tutar (₺)", 16)
    ]
    ws.row_dimensions[r_idx].height = 20
    for col_i, (c_name, c_w) in enumerate(fab_cols, start=1):
        c = ws.cell(row=r_idx, column=col_i, value=c_name)
        c.font = f_th
        c.fill = fill_subth
        c.alignment = Alignment(horizontal="center" if col_i in (3,6) else ("right" if col_i in (4,5,7) else "left"), vertical="center")
        c.border = cell_border
    r_idx += 1

    kumas_subtot = 0.0
    for fb in fabrics:
        ws.row_dimensions[r_idx].height = 19
        title = fb.get("title") or "Kumaş"
        desc = f"{fb.get('quality') or ''} {fb.get('code') or ''}".strip() or "-"
        ut = fb.get("unit_type") or "M"
        sarf = float(fb.get("grams") or 0.0) if ut == "KG" else float(fb.get("meters") or 0.0)
        pr = float(fb.get("price") or 0.0)
        curr = fb.get("currency") or "EUR"
        wastage = float(fb.get("wastage_percent") or 0.0)
        tot_tl = float(fb.get("total_tl") or 0.0)
        kumas_subtot += tot_tl

        ws.cell(row=r_idx, column=1, value=title).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r_idx, column=2, value=desc).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r_idx, column=3, value=ut).alignment = Alignment(horizontal="center", vertical="center")
        
        c4 = ws.cell(row=r_idx, column=4, value=sarf)
        c4.number_format = "#,##0.00"
        c4.alignment = Alignment(horizontal="right", vertical="center")

        c5 = ws.cell(row=r_idx, column=5, value=f"{pr:.2f} {curr}")
        c5.alignment = Alignment(horizontal="right", vertical="center")

        c6 = ws.cell(row=r_idx, column=6, value=f"%{wastage:.0f}")
        c6.alignment = Alignment(horizontal="center", vertical="center")

        c7 = ws.cell(row=r_idx, column=7, value=tot_tl)
        c7.number_format = "₺ #,##0.00"
        c7.alignment = Alignment(horizontal="right", vertical="center")
        c7.font = f_td_bold

        for col_i in range(1, 8):
            c_item = ws.cell(row=r_idx, column=col_i)
            c_item.border = cell_border
            if col_i != 7:
                c_item.font = f_td
        r_idx += 1

    # Kumaş Ara Toplam
    ws.row_dimensions[r_idx].height = 20
    ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=6)
    c_tot1 = ws.cell(row=r_idx, column=1, value="KUMAŞ ARA TOPLAMI (₺):")
    c_tot1.font = f_total_lbl
    c_tot1.alignment = Alignment(horizontal="right", vertical="center")
    c_tot1.fill = fill_tot
    c_tot1_v = ws.cell(row=r_idx, column=7, value=totals.get("toplam_kumas_tl") or kumas_subtot)
    c_tot1_v.font = f_td_bold
    c_tot1_v.number_format = "₺ #,##0.00"
    c_tot1_v.alignment = Alignment(horizontal="right", vertical="center")
    c_tot1_v.fill = fill_tot
    for col_i in range(1, 8):
        ws.cell(row=r_idx, column=col_i).border = tot_border
    r_idx += 2

    # 4. İmalat, Aksesuar ve Operasyon Kalemleri
    ws.merge_cells(f"A{r_idx}:G{r_idx}")
    s3 = ws[f"A{r_idx}"]
    s3.value = f"2. İMALAT, AKSESUAR VE OPERASYON KALEMLERİ ({len(items)} Kalem)"
    s3.font = f_sec
    s3.fill = fill_sec
    s3.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[r_idx].height = 20
    r_idx += 1

    item_cols = [
        ("Kalem Adı", 26),
        ("Miktar", 12),
        ("Birim Fiyat", 14),
        ("Para Birimi", 12),
        ("Açıklama / Formül", 24),
        ("Kur", 10),
        ("Tutar (₺)", 16)
    ]
    ws.row_dimensions[r_idx].height = 20
    for col_i, (c_name, c_w) in enumerate(item_cols, start=1):
        c = ws.cell(row=r_idx, column=col_i, value=c_name)
        c.font = f_th
        c.fill = fill_subth
        c.alignment = Alignment(horizontal="center" if col_i in (2,4,6) else ("right" if col_i in (3,7) else "left"), vertical="center")
        c.border = cell_border
    r_idx += 1

    imalat_subtot = 0.0
    for it in items:
        ws.row_dimensions[r_idx].height = 19
        name = it.get("name") or "Kalem"
        qty = float(it.get("qty") or 1.0)
        pr = float(it.get("price") or 0.0)
        curr = it.get("currency") or "TL"
        note = it.get("note") or ""
        tot_tl = float(it.get("total_tl") or 0.0)
        imalat_subtot += tot_tl

        ws.cell(row=r_idx, column=1, value=name).alignment = Alignment(horizontal="left", vertical="center")
        c_q = ws.cell(row=r_idx, column=2, value=qty)
        c_q.alignment = Alignment(horizontal="center", vertical="center")
        c_q.number_format = "#,##0.00"

        c_p = ws.cell(row=r_idx, column=3, value=pr)
        c_p.alignment = Alignment(horizontal="right", vertical="center")
        c_p.number_format = "#,##0.00"

        ws.cell(row=r_idx, column=4, value=curr).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=5, value=note).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r_idx, column=6, value=cur_rates.get(curr, 1.0)).alignment = Alignment(horizontal="center", vertical="center")

        c_tot = ws.cell(row=r_idx, column=7, value=tot_tl)
        c_tot.number_format = "₺ #,##0.00"
        c_tot.alignment = Alignment(horizontal="right", vertical="center")
        c_tot.font = f_td_bold

        for col_i in range(1, 8):
            c_item = ws.cell(row=r_idx, column=col_i)
            c_item.border = cell_border
            if col_i != 7:
                c_item.font = f_td
        r_idx += 1

    # İmalat Ara Toplam & Fire
    ws.row_dimensions[r_idx].height = 20
    ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=6)
    c_tot2 = ws.cell(row=r_idx, column=1, value=f"İMALAT & AKSESUAR TOPLAMI (+%{rates.get('imalat_fire', 5)} Fire Dahil) (₺):")
    c_tot2.font = f_total_lbl
    c_tot2.alignment = Alignment(horizontal="right", vertical="center")
    c_tot2.fill = fill_tot
    c_tot2_v = ws.cell(row=r_idx, column=7, value=totals.get("toplam_imalat_tl") or imalat_subtot)
    c_tot2_v.font = f_td_bold
    c_tot2_v.number_format = "₺ #,##0.00"
    c_tot2_v.alignment = Alignment(horizontal="right", vertical="center")
    c_tot2_v.fill = fill_tot
    for col_i in range(1, 8):
        ws.cell(row=r_idx, column=col_i).border = tot_border
    r_idx += 2

    # 5. Maliyet, Kâr, Komisyon ve Nihai Satış Fiyatı Matrisi
    ws.merge_cells(f"A{r_idx}:G{r_idx}")
    s4 = ws[f"A{r_idx}"]
    s4.value = "3. ORANLAR, KÂR MARJI, KOMİSYON VE SATIŞ FİYATI"
    s4.font = f_sec
    s4.fill = fill_sec
    s4.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[r_idx].height = 20
    r_idx += 1

    summary_rows = [
        ("Maliyet 2 (Kumaş + Net İmalat)", totals.get("maliyet_2_tl") or 0.0, totals.get("maliyet_2_doviz") or 0.0),
        (f"KDV (+%{rates.get('kdv', 10)})", totals.get("tkdv_tl") or 0.0, totals.get("tkdv_doviz") or 0.0),
        ("Maliyet 1 (Maliyet 2 + KDV)", totals.get("maliyet_1_tl") or 0.0, totals.get("maliyet_1_doviz") or 0.0),
        (f"Genel Gider (+%{rates.get('gg', 15)})", totals.get("gg_tl") or 0.0, totals.get("gg_doviz") or 0.0),
        (f"Kâr Marjı / MUP (+%{rates.get('mup', 15)})", totals.get("mup_tl") or 0.0, totals.get("mup_doviz") or 0.0),
        (f"Komisyon (+%{rates.get('komis', 0)})", totals.get("komis_tl") or 0.0, totals.get("komis_doviz") or 0.0),
        ("BİRİM SATIŞ FİYATI", totals.get("toplam_satis_tl") or 0.0, totals.get("toplam_satis_doviz") or 0.0)
    ]

    for idx, (label, tl_val, doviz_val) in enumerate(summary_rows):
        is_grand = idx == len(summary_rows) - 1
        ws.row_dimensions[r_idx].height = 24 if is_grand else 20
        ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=4)
        c_l = ws.cell(row=r_idx, column=1, value=label)
        c_l.font = f_grand if is_grand else f_total_lbl
        c_l.alignment = Alignment(horizontal="right", vertical="center")
        if is_grand:
            c_l.fill = fill_grand

        ws.cell(row=r_idx, column=5, value="TL:").alignment = Alignment(horizontal="right", vertical="center")
        c_tl = ws.cell(row=r_idx, column=6, value=tl_val)
        c_tl.number_format = "₺ #,##0.00"
        c_tl.alignment = Alignment(horizontal="right", vertical="center")
        c_tl.font = f_grand if is_grand else f_td_bold
        if is_grand:
            c_tl.fill = fill_grand

        c_dov = ws.cell(row=r_idx, column=7, value=doviz_val)
        c_dov.number_format = f"[{target_curr}] #,##0.00" if target_curr != "TL" else "₺ #,##0.00"
        c_dov.alignment = Alignment(horizontal="right", vertical="center")
        c_dov.font = f_grand if is_grand else f_td_bold
        if is_grand:
            c_dov.fill = fill_grand

        for col_i in range(1, 8):
            ws.cell(row=r_idx, column=col_i).border = tot_border if is_grand else cell_border

        r_idx += 1

    # Sütun Genişlikleri Ayarı
    for col_i, (c_name, c_w) in enumerate(fab_cols, start=1):
        col_letter = get_column_letter(col_i)
        ws.column_dimensions[col_letter].width = max(c_w, 16)

    buf = io.BytesIO()
    wb.save(buf)
    val = buf.getvalue()
    m_no = (header.get("model_no") or "Model").replace(" ", "_")
    c_name = (header.get("customer_name") or "Musteri").replace(" ", "_")
    dStr = datetime.now().strftime('%Y%m%d')
    filename = f"Maliyet_Tablosu_{c_name}_{m_no}_{dStr}.xlsx"
    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Content-Length": str(len(val))
    }
    return Response(
        content=val,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )

@app.get("/api/free-cost-studies/{study_id}")
def api_get_free_cost_study(study_id: int, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT * FROM free_cost_studies WHERE id = ?", (study_id,))
    row = c.fetchone()
    conn.close()
    if not row:
        raise HTTPException(status_code=404, detail="Maliyet çalışması bulunamadı.")
    res = dict(row)
    try:
        res["cost_data"] = json.loads(res.get("cost_data_json") or "{}")
    except:
        res["cost_data"] = {}
    return {"status": "success", "study": res}

@app.put("/api/free-cost-studies/{study_id}")
def api_update_free_cost_study(study_id: int, req: FreeCostStudySaveRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT id FROM free_cost_studies WHERE id = ?", (study_id,))
    if not c.fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Maliyet çalışması bulunamadı.")
    
    cost_str = json.dumps(req.cost_data)
    c.execute("""
    UPDATE free_cost_studies SET
        season = ?, customer_name = ?, brand = ?, model_no = ?, color = ?,
        fabric_supplier = ?, fabric_name = ?, fabric_code = ?, fabric_color = ?, fabric_width = ?,
        cost_data_json = ?, kumas_total_tl = ?, imalat_total_tl = ?, unit_cost_tl = ?, unit_price_target = ?,
        target_currency = ?, updated_at = CURRENT_TIMESTAMP
    WHERE id = ?
    """, (
        req.season or "", req.customer_name or "", req.brand or "", req.model_no or "", req.color or "",
        req.fabric_supplier or "", req.fabric_name or "", req.fabric_code or "", req.fabric_color or "", req.fabric_width or "",
        cost_str, req.kumas_total_tl or 0.0, req.imalat_total_tl or 0.0, req.unit_cost_tl or 0.0, req.unit_price_target or 0.0,
        req.target_currency or "EUR", study_id
    ))
    conn.commit()
    conn.close()
    return {"status": "success", "message": "Serbest maliyet çalışması başarıyla güncellendi."}

@app.delete("/api/free-cost-studies/{study_id}")
def api_delete_free_cost_study(study_id: int, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    c.execute("DELETE FROM free_cost_studies WHERE id = ?", (study_id,))
    conn.commit()
    conn.close()
    return {"status": "success", "message": "Maliyet çalışması silindi."}
# ----------------- AUDIT LOGS -----------------



@app.get("/api/audit-logs")
def api_get_audit_logs(style_id: Optional[int] = Query(None), limit: int = 100, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1

    if style_id:
        c.execute("""
        SELECT a.*, s.style_no, s.color_name, o.po_number, o.customer_name, o.brand
        FROM audit_logs a
        JOIN styles s ON a.style_id = s.id
        JOIN orders o ON s.order_id = o.id
        WHERE a.style_id = ?
        ORDER BY a.id DESC LIMIT ?
        """, (style_id, limit))
    else:
        c.execute("""
        SELECT a.*, s.style_no, s.color_name, o.po_number, o.customer_name, o.brand
        FROM audit_logs a
        JOIN styles s ON a.style_id = s.id
        JOIN orders o ON s.order_id = o.id
        WHERE s.company_id = ?
        ORDER BY a.id DESC LIMIT ?
        """, (company_id, limit))

    logs = [dict(r) for r in c.fetchall()]
    conn.close()
    return logs

# ----------------- STATUS & MESSAGE LOGS -----------------

@app.get("/api/styles/{style_id}/logs")
def api_get_style_logs(style_id: int, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT * FROM production_status_logs WHERE style_id = ? ORDER BY id DESC", (style_id,))
    logs = [dict(r) for r in c.fetchall()]
    conn.close()
    return logs

class NewLogRequest(BaseModel):
    category: str = "Genel"
    title: Optional[str] = "Durum Notu"
    message: str

@app.post("/api/styles/{style_id}/logs")
def api_add_style_log(style_id: int, req: NewLogRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    c.execute("""
    INSERT INTO production_status_logs (style_id, user_id, user_name, category, title, message)
    VALUES (?, ?, ?, ?, ?, ?)
    """, (style_id, user.get("user_id"), user.get("full_name") or user.get("username"), req.category, req.title, req.message))
    conn.commit()
    conn.close()
    return {"status": "success", "message": "Not eklendi"}

# ----------------- UPLOAD & PARSER -----------------

@app.post("/api/upload/parse-preview")
async def api_upload_and_parse(request: Request, user: Dict[str, Any] = Depends(require_user)):
    form = await request.form()
    upload_list = []
    
    # Check 'files' first, then 'file' without combining both for the same upload
    files_list = form.getlist("files")
    if files_list:
        upload_list = files_list
    else:
        file_single = form.getlist("file")
        if file_single:
            upload_list = file_single
        else:
            for v in form.values():
                if hasattr(v, "filename") and v.filename:
                    upload_list.append(v)


    if not upload_list:
        raise HTTPException(status_code=400, detail="Lütfen en az bir sipariş dosyası seçin.")

    combined_data = []
    file_info = []
    first_type = "MULTI_UPLOAD"

    for f_obj in upload_list:
        if not hasattr(f_obj, "filename") or not f_obj.filename:
            continue
        saved_path = UPLOADS_DIR / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{f_obj.filename}"
        with open(saved_path, "wb") as buffer:
            shutil.copyfileobj(f_obj.file, buffer)
        
        try:
            parsed_result = OrderParserEngine.auto_detect_and_parse(str(saved_path))
            first_type = parsed_result.get("type", "UPLOAD")
            items = parsed_result.get("data", [])
            for it in items:
                it["source_file"] = f_obj.filename
                combined_data.append(it)
            file_info.append(f_obj.filename)
        except Exception as e:
            print(f"Parsing error for {f_obj.filename}: {e}")

    if not combined_data and upload_list:
        raise HTTPException(status_code=400, detail="Seçilen dosyalardan sipariş ayrıştırılamadı.")

    return {
        "type": "MULTI_UPLOAD" if len(file_info) > 1 else first_type,
        "count": len(combined_data),
        "data": combined_data,
        "filename": ", ".join(file_info),
        "files_count": len(file_info)
    }



class ConfirmImportRequest(BaseModel):
    file_path: Optional[str] = ""
    filename: Optional[str] = ""
    source_type: Optional[str] = "MULTI_UPLOAD"
    items: List[Dict[str, Any]]


@app.post("/api/upload/confirm-import")
def api_confirm_import(req: ConfirmImportRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1
    
    inserted_orders = 0
    inserted_styles = 0

    try:
        # Group items by PO Number to attach to unique order parents or create individual ones
        po_groups: Dict[str, List[Any]] = {}
        for item in req.items:
            po_num = str(item.get("po_number") or f"PO-{datetime.now().strftime('%Y%m%d')}").strip()
            po_groups.setdefault(po_num, []).append(item)

        for po_num, items_list in reversed(list(po_groups.items())):
            first_item = items_list[0]
            cust_name = first_item.get("customer_name") or "Müşteri"
            brand = first_item.get("brand") or "Marka"
            season = first_item.get("season") or "2026/2027"
            order_date = first_item.get("order_date") or datetime.now().strftime("%d.%m.%Y")
            delivery_date = first_item.get("delivery_date") or ""
            delivery_terms = first_item.get("delivery_terms") or ""
            payment_terms = first_item.get("payment_terms") or ""
            currency = first_item.get("currency") or "EUR"

            total_q = sum(int(it.get("total_quantity") or 0) for it in items_list)
            total_a = sum(float(it.get("total_quantity") or 0) * float(it.get("unit_price") or 0.0) for it in items_list)

            c.execute("""
            INSERT INTO orders (
                company_id, po_number, customer_name, brand, season, order_date, delivery_date,
                delivery_terms, payment_terms, total_quantity, total_amount, currency, raw_file_name
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                company_id, po_num, cust_name, brand, season, order_date, delivery_date,
                delivery_terms, payment_terms, total_q, total_a, currency, req.filename
            ))
            order_id = c.lastrowid
            inserted_orders += 1

            for it in items_list:
                style_no = str(it.get("style_no") or "STYLE-01").strip()
                desc = it.get("description") or ""
                comp = it.get("fabric_composition") or ""
                article = it.get("fabric_article") or ""
                ftype = it.get("fabric_type") or ""
                color_code = str(it.get("color_code") or "").strip()
                raw_color_name = str(it.get("color_name") or "").strip()
                if color_code and not raw_color_name.startswith(color_code):
                    color_name = f"{color_code} {raw_color_name}".strip()
                else:
                    color_name = raw_color_name or color_code
                price = float(it.get("unit_price") or 0.0)
                
                sizes_dict = it.get("size_distribution") or {}
                # Calculate real total from sizes if present
                if sizes_dict and any(int(v or 0) > 0 for v in sizes_dict.values()):
                    item_qty = sum(int(v or 0) for v in sizes_dict.values())
                else:
                    item_qty = int(it.get("total_quantity") or 0)
                    
                item_amt = item_qty * price

                style_brand = str(it.get("brand") or brand or "").strip()
                style_channel = str(it.get("channel") or "").strip()
                if not style_channel:
                    style_channel = OrderParserEngine.clean_channel_value(it.get("notes") or it.get("comments") or "")

                c.execute("""
                INSERT INTO styles (
                    order_id, company_id, style_no, description, fabric_composition, fabric_article,
                    fabric_type, color_code, color_name, unit_price, total_quantity, total_amount,
                    currency, status, fabric_order_status, last_status_updated_by, last_status_updated_at,
                    brand, channel
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Planlamada', ?, 'Sipariş Yükleme', datetime('now'), ?, ?)
                """, (
                    order_id, company_id, style_no, desc, comp, article,
                    ftype, color_code, color_name, price, item_qty, item_amt,
                    currency, it.get("fabric_order_status") or "",
                    style_brand, style_channel
                ))
                style_id = c.lastrowid
                inserted_styles += 1

                for idx, (sz_name, sz_qty) in enumerate(sizes_dict.items()):
                    if str(sz_name).strip():
                        c.execute("""
                        INSERT INTO size_distributions (style_id, size_name, quantity, sort_order)
                        VALUES (?, ?, ?, ?)
                        """, (style_id, str(sz_name).strip(), int(sz_qty or 0), idx + 1))

                for log in it.get("status_logs", []):
                    c.execute("""
                    INSERT INTO production_status_logs (style_id, user_name, category, title, message)
                    VALUES (?, ?, 'Sipariş Notu', ?, ?)
                    """, (style_id, "Sipariş Yükleme", log.get("title") or "Not", log.get("message") or ""))

        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        raise HTTPException(status_code=500, detail=f"Kayıt hatası: {str(e)}")

    conn.close()
    return {
        "status": "success",
        "message": f"{inserted_orders} Sipariş ve {inserted_styles} Model başarıyla sisteme aktarıldı."
    }

# ----------------- MANUEL SİPARİŞ / MODEL EKLEME -----------------

class CreateManualStyleRequest(BaseModel):
    customer_name: Optional[str] = ""
    brand: Optional[str] = ""
    channel: Optional[str] = ""
    po_number: Optional[str] = ""
    style_no: Optional[str] = ""
    description: Optional[str] = ""
    color_name: Optional[str] = ""
    color_code: Optional[str] = ""
    fabric_article: Optional[str] = ""
    fabric_composition: Optional[str] = ""
    unit_price: Optional[float] = 0.0
    currency: Optional[str] = "EUR"
    unit_meters: Optional[str] = ""
    unit_grams: Optional[str] = ""
    fabric_wastage_percent: Optional[float] = 5.0
    fabric_ordered_meters: Optional[float] = 0.0
    fabric_order_unit: Optional[str] = "M"
    order_date: Optional[str] = ""
    delivery_date: Optional[str] = ""
    status: Optional[str] = "Planlamada"
    size_map: Optional[Dict[str, int]] = None

@app.post("/api/styles/manual-create")
def api_create_manual_style(req: CreateManualStyleRequest, user: Dict[str, Any] = Depends(require_user)):
    conn = get_db()
    c = conn.cursor()
    company_id = user.get("company_id") or 1
    user_name = user.get("full_name") or user.get("username") or "Kullanıcı"

    po = (req.po_number or "MANUEL-PO").strip()
    cust = (req.customer_name or "").strip()
    brand = (req.brand or "").strip()
    channel = (req.channel or "").strip()
    order_dt = req.order_date or datetime.now().strftime("%d.%m.%Y")
    del_dt = req.delivery_date or ""
    curr = req.currency or "EUR"

    sizes_dict = req.size_map or {}
    total_qty = sum(int(v or 0) for v in sizes_dict.values()) if sizes_dict else 0
    unit_price = float(req.unit_price or 0.0)
    total_amount = total_qty * unit_price

    try:
        c.execute("""
        SELECT id FROM orders 
        WHERE company_id = ? AND po_number = ? AND customer_name = ?
        """, (company_id, po, cust))
        row = c.fetchone()
        if row:
            order_id = row[0]
            c.execute("""
            UPDATE orders 
            SET total_quantity = total_quantity + ?, total_amount = total_amount + ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """, (total_qty, total_amount, order_id))
        else:
            c.execute("""
            INSERT INTO orders (
                company_id, po_number, customer_name, brand, order_date, delivery_date,
                total_quantity, total_amount, currency, status, raw_file_name
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'Yeni', 'Manuel Giriş')
            """, (company_id, po, cust, brand, order_dt, del_dt, total_qty, total_amount, curr))
            order_id = c.lastrowid

        style_no = (req.style_no or "STYLE-01").strip()
        c.execute("""
        INSERT INTO styles (
            order_id, company_id, style_no, description, fabric_composition, fabric_article,
            color_code, color_name, unit_price, currency, total_quantity, total_amount,
            unit_meters, unit_grams, fabric_wastage_percent, fabric_ordered_meters, fabric_order_unit,
            status, last_status_updated_by, last_status_updated_at, channel
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, datetime('now'), ?)
        """, (
            order_id, company_id, style_no, req.description or "", req.fabric_composition or "",
            req.fabric_article or "", req.color_code or "", req.color_name or "", unit_price, curr,
            total_qty, total_amount, req.unit_meters or "", req.unit_grams or "",
            req.fabric_wastage_percent or 5.0, req.fabric_ordered_meters or 0.0, req.fabric_order_unit or "M",
            req.status or "Planlamada", user_name, channel
        ))
        style_id = c.lastrowid


        for idx, (sz_name, sz_qty) in enumerate(sizes_dict.items()):
            if str(sz_name).strip():
                c.execute("""
                INSERT INTO size_distributions (style_id, size_name, quantity, sort_order)
                VALUES (?, ?, ?, ?)
                """, (style_id, str(sz_name).strip(), int(sz_qty or 0), idx + 1))

        c.execute("""
        INSERT INTO audit_logs (style_id, order_id, user_id, user_name, field_name, old_value, new_value, note)
        VALUES (?, ?, ?, ?, 'manual_create', '', ?, 'Manuel Sipariş / Model Eklendi')
        """, (style_id, order_id, user.get("user_id"), user_name, f"{style_no} - {req.color_name or ''} ({total_qty} Adet)"))

        c.execute("""
        INSERT INTO production_status_logs (style_id, user_name, category, title, message)
        VALUES (?, ?, 'Sipariş Notu', 'Manuel Giriş Yapıldı', ?)
        """, (style_id, user_name, f"Model {style_no} manuel olarak oluşturuldu."))

        conn.commit()
        conn.close()
        return {"status": "success", "style_id": style_id, "order_id": order_id}
    except Exception as e:
        conn.rollback()
        conn.close()
        raise HTTPException(status_code=500, detail=f"Manuel sipariş kaydetme hatası: {str(e)}")



# ----------------- FABRICTAG ENDPOINTS -----------------

@app.get("/api/fabrictag/fabrics")
def api_fabrictag_list(q: Optional[str] = Query(None), user: Dict[str, Any] = Depends(require_user)):
    return FabricTagConnector.search_fabrics(q or "")

class UpdateFabricReceivedRequest(BaseModel):
    received_meters: Optional[str] = ""
    received_date: Optional[str] = None

@app.post("/api/fabrictag/fabrics/{fabric_id}/received")
def api_fabrictag_update_received(fabric_id: int, req: UpdateFabricReceivedRequest, user: Dict[str, Any] = Depends(require_user)):
    from datetime import datetime
    date_str = req.received_date
    if not date_str and req.received_meters and req.received_meters.strip():
        date_str = datetime.now().strftime("%d.%m.%Y")
    ok = FabricTagConnector.update_received_meters(fabric_id, req.received_meters or "", date_str)
    if not ok:
        raise HTTPException(status_code=500, detail="Kumaş gelen miktarı güncellenemedi.")
    return {"status": "success", "received_meters": req.received_meters, "received_date": date_str}

class AddFabricBatchRequest(BaseModel):
    amount: str
    date: Optional[str] = None
    note: Optional[str] = ""
    unit_price: Optional[str] = ""
    currency: Optional[str] = "EUR"
    exchange_rate: Optional[str] = ""
    vat_rate: Optional[float] = 10.0

@app.post("/api/fabrictag/fabrics/{fabric_id}/batches")
def api_fabrictag_add_batch(fabric_id: int, req: AddFabricBatchRequest, user: Dict[str, Any] = Depends(require_user)):
    res = FabricTagConnector.add_received_batch(
        fabric_id, req.amount, req.date, req.note or "",
        unit_price=req.unit_price or "",
        currency=req.currency or "EUR",
        exchange_rate=req.exchange_rate or "",
        vat_rate=req.vat_rate if req.vat_rate is not None else 10.0
    )
    if not res:
        raise HTTPException(status_code=500, detail="Parti eklenemedi.")
    return {"status": "success", **res}

class UpdateFabricBatchRequest(BaseModel):
    amount: Optional[str] = None
    date: Optional[str] = None
    note: Optional[str] = None
    unit_price: Optional[str] = None
    currency: Optional[str] = None
    exchange_rate: Optional[str] = None
    vat_rate: Optional[float] = None

@app.put("/api/fabrictag/fabrics/{fabric_id}/batches/{batch_id}")
def api_fabrictag_update_batch(fabric_id: int, batch_id: int, req: UpdateFabricBatchRequest, user: Dict[str, Any] = Depends(require_user)):
    res = FabricTagConnector.update_received_batch(fabric_id, batch_id, req.dict(exclude_unset=True))
    if not res:
        raise HTTPException(status_code=500, detail="Parti güncellenemedi.")
    return {"status": "success", **res}

@app.delete("/api/fabrictag/fabrics/{fabric_id}/batches/{batch_id}")
def api_fabrictag_delete_batch(fabric_id: int, batch_id: int, user: Dict[str, Any] = Depends(require_user)):
    res = FabricTagConnector.delete_received_batch(fabric_id, batch_id)
    if not res:
        raise HTTPException(status_code=500, detail="Parti silinemedi.")
    return {"status": "success", **res}


def parse_date_sort_key(dt_str: str):
    """Robust date parsing for chronological sorting (earliest to latest)."""
    if not dt_str:
        return "0000-00-00 00:00"
    s = str(dt_str).strip()
    try:
        # Check DD.MM.YYYY HH:MM or DD.MM.YYYY
        if "." in s:
            parts = s.split(" ")
            dparts = parts[0].split(".")
            if len(dparts) == 3:
                day, month, year = dparts[0].zfill(2), dparts[1].zfill(2), dparts[2]
                time_part = parts[1] if len(parts) > 1 else "00:00"
                return f"{year}-{month}-{day} {time_part}"
        # Check YYYY-MM-DD
        if "-" in s:
            return s
    except Exception:
        pass
    return s


def get_fabric_ledger_data(fabric_id: int, company_id: int = 1):
    """Compiles chronological stock ledger with running balance, incoming purchases, and cutting consumption."""
    ft_conn = FabricTagConnector.get_connection()
    if not ft_conn:
        return None
    ft_c = ft_conn.cursor()
    ft_c.execute("""
        SELECT id, internal_code, company_name, quality_code, quality_name, 
               design_code, color, width, weight, composition, received_meters, 
               received_date, received_batches_json, barcode_or_qr
        FROM fabrics WHERE id = ?
    """, (fabric_id,))
    frow = ft_c.fetchone()
    ft_conn.close()
    if not frow:
        return None

    fabric = dict(frow)
    int_code = (fabric.get("internal_code") or "").strip().upper()
    q_code = (fabric.get("quality_code") or "").strip().lower()

    # 1. Gelen Kumaş Partileri
    batches = []
    b_json = fabric.get("received_batches_json")
    if b_json:
        try:
            batches = json.loads(b_json)
        except Exception:
            batches = []
    elif fabric.get("received_meters"):
        batches = [{
            "id": 1,
            "amount": fabric.get("received_meters"),
            "date": fabric.get("received_date") or "",
            "note": "1. Giriş / Devir Stoku",
            "unit_price": "",
            "currency": "EUR",
            "exchange_rate": "",
            "vat_rate": 10.0
        }]

    movements = []
    for b in batches:
        try:
            amt_raw = str(b.get("amount", 0)).replace(",", ".").replace("M", "").replace("m", "").replace("KG", "").replace("kg", "").strip()
            amt = float(amt_raw)
        except Exception:
            amt = 0.0
        if amt <= 0:
            continue

        b_id = b.get("id")
        dt = b.get("date") or ""
        note = b.get("note") or "Kumaş Girişi"
        u_price_str = str(b.get("unit_price") or "").strip()
        curr = str(b.get("currency") or "EUR").strip()
        ex_rate_str = str(b.get("exchange_rate") or "").strip()
        vat_r = float(b.get("vat_rate") if b.get("vat_rate") is not None else 10.0)

        price_f = 0.0
        try:
            price_f = float(u_price_str.replace(",", ".")) if u_price_str else 0.0
        except Exception:
            price_f = 0.0

        rate_f = 1.0
        try:
            if ex_rate_str:
                rate_f = float(ex_rate_str.replace(",", "."))
            elif curr in ["₺", "TL"]:
                rate_f = 1.0
        except Exception:
            rate_f = 1.0

        # Tutar formülü: Metraj * Birim Fiyat * Kur * (1 + KDV/100)
        # Eğer kur veya fiyat girilmemişse tutar 0
        base_amt = amt * price_f * rate_f
        total_tl = base_amt * (1.0 + (vat_r / 100.0)) if price_f > 0 else 0.0

        movements.append({
            "type": "gelen",
            "id": b_id,
            "batch_id": b_id,
            "date": dt,
            "sort_key": parse_date_sort_key(dt),
            "amount": amt,
            "note": note,
            "unit_price": u_price_str,
            "currency": curr,
            "exchange_rate": ex_rate_str,
            "vat_rate": vat_r,
            "total_amount": round(total_tl, 2)
        })

    # 2. Kesilen Kumaşlar (TexFlow styles tablosundan)
    conn = get_db()
    c = conn.cursor()
    c.execute("""
        SELECT s.id, s.order_id, s.style_no, s.color_name, s.cut_meters, s.cutting_date,
               s.created_at, s.fabric_article, s.fabric_article_2,
               s.fabric_type, s.fabric_type_2,
               o.po_number, o.customer_name, o.brand, o.order_date
        FROM styles s
        LEFT JOIN orders o ON s.order_id = o.id
        WHERE s.cut_meters IS NOT NULL AND s.cut_meters != '' AND s.cut_meters != '0'
    """)
    all_cuts = [dict(r) for r in c.fetchall()]

    # fabric_links tablosu
    fl_map = {}
    if all_cuts:
        c.execute("SELECT style_id, fabrictag_id, fabric_code FROM fabric_links")
        for fl in c.fetchall():
            sid = fl["style_id"]
            if sid not in fl_map:
                fl_map[sid] = []
            fl_map[sid].append(dict(fl))

    for s in all_cuts:
        try:
            cut_val = float(str(s.get("cut_meters") or 0).replace(",", "."))
        except Exception:
            cut_val = 0.0
        if cut_val <= 0:
            continue

        sid = s["id"]
        links = fl_map.get(sid, [])
        is_match = False

        # 1. fabrictag_id match
        for lk in links:
            if lk.get("fabrictag_id") == fabric_id:
                is_match = True
                break
            if int_code and str(lk.get("fabric_code") or "").strip().upper() == int_code:
                is_match = True
                break

        # 2. internal_code in fabric fields
        if not is_match and int_code:
            fa1 = (s.get("fabric_article") or s.get("fabric_type") or "").upper()
            fa2 = (s.get("fabric_article_2") or s.get("fabric_type_2") or "").upper()
            if int_code in fa1 or int_code in fa2:
                is_match = True

        # 3. quality_code in fabric fields
        if not is_match and q_code and len(q_code) >= 3 and q_code != "kodsuz":
            fa1_low = (s.get("fabric_article") or s.get("fabric_type") or "").lower()
            fa2_low = (s.get("fabric_article_2") or s.get("fabric_type_2") or "").lower()
            if q_code in fa1_low or q_code in fa2_low:
                is_match = True

        if is_match:
            c_date = s.get("cutting_date") or s.get("order_date") or (str(s.get("created_at") or "")[:10]) or ""
            st_no = s.get("style_no") or ""
            col_nm = s.get("color_name") or ""
            po_num = s.get("po_number") or ""
            cust_nm = s.get("customer_name") or s.get("brand") or ""
            
            note_parts = [f"Kesim: Model {st_no}"]
            if col_nm:
                note_parts.append(col_nm)
            if po_num:
                note_parts.append(f"PO: {po_num}")
            if cust_nm:
                note_parts.append(f"({cust_nm})")

            movements.append({
                "type": "kesilen",
                "id": f"cut_{sid}",
                "style_id": sid,
                "date": c_date,
                "sort_key": parse_date_sort_key(c_date),
                "amount": cut_val,
                "note": " - ".join(note_parts),
                "unit_price": "",
                "currency": "",
                "exchange_rate": "",
                "vat_rate": "",
                "total_amount": 0.0
            })

    # 3. Kronolojik Sıralama: En eskiden en yeniye doğru
    # Aynı tarihte gelenler önce, kesilenler sonra sıralansın
    movements.sort(key=lambda m: (m["sort_key"], 0 if m["type"] == "gelen" else 1))

    # 4. Yürüyen Bakiye Hesaplama
    running_balance = 0.0
    total_received = 0.0
    total_cut = 0.0
    total_spend_tl = 0.0

    for m in movements:
        if m["type"] == "gelen":
            running_balance += m["amount"]
            total_received += m["amount"]
            total_spend_tl += m.get("total_amount", 0.0)
            m["gelen"] = m["amount"]
            m["kesilen"] = 0.0
        else:
            running_balance -= m["amount"]
            total_cut += m["amount"]
            m["gelen"] = 0.0
            m["kesilen"] = m["amount"]
        
        m["balance"] = round(running_balance, 2)

    return {
        "fabric": fabric,
        "movements": movements,
        "summary": {
            "total_received": round(total_received, 2),
            "total_cut": round(total_cut, 2),
            "remaining": round(running_balance, 2),
            "total_spend_tl": round(total_spend_tl, 2)
        }
    }


@app.get("/api/fabrictag/fabrics/{fabric_id}/ledger")
def api_get_fabric_ledger(fabric_id: int, user: Dict[str, Any] = Depends(require_user)):
    data = get_fabric_ledger_data(fabric_id, company_id=user.get("company_id", 1))
    if not data:
        raise HTTPException(status_code=404, detail="Kumaş kaydı bulunamadı.")
    return {"status": "success", **data}


@app.get("/api/reports/fabric-ledger-excel/{fabric_id}")
def api_export_fabric_ledger_excel(fabric_id: int, user: Dict[str, Any] = Depends(require_user)):
    data = get_fabric_ledger_data(fabric_id, company_id=user.get("company_id", 1))
    if not data:
        raise HTTPException(status_code=404, detail="Kumaş kaydı bulunamadı.")

    fabric = data["fabric"]
    movements = data["movements"]
    summary = data["summary"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stok Hareketleri"
    ws.views.sheetView[0].showGridLines = True

    # Font ve Stiller
    f_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    f_meta_label = Font(name="Segoe UI", size=9, bold=True, color="475569")
    f_meta_val = Font(name="Segoe UI", size=9, bold=False, color="0F172A")
    f_th = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    f_td = Font(name="Segoe UI", size=9, color="1E293B")
    f_td_bold = Font(name="Segoe UI", size=9, bold=True, color="0F172A")
    f_total = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

    fill_banner = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_th = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_meta = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    # KPI kutuları dolguları
    fill_kpi_rec = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")   # Sky
    fill_kpi_cut = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")   # Amber
    fill_kpi_rem = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")   # Emerald
    fill_kpi_ttr = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")   # Indigo

    bd_thin = Side(style="thin", color="CBD5E1")
    bd_thick = Side(style="medium", color="475569")
    bd_double = Side(style="double", color="1E293B")
    cell_border = Border(left=bd_thin, right=bd_thin, top=bd_thin, bottom=bd_thin)
    total_border = Border(left=bd_thin, right=bd_thin, top=bd_thick, bottom=bd_double)

    # 1. Başlık Banner (Row 2, A2:I2)
    ws.merge_cells("A2:I2")
    cell_b = ws["A2"]
    cell_b.value = "TEXFLOW  |  KUMAŞ STOK HAREKETLERİ & DETAYLI EKSTRE"
    cell_b.font = f_title
    cell_b.fill = fill_banner
    cell_b.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 36

    # 2. Kumaş Künyesi (Row 4-5)
    f_code = fabric.get("internal_code") or "-"
    f_comp_name = fabric.get("company_name") or "-"
    f_qcode = fabric.get("quality_code") or "-"
    f_qname = fabric.get("quality_name") or "-"
    f_col = fabric.get("color") or "-"
    f_des = fabric.get("design_code") or "-"
    f_comp = fabric.get("composition") or "-"
    f_w = fabric.get("width") or "-"
    f_g = fabric.get("weight") or "-"

    meta_line1 = f"İç Kod: {f_code}  |  Firma: {f_comp_name}  |  Kalite Kodu: {f_qcode}  |  Kalite Adı: {f_qname}"
    meta_line2 = f"Desen: {f_des}  |  Renk: {f_col}  |  Kompozisyon: {f_comp}  |  En: {f_w}  |  Gramaj: {f_g}  |  Rapor: {datetime.now().strftime('%d.%m.%Y %H:%M')}"

    ws.merge_cells("A4:I4")
    ws["A4"].value = meta_line1
    ws["A4"].font = f_meta_label
    ws["A4"].fill = fill_meta
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells("A5:I5")
    ws["A5"].value = meta_line2
    ws["A5"].font = f_meta_val
    ws["A5"].fill = fill_meta
    ws["A5"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 20

    # 3. KPI Özet Kutuları (Row 7)
    # A7:B7 Toplam Gelen, C7:D7 Toplam Kesilen, E7:F7 Kalan Stok, G7:I7 Toplam Tutar
    ws.merge_cells("A7:B7")
    ws["A7"].value = f"Toplam Gelen: {summary['total_received']:,.2f} m".replace(",", "X").replace(".", ",").replace("X", ".")
    ws["A7"].font = Font(name="Segoe UI", size=10, bold=True, color="0369A1")
    ws["A7"].fill = fill_kpi_rec
    ws["A7"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("C7:D7")
    ws["C7"].value = f"Toplam Kesilen: {summary['total_cut']:,.2f} m".replace(",", "X").replace(".", ",").replace("X", ".")
    ws["C7"].font = Font(name="Segoe UI", size=10, bold=True, color="B45309")
    ws["C7"].fill = fill_kpi_cut
    ws["C7"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("E7:F7")
    ws["E7"].value = f"Kalan Stok: {summary['remaining']:,.2f} m".replace(",", "X").replace(".", ",").replace("X", ".")
    ws["E7"].font = Font(name="Segoe UI", size=10, bold=True, color="15803D" if summary['remaining'] >= 0 else "B91C1C")
    ws["E7"].fill = fill_kpi_rem
    ws["E7"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("G7:I7")
    ws["G7"].value = f"Toplam Alış: ₺ {summary['total_spend_tl']:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    ws["G7"].font = Font(name="Segoe UI", size=10, bold=True, color="4338CA")
    ws["G7"].fill = fill_kpi_ttr
    ws["G7"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[7].height = 26

    # 4. Tablo Başlıkları (Row 9)
    headers = [
        "Tarih",
        "Hareket Türü / Açıklama",
        "Gelen Kumaş (m)",
        "Kesilen Kumaş (m)",
        "Kalan Kumaş (m)",
        "Birim Fiyat",
        "Döviz Kuru",
        "KDV (%)",
        "Tutar (₺)"
    ]
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=9, column=col_idx, value=h_text)
        cell.font = f_th
        cell.fill = fill_th
        cell.alignment = Alignment(
            horizontal="center" if col_idx in [1, 7, 8] else ("right" if col_idx in [3, 4, 5, 6, 9] else "left"),
            vertical="center"
        )
        cell.border = cell_border
    ws.row_dimensions[9].height = 24

    # 5. Hareket Satırları
    current_row = 10
    num_fmt_meters = "#,##0.00"
    num_fmt_curr = "₺ #,##0.00"

    for idx, m in enumerate(movements):
        r_fill = fill_zebra if idx % 2 == 1 else None

        c1 = ws.cell(row=current_row, column=1, value=m.get("date") or "-")
        c1.alignment = Alignment(horizontal="center", vertical="center")

        c2 = ws.cell(row=current_row, column=2, value=m.get("note") or "-")
        c2.alignment = Alignment(horizontal="left", vertical="center")

        # Gelen (m)
        c3 = ws.cell(row=current_row, column=3, value=m.get("gelen", 0.0) if m.get("gelen", 0.0) > 0 else "")
        c3.number_format = num_fmt_meters
        c3.alignment = Alignment(horizontal="right", vertical="center")
        if m.get("gelen", 0.0) > 0:
            c3.font = Font(name="Segoe UI", size=9, bold=True, color="0369A1")

        # Kesilen (m)
        c4 = ws.cell(row=current_row, column=4, value=m.get("kesilen", 0.0) if m.get("kesilen", 0.0) > 0 else "")
        c4.number_format = num_fmt_meters
        c4.alignment = Alignment(horizontal="right", vertical="center")
        if m.get("kesilen", 0.0) > 0:
            c4.font = Font(name="Segoe UI", size=9, bold=True, color="B45309")

        # Kalan (m)
        c5 = ws.cell(row=current_row, column=5, value=m.get("balance", 0.0))
        c5.number_format = num_fmt_meters
        c5.alignment = Alignment(horizontal="right", vertical="center")
        bal_color = "15803D" if m.get("balance", 0.0) >= 0 else "B91C1C"
        c5.font = Font(name="Segoe UI", size=9, bold=True, color=bal_color)

        # Birim Fiyat & Kur
        p_val = m.get("unit_price")
        p_curr = m.get("currency") or ""
        if p_val:
            try:
                c6_num = float(str(p_val).replace(",", "."))
                c6 = ws.cell(row=current_row, column=6, value=f"{p_curr} {c6_num:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
            except Exception:
                c6 = ws.cell(row=current_row, column=6, value=f"{p_curr} {p_val}")
        else:
            c6 = ws.cell(row=current_row, column=6, value="-")
        c6.alignment = Alignment(horizontal="right", vertical="center")

        # Kur
        ex_val = m.get("exchange_rate")
        if ex_val:
            try:
                c7_num = float(str(ex_val).replace(",", "."))
                c7 = ws.cell(row=current_row, column=7, value=c7_num)
                c7.number_format = num_fmt_meters
            except Exception:
                c7 = ws.cell(row=current_row, column=7, value=ex_val)
        else:
            c7 = ws.cell(row=current_row, column=7, value="1,00" if p_curr in ["₺", "TL"] else "-")
        c7.alignment = Alignment(horizontal="center", vertical="center")

        # KDV (%)
        vat_val = m.get("vat_rate")
        if vat_val is not None and vat_val != "" and m.get("type") == "gelen":
            c8 = ws.cell(row=current_row, column=8, value=f"%{int(vat_val) if float(vat_val).is_integer() else vat_val}")
        else:
            c8 = ws.cell(row=current_row, column=8, value="-")
        c8.alignment = Alignment(horizontal="center", vertical="center")

        # Tutar (₺)
        tot_val = m.get("total_amount", 0.0)
        c9 = ws.cell(row=current_row, column=9, value=tot_val if tot_val > 0 else "")
        c9.number_format = num_fmt_curr
        c9.alignment = Alignment(horizontal="right", vertical="center")
        if tot_val > 0:
            c9.font = Font(name="Segoe UI", size=9, bold=True, color="4338CA")

        for col_i in range(1, 10):
            c_item = ws.cell(row=current_row, column=col_i)
            c_item.border = cell_border
            if r_fill and not c_item.fill.start_color.rgb:
                c_item.fill = r_fill

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # 6. Alt Toplam Satırı (Total Row)
    c_tot_label = ws.cell(row=current_row, column=1, value="GENEL TOPLAM")
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
    c_tot_label.font = f_total
    c_tot_label.alignment = Alignment(horizontal="center", vertical="center")

    c_tot_rec = ws.cell(row=current_row, column=3, value=summary["total_received"])
    c_tot_rec.font = Font(name="Segoe UI", size=10, bold=True, color="0369A1")
    c_tot_rec.number_format = num_fmt_meters
    c_tot_rec.alignment = Alignment(horizontal="right", vertical="center")

    c_tot_cut = ws.cell(row=current_row, column=4, value=summary["total_cut"])
    c_tot_cut.font = Font(name="Segoe UI", size=10, bold=True, color="B45309")
    c_tot_cut.number_format = num_fmt_meters
    c_tot_cut.alignment = Alignment(horizontal="right", vertical="center")

    c_tot_rem = ws.cell(row=current_row, column=5, value=summary["remaining"])
    c_tot_rem.font = Font(name="Segoe UI", size=10, bold=True, color="15803D" if summary["remaining"] >= 0 else "B91C1C")
    c_tot_rem.number_format = num_fmt_meters
    c_tot_rem.alignment = Alignment(horizontal="right", vertical="center")

    ws.cell(row=current_row, column=6, value="")
    ws.cell(row=current_row, column=7, value="")
    ws.cell(row=current_row, column=8, value="")

    c_tot_spend = ws.cell(row=current_row, column=9, value=summary["total_spend_tl"])
    c_tot_spend.font = Font(name="Segoe UI", size=10, bold=True, color="4338CA")
    c_tot_spend.number_format = num_fmt_curr
    c_tot_spend.alignment = Alignment(horizontal="right", vertical="center")

    for col_i in range(1, 10):
        cell_t = ws.cell(row=current_row, column=col_i)
        cell_t.fill = fill_total
        cell_t.border = total_border

    ws.row_dimensions[current_row].height = 24

    # Kolon genişlikleri
    col_widths = {
        1: 14,  # Tarih
        2: 36,  # Açıklama / Belge / Model
        3: 17,  # Gelen Kumaş (m)
        4: 17,  # Kesilen Kumaş (m)
        5: 17,  # Kalan Kumaş (m)
        6: 15,  # Birim Fiyat
        7: 13,  # Döviz Kuru
        8: 11,  # KDV (%)
        9: 20   # Tutar (₺)
    }
    for c_idx, width in col_widths.items():
        col_letter = get_column_letter(c_idx)
        ws.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    val = buf.getvalue()
    filename = f"Kumas_Stok_Ekstresi_{f_code or fabric_id}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Content-Length": str(len(val))
    }
    return Response(
        content=val,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )



# ----------------- SUPERADMIN ENDPOINTS -----------------

@app.get("/api/superadmin/companies")
def api_superadmin_get_companies(user: Dict[str, Any] = Depends(require_superadmin)):
    conn = get_db()
    c = conn.cursor()
    c.execute("""
    SELECT c.*, 
           (SELECT COUNT(*) FROM users WHERE company_id = c.id) as user_count,
           (SELECT COUNT(*) FROM orders WHERE company_id = c.id) as order_count
    FROM companies c
    ORDER BY c.id DESC
    """)
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    return rows

class CreateCompanyRequest(BaseModel):
    name: str
    code: str
    license_key: str
    license_expires_at: str
    max_users: int = 10

@app.post("/api/superadmin/companies")
def api_superadmin_create_company(req: CreateCompanyRequest, user: Dict[str, Any] = Depends(require_superadmin)):
    conn = get_db()
    c = conn.cursor()
    try:
        c.execute("""
        INSERT INTO companies (name, code, license_key, license_expires_at, is_active, max_users)
        VALUES (?, ?, ?, ?, 1, ?)
        """, (req.name, req.code, req.license_key, req.license_expires_at, req.max_users))
        cid = c.lastrowid
        
        c.execute("SELECT menu_key, title_tr, title_en, icon, path, sort_order, is_visible, allowed_roles_json FROM dynamic_menus WHERE company_id = 1")
        menus = c.fetchall()
        for m in menus:
            c.execute("""
            INSERT INTO dynamic_menus (company_id, menu_key, title_tr, title_en, icon, path, sort_order, is_visible, allowed_roles_json)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (cid, m[0], m[1], m[2], m[3], m[4], m[5], m[6], m[7]))
            
        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        raise HTTPException(status_code=400, detail=f"Firma oluşturulamadı: {str(e)}")
# ----------------- HAFTALIK PROGRAM EXCEL ENDPOINT -----------------

@app.get("/api/reports/haftalik-program-excel")
def api_haftalik_program_excel(
    start_key: Optional[str] = Query(None),
    end_key: Optional[str] = Query(None),
    only_with_orders: bool = Query(True),
    hide_headers: bool = Query(False),
    show_images: bool = Query(True)
):
    """
    Patron Takip Çıktısı için A4 Yatay 7'şer model sütunlu Haftalık Program Excel üreticisi
    """
    from datetime import timedelta
    from collections import defaultdict

    conn = get_db()
    c = conn.cursor()
    c.execute("""
        SELECT 
            s.id, s.style_no, s.brand, s.description,
            o.customer_name, o.brand as order_brand,
            s.color_code, s.color_name, s.total_quantity,
            o.delivery_date, s.image_url, s.image_url_2
        FROM styles s
        LEFT JOIN orders o ON s.order_id = o.id
        WHERE o.delivery_date IS NOT NULL AND o.delivery_date != ''
        ORDER BY o.delivery_date ASC, s.id ASC
    """)
    rows = c.fetchall()
    conn.close()

    tr_months = ['Oca', 'Şub', 'Mar', 'Nis', 'May', 'Haz', 'Tem', 'Ağu', 'Eyl', 'Eki', 'Kas', 'Ara']
    week_map = {}

    for r in rows:
        raw_date = str(r[9]).strip()
        dt = None
        for fmt in ('%d.%m.%Y', '%Y-%m-%d', '%d/%m/%Y'):
            try:
                dt = datetime.strptime(raw_date, fmt)
                break
            except Exception:
                pass
        if not dt:
            continue

        year, week, _ = dt.isocalendar()
        key = f"{year}_{week:02d}"

        if start_key and key < start_key:
            continue
        if end_key and key > end_key:
            continue

        if key not in week_map:
            jan4 = datetime(year, 1, 4)
            start_of_year = jan4 - timedelta(days=jan4.isoweekday() - 1)
            monday = start_of_year + timedelta(weeks=week - 1)
            sunday = monday + timedelta(days=6)
            label = f"{monday.day} {tr_months[monday.month - 1]} - {sunday.day} {tr_months[sunday.month - 1]} {year}"

            week_map[key] = {
                'key': key,
                'year': year,
                'week': week,
                'label': label,
                'models': {}
            }

        style_no = (r[1] or 'Bilinmeyen').strip()
        brand = (r[2] or r[5] or r[4] or '').strip()
        color = (r[7] or r[6] or 'Standart').strip()
        qty = int(r[8] or 0)
        img = r[10] or r[11] or ''

        w = week_map[key]
        if style_no not in w['models']:
            w['models'][style_no] = {
                'style_no': style_no,
                'brand': brand,
                'image_url': img,
                'colors': defaultdict(int),
                'total_quantity': 0
            }

        m = w['models'][style_no]
        if brand and not m['brand']:
            m['brand'] = brand
        if img and not m['image_url']:
            m['image_url'] = img
        m['colors'][color] += qty
        m['total_quantity'] += qty

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Haftalık Program'

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    try:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    except Exception:
        pass
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35

    # Sol tarafta 55px boşluk (Sütun A: ~55px)
    ws.column_dimensions['A'].width = 7.0
    for c_idx in range(2, 16):
        ws.column_dimensions[get_column_letter(c_idx)].width = 14.0

    thin_black = Side(style='thin', color='000000')
    thick_black = Side(style='medium', color='000000')

    current_row = 1
    sorted_weeks = sorted(week_map.values(), key=lambda x: x['key'])

    for w in sorted_weeks:
        models_list = list(w['models'].values())
        if not models_list and only_with_orders:
            continue

        week_total_qty = sum(m['total_quantity'] for m in models_list)
        week_total_models = len(models_list)

        for chunk_idx in range(0, max(1, len(models_list)), 7):
            chunk = models_list[chunk_idx:chunk_idx + 7]

            if chunk_idx == 0 and not hide_headers:
                ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=15)
                h_cell = ws.cell(current_row, 2)
                h_cell.value = f"{w['week']}. HAFTA ({w['label']})   •   Toplam {week_total_models} Model   •   {week_total_qty:,} Adet"
                h_cell.font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
                h_cell.fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
                h_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
                ws.row_dimensions[current_row].height = 26.0
                current_row += 1

            brand_row = current_row
            model_row = current_row + 1
            img_start_row = current_row + 2
            img_end_row = current_row + 7
            ws.row_dimensions[brand_row].height = 22.0
            ws.row_dimensions[model_row].height = 22.0
            for ir in range(img_start_row, img_end_row + 1):
                ws.row_dimensions[ir].height = 20.0

            max_colors = 2
            for m in chunk:
                max_colors = max(max_colors, len(m['colors']))

            for slot_idx in range(7):
                col_left = 2 + (slot_idx * 2)
                col_right = col_left + 1
                col_left_let = get_column_letter(col_left)

                if slot_idx < len(chunk):
                    model = chunk[slot_idx]
                    # Brand row (12pt)
                    ws.merge_cells(start_row=brand_row, start_column=col_left, end_row=brand_row, end_column=col_right)
                    b_cell = ws.cell(brand_row, col_left)
                    b_cell.value = (model['brand'] or '').upper()
                    b_cell.font = Font(name='Calibri', size=12, bold=True, color='1E293B')
                    b_cell.fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
                    b_cell.alignment = Alignment(horizontal='center', vertical='center')

                    # Model row (12pt)
                    ws.merge_cells(start_row=model_row, start_column=col_left, end_row=model_row, end_column=col_right)
                    m_cell = ws.cell(model_row, col_left)
                    m_cell.value = model['style_no']
                    m_cell.font = Font(name='Calibri', size=12, bold=True, color='0F172A')
                    m_cell.fill = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
                    m_cell.alignment = Alignment(horizontal='center', vertical='center')

                    # Image area
                    ws.merge_cells(start_row=img_start_row, start_column=col_left, end_row=img_end_row, end_column=col_right)
                    img_cell = ws.cell(img_start_row, col_left)
                    img_cell.alignment = Alignment(horizontal='center', vertical='center')

                    if show_images and model['image_url']:
                        img_path = UPLOADS_DIR / 'styles' / Path(model['image_url']).name
                        if img_path.exists():
                            try:
                                pil_img = PILImage.open(img_path)
                                orig_w, orig_h = pil_img.size
                                if orig_w > 0 and orig_h > 0:
                                    # Hücre kutusu: ~204px genişlik, ~120pt = 160px yükseklik
                                    # En ve boydan hangisi sınırlandırıyorsa o oranda büyüt / sığdır
                                    scale = min(184.0 / orig_w, 144.0 / orig_h)
                                    new_w = max(1, int(orig_w * scale))
                                    new_h = max(1, int(orig_h * scale))

                                    openpyxl_img = OpenpyxlImage(str(img_path))
                                    openpyxl_img.width = new_w
                                    openpyxl_img.height = new_h

                                    # Hem enden hem boydan tam ortala
                                    offset_x_px = max(0, int((204 - new_w) / 2))
                                    offset_y_px = max(0, int((160 - new_h) / 2))

                                    marker = AnchorMarker(
                                        col=col_left - 1,
                                        colOff=pixels_to_EMU(offset_x_px),
                                        row=img_start_row - 1,
                                        rowOff=pixels_to_EMU(offset_y_px)
                                    )
                                    size = XDRPositiveSize2D(pixels_to_EMU(new_w), pixels_to_EMU(new_h))
                                    openpyxl_img.anchor = OneCellAnchor(_from=marker, ext=size)
                                    ws.add_image(openpyxl_img)
                            except Exception:
                                pass

                    # Colors and quantities (12pt, ince siyah iç çizgiler)
                    colors_list = list(model['colors'].items())
                    for c_idx in range(max_colors):
                        c_row = img_end_row + 1 + c_idx
                        ws.row_dimensions[c_row].height = 20.0
                        c_col_cell = ws.cell(c_row, col_left)
                        c_qty_cell = ws.cell(c_row, col_right)

                        if c_idx < len(colors_list):
                            c_name, c_qty = colors_list[c_idx]
                            c_col_cell.value = c_name
                            c_col_cell.font = Font(name='Calibri', size=12, color='000000')
                            c_col_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)

                            c_qty_cell.value = c_qty
                            c_qty_cell.font = Font(name='Calibri', size=12, bold=True, color='000000')
                            c_qty_cell.alignment = Alignment(horizontal='center', vertical='center')
                        else:
                            c_col_cell.value = ''
                            c_qty_cell.value = ''

                    # Total row (12pt)
                    tot_row = img_end_row + 1 + max_colors
                    ws.row_dimensions[tot_row].height = 22.0
                    tot_label_cell = ws.cell(tot_row, col_left)
                    tot_label_cell.value = 'TOPLAM'
                    tot_label_cell.font = Font(name='Calibri', size=12, bold=True, color='047857')
                    tot_label_cell.fill = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')
                    tot_label_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)

                    tot_val_cell = ws.cell(tot_row, col_right)
                    tot_val_cell.value = model['total_quantity']
                    tot_val_cell.font = Font(name='Calibri', size=12, bold=True, color='047857')
                    tot_val_cell.fill = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')
                    tot_val_cell.alignment = Alignment(horizontal='center', vertical='center')

                else:
                    for r_i in range(brand_row, img_end_row + 1 + max_colors + 1):
                        ws.cell(r_i, col_left).value = None
                        ws.cell(r_i, col_right).value = None

                slot_end_row = img_end_row + 1 + max_colors
                tot_row = slot_end_row
                for r_idx in range(brand_row, slot_end_row + 1):
                    for cl in (col_left, col_right):
                        top_s = thick_black if (r_idx == brand_row or r_idx == tot_row) else thin_black
                        bot_s = thick_black if r_idx == slot_end_row else thin_black
                        left_s = thick_black if cl == col_left else thin_black
                        right_s = thick_black if cl == col_right else thin_black
                        ws.cell(r_idx, cl).border = Border(top=top_s, bottom=bot_s, left=left_s, right=right_s)

            current_row = img_end_row + 1 + max_colors + 2

    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    filename = f"Haftalik_Program_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return Response(
        content=output_stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

# ----------------- SPA ROOT -----------------

@app.get("/")
def serve_index():
    index_path = STATIC_DIR / "index.html"
    if index_path.exists():
        return FileResponse(
            str(index_path),
            headers={
                "Cache-Control": "no-cache, no-store, must-revalidate, max-age=0",
                "Pragma": "no-cache",
                "Expires": "0"
            }
        )
    return HTMLResponse("<h1>TexFlow API is running.</h1>")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8050)
